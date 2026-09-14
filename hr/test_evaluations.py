import copy
from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook

from core.models import Role, OrganizationalUnit
from hr.models import (Employee, EvaluationGroup, EvaluationCriterion, EvaluationScale, EvaluationUnitSetup,
    EvaluationEmployeeSetup, EmployeeEvaluation, EvaluationApproval)
from hr.services.evaluations import (import_evaluation_catalog, build_snapshot, sign_snapshot, read_snapshot,
    save_evaluation, approve_evaluation, approval_state, visible_evaluations, editable_employees, cyrillic)


def workbook_fixture(path):
    book = Workbook()
    scale = book.active
    scale.title = 'Max_vrednosti'
    scale.append(['menadzment','merilo','bodovi','max_vrednost'])
    groups = {
        0:[[.075,.045,.02,0,-.02]]*2+[[.04,.03,.02,0,-.02]]*2+[[.035,.02,.01,0,-.02]]*2,
        1:[[.125,.07,.035,0,-.035]]*2+[[.07,.055,.035,0,-.035]]*2+[[.055,.035,.02,0,-.035]]*2,
    }
    for management, criteria in groups.items():
        for number, values in enumerate(criteria,1):
            for points,value in zip([4,3,2,1,0],values):
                scale.append([management,number,points,value])
    form = book.create_sheet('obrazac')
    criteria = book.create_sheet('kriterijumi')
    for number in range(1,7):
        form.cell(18+number,1,f'{number}.Мерило {number}')
        for points in range(5):
            row = (4 if number<=3 else 15)+(4-points)
            col = [1,7,13][(number-1)%3]
            criteria.cell(row,col,f'Опис {number}/{points}')
    book.save(path)
    book.close()


class EvaluationTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        with TemporaryDirectory() as directory:
            path = Path(directory)/'fixture.xlsx'
            workbook_fixture(path)
            import_evaluation_catalog(path)
        cls.employee = Employee.objects.create(employee_code=810,first_name='Petar',last_name='Petrović',department_code=100,
            org_unit_code='100',position='Inženjer',education='VII',date_of_birth=date(1990,1,1),date_of_joining=date(2020,1,1))
        cls.people = []
        cls.reviewers = []
        role = Role.objects.get(slug='ocenjivac')
        for index in range(3):
            person = Employee.objects.create(employee_code=811+index,first_name='Ocenjivač',last_name=str(index+1),department_code=100,
                org_unit_code='100',date_of_birth=date(1980,1,1),date_of_joining=date(2010,1,1))
            user = get_user_model().objects.create_user(f'evaluator-{index}',employee=person)
            user.roles.add(role)
            cls.people.append(person)
            cls.reviewers.append(user)
        cls.admin = get_user_model().objects.create_superuser('evaluation-admin','admin@example.test','test')
        cls.other = get_user_model().objects.create_user('unassigned-reviewer')
        cls.other.roles.add(role)
        cls.group = EvaluationGroup.objects.get(code='employee')
        cls.unit = EvaluationUnitSetup.objects.create(unit_code='100',supervisor=cls.people[0],director=cls.people[1],general_director=cls.people[2])
        EvaluationEmployeeSetup.objects.create(employee=cls.employee,group=cls.group)
        OrganizationalUnit.objects.create(code='100',name='Primer jedinice',center='10')

    def snapshot(self,group=None):
        return build_snapshot(employee=self.employee,group=group or self.group,year=2026,month=8,
            supervisor=self.people[0],director=self.people[1],general_director=self.people[2])

    def save(self,points=1,source=None,user=None):
        user = user or self.admin
        snapshot = copy.deepcopy(source.snapshot) if source else self.snapshot()
        token = sign_snapshot(snapshot,user,source.pk if source else None)
        envelope = read_snapshot(token,user)
        values = {f'criterion_{item["code"]}':str(points) for item in snapshot['items']}
        return save_evaluation(envelope=envelope,values=values,user=user)

    def approve(self,item,stage,index,**kwargs):
        return approve_evaluation(evaluation=item,stage=stage,user=self.reviewers[index],document_date=date(2026,9,15),**kwargs)

    def test_imports_exact_grids_and_preserves_local_edits(self):
        self.assertEqual(EvaluationCriterion.objects.count(),6)
        self.assertEqual(EvaluationScale.objects.count(),60)
        scale = EvaluationScale.objects.first()
        scale.description = 'Локални опис'
        scale.save()
        with TemporaryDirectory() as directory:
            path = Path(directory)/'fixture.xlsx'
            workbook_fixture(path)
            counts = import_evaluation_catalog(path)
        self.assertEqual(counts['preserved'],60)
        scale.refresh_from_db()
        self.assertEqual(scale.description,'Локални опис')

    def test_incomplete_catalog_rejects_whole_import(self):
        with TemporaryDirectory() as directory:
            path = Path(directory)/'fixture.xlsx'
            workbook_fixture(path)
            from openpyxl import load_workbook
            book = load_workbook(path)
            book['Max_vrednosti'].delete_rows(61)
            book.save(path)
            book.close()
            with self.assertRaises(ValidationError):
                import_evaluation_catalog(path)
        self.assertEqual(EvaluationScale.objects.count(),60)

    def test_neutral_is_one_point_and_maxima_match_workbook(self):
        snapshot = self.snapshot()
        self.assertEqual(Decimal(snapshot['maximum']),Decimal('1.30000'))
        self.assertEqual(Decimal(snapshot['minimum']),Decimal('.88000'))
        self.assertTrue(all(item['points']==1 and Decimal(item['value'])==0 for item in snapshot['items']))
        management = self.snapshot(EvaluationGroup.objects.get(code='management'))
        self.assertEqual(Decimal(management['maximum']),Decimal('1.50000'))
        self.assertEqual(Decimal(management['minimum']),Decimal('.79000'))
        self.assertEqual(self.save().personal_coefficient,Decimal('1'))
        self.assertEqual(self.save(points=4).personal_coefficient,Decimal('1.3'))
        self.assertEqual(self.save(points=0).personal_coefficient,Decimal('.88'))

    def test_catalog_employee_and_unit_changes_do_not_change_saved_snapshot_or_print(self):
        item = self.save(points=4)
        before = copy.deepcopy(item.snapshot)
        EvaluationScale.objects.update(coefficient=Decimal('0'))
        EvaluationCriterion.objects.update(name='Промењено мерило')
        Employee.objects.filter(pk=self.employee.pk).update(first_name='Changed',position='Changed')
        self.client.force_login(self.admin)
        response = self.client.get(reverse('hr:evaluation_print',args=[item.pk]))
        self.assertContains(response,'Петар Петровић')
        self.assertContains(response,'1,30000')
        self.assertNotContains(response,'Changed')
        self.assertNotContains(response,'Промењено мерило')
        item.refresh_from_db()
        self.assertEqual(item.snapshot,before)

    def test_signed_form_survives_catalog_change_and_rejects_tampering_or_other_user(self):
        token = sign_snapshot(self.snapshot(),self.admin)
        EvaluationScale.objects.update(coefficient=Decimal('0'))
        envelope = read_snapshot(token,self.admin)
        values = {f'criterion_{i}':'4' for i in range(1,7)}
        item = save_evaluation(envelope=envelope,values=values,user=self.admin)
        self.assertEqual(item.personal_coefficient,Decimal('1.3'))
        with self.assertRaises(ValidationError):
            read_snapshot(token+'corrupt',self.admin)
        with self.assertRaises(ValidationError):
            read_snapshot(token,self.other)

    def test_repeated_post_is_idempotent_and_new_revision_preserves_history(self):
        envelope = read_snapshot(sign_snapshot(self.snapshot(),self.admin),self.admin)
        values = {f'criterion_{i}':'1' for i in range(1,7)}
        first = save_evaluation(envelope=envelope,values=values,user=self.admin)
        again = save_evaluation(envelope=envelope,values=values,user=self.admin)
        self.assertEqual(first.pk,again.pk)
        self.approve(first,'supervisor',0)
        second = self.save(points=4,source=first)
        self.assertEqual(second.revision,2)
        self.assertEqual(second.approvals.count(),0)
        self.assertEqual(first.approvals.count(),1)
        self.assertEqual(first.personal_coefficient,Decimal('1'))

    def test_only_named_reviewers_can_approve_in_order(self):
        item = self.save()
        with self.assertRaises(PermissionDenied):
            approve_evaluation(evaluation=item,stage='supervisor',user=self.admin,document_date=date(2026,9,15))
        with self.assertRaises(ValidationError):
            self.approve(item,'director',1)
        self.approve(item,'supervisor',0)
        with self.assertRaises(ValidationError):
            self.approve(item,'supervisor',0)
        self.approve(item,'director',1,coefficient=Decimal('1.2'),comment='Korekcija')
        self.approve(item,'general_director',2)
        approvals,value = approval_state(item)
        self.assertEqual(value,Decimal('1.2'))
        self.assertEqual(len(approvals),3)
        self.assertEqual(approvals['director'].comment,'Корекција')

    def test_director_override_requires_comment_and_respects_snapshot_limits(self):
        item = self.save()
        self.approve(item,'supervisor',0)
        with self.assertRaises(ValidationError):
            self.approve(item,'director',1,coefficient=Decimal('1.2'))
        with self.assertRaises(ValidationError):
            self.approve(item,'director',1,coefficient=Decimal('9'),comment='Previše')
        self.assertEqual(item.approvals.count(),1)

    def test_old_revision_cannot_receive_new_approvals(self):
        item = self.save()
        self.save(source=item)
        with self.assertRaises(ValidationError):
            self.approve(item,'supervisor',0)

    def test_scope_blocks_unassigned_employee_and_snapshot_reads(self):
        item = self.save()
        self.assertTrue(visible_evaluations(self.reviewers[0]).filter(pk=item.pk).exists())
        self.assertFalse(visible_evaluations(self.other).exists())
        self.assertTrue(editable_employees(self.reviewers[0]).filter(pk=self.employee.pk).exists())
        self.assertFalse(editable_employees(self.other).exists())
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(reverse('hr:evaluation_detail',args=[item.pk])).status_code,404)
        self.assertEqual(self.client.get(reverse('hr:evaluation_print',args=[item.pk])).status_code,404)
        self.assertEqual(self.client.get(reverse('hr:evaluation_catalog')).status_code,403)

    def test_create_flow_prefills_setup_and_saves_values(self):
        self.client.force_login(self.admin)
        url = reverse('hr:evaluation_create')
        response = self.client.get(url,{'employee':self.employee.pk})
        self.assertEqual(response.context['header_form'].initial['group'],self.group.pk)
        self.assertEqual(response.context['header_form'].initial['director'],self.people[1].pk)
        response = self.client.get(url,{'employee':self.employee.pk,'year':2026,'month':8,'prepare':1})
        self.assertEqual(response.status_code,200)
        token = response.context['values_form'].fields['snapshot_token'].initial
        data = {'snapshot_token':token,**{f'criterion_{i}':'4' for i in range(1,7)}}
        response = self.client.post(url,data)
        self.assertEqual(response.status_code,302)
        item = EmployeeEvaluation.objects.get()
        self.assertEqual(item.personal_coefficient,Decimal('1.3'))
        self.assertEqual(self.client.get(response.url).status_code,200)

    def test_catalog_edit_and_all_tabs_render(self):
        self.client.force_login(self.admin)
        for kind in ['groups','criteria','scales','units','employees']:
            self.assertEqual(self.client.get(reverse('hr:evaluation_catalog'),{'kind':kind}).status_code,200)
            self.assertEqual(self.client.get(reverse('hr:evaluation_catalog_create',args=[kind])).status_code,200)
        scale = EvaluationScale.objects.filter(group=self.group,criterion__code=1,points=4).get()
        response = self.client.post(reverse('hr:evaluation_catalog_edit',args=['scales',scale.pk]),{
            'group':self.group.pk,'criterion':scale.criterion_id,'points':4,'coefficient':'0,07000','description':'Ново','is_active':'on'})
        self.assertEqual(response.status_code,302)
        scale.refresh_from_db()
        self.assertEqual(scale.coefficient,Decimal('.07'))

    def test_monthly_list_keeps_latest_revision(self):
        first = self.save()
        second = self.save(points=4,source=first)
        self.client.force_login(self.admin)
        response = self.client.get(reverse('hr:evaluation_list'),{'year':2026,'month':8})
        self.assertEqual([item.pk for item in response.context['evaluations']],[second.pk])

    def test_bulk_approval_is_atomic_when_one_selected_item_is_not_ready(self):
        first = self.save()
        # Another month is a distinct current evaluation; the second lacks supervisor consent.
        snapshot = self.snapshot()
        snapshot['month'] = 7
        second = save_evaluation(envelope=read_snapshot(sign_snapshot(snapshot,self.admin),self.admin),values={f'criterion_{i}':'1' for i in range(1,7)},user=self.admin)
        self.approve(first,'supervisor',0)
        self.client.force_login(self.reviewers[1])
        response = self.client.post(reverse('hr:evaluation_bulk_approve'),{'evaluations':[first.pk,second.pk],
            'stage':'director','document_date':'2026-09-15','confirm':'on'})
        self.assertEqual(response.status_code,302)
        self.assertFalse(EvaluationApproval.objects.filter(stage='director').exists())

    def test_cyrillic_transliteration_preserves_existing_cyrillic(self):
        self.assertEqual(cyrillic('Ljiljana Petrović — ИМС'),'Љиљана Петровић — ИМС')
