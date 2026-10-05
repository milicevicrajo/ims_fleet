"""Spisak zaposlenih: klik na šifru i ime vodi na detalj, štampa (A4/PDF) i Excel prate filtere ekrana."""
import datetime
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from hr.models import Employee


def zaposleni(sifra, prezime, aktivan=True, oj="430"):
    return Employee.objects.create(employee_code=sifra, first_name="Ana", last_name=prezime, position="Inženjer",
                                   department_code=430, org_unit_code=oj, gender="F", is_active=aktivan,
                                   date_of_birth=datetime.date(1990, 5, 4), date_of_joining=datetime.date(2020, 1, 2))


class SpisakZaposlenihTests(TestCase):
    def setUp(self):
        self.anic, self.bozic = zaposleni(1, "Anić"), zaposleni(2, "Božić", oj="410")
        self.bivsi = zaposleni(3, "Bivši", aktivan=False)
        self.client.force_login(get_user_model().objects.create_superuser("kadrovik", password="x"))

    def test_ime_i_sifra_vode_na_detalj_bez_kolone_detalji(self):
        odgovor = self.client.get(reverse("employee_list"))
        detalj = reverse("employee_detail", args=[self.anic.pk])
        self.assertContains(odgovor, f'href="{detalj}"', count=2)  # šifra i „Prezime i ime”
        self.assertContains(odgovor, f'<a class="hr-detail-link" href="{detalj}">Anić Ana</a>', html=True)
        self.assertContains(odgovor, "<th>Prezime i ime</th>", html=True)
        self.assertNotContains(odgovor, "mdi-account-details")
        self.assertContains(odgovor, "scrollX: false")
        self.assertContains(odgovor, "izvoz=stampa")
        self.assertContains(odgovor, "izvoz=xlsx")

    def test_dugmad_izvoza_cuvaju_filtere(self):
        odgovor = self.client.get(reverse("employee_list"), {"status": "all", "oj": "430"})
        self.assertContains(odgovor, "?status=all&amp;oj=430&amp;izvoz=xlsx")

    def test_stampa_a4_prati_filtere(self):
        odgovor = self.client.get(reverse("employee_list"), {"status": "active", "oj": "430", "izvoz": "stampa"})
        self.assertEqual(odgovor.status_code, 200)
        self.assertContains(odgovor, "size: A4 portrait")
        self.assertContains(odgovor, "Aktivni zaposleni · OJ 430")
        self.assertContains(odgovor, "Anić Ana")
        self.assertNotContains(odgovor, "Božić")
        self.assertNotContains(odgovor, "Bivši")

    def test_excel_ima_zaglavlje_i_redove(self):
        odgovor = self.client.get(reverse("employee_list"), {"status": "all", "izvoz": "xlsx"})
        self.assertEqual(odgovor.status_code, 200)
        self.assertIn("spreadsheetml", odgovor["Content-Type"])
        ws = load_workbook(BytesIO(odgovor.content)).active
        self.assertEqual(ws["A3"].value, "Spisak zaposlenih")
        self.assertEqual([c.value for c in ws[6]][:4], ["R.br.", "Šifra", "Prezime i ime", "Radno mesto"])
        self.assertEqual([ws.cell(row=r, column=3).value for r in (7, 8, 9)], ["Anić Ana", "Bivši Ana", "Božić Ana"])
        self.assertEqual(ws.cell(row=7, column=7).value.date(), datetime.date(1990, 5, 4))
        self.assertEqual(ws.freeze_panes, "A7")

    def test_izvoz_trazi_dozvolu_spiska(self):
        korisnik = get_user_model().objects.create_user("obican", password="x", employee=self.anic)
        self.client.force_login(korisnik)
        self.assertEqual(self.client.get(reverse("employee_list"), {"izvoz": "xlsx"}).status_code, 403)
        self.assertEqual(self.client.get(reverse("employee_list"), {"izvoz": "stampa"}).status_code, 403)
