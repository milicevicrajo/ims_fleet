"""Spisak zaposlenih za štampu (A4) i Excel — isti filteri i redosled kao ekran „Spisak zaposlenih”."""
from django.utils import timezone
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.exporting import create_xlsx_workbook, workbook_response

INSTITUT = 'Institut za ispitivanje materijala a.d.'
ADRESA = 'Bulevar vojvode Mišića 43, 11000 Beograd'
STATUSI = {'active': 'Aktivni zaposleni', 'inactive': 'Neaktivni zaposleni', 'all': 'Svi zaposleni'}

# (naslov, širina kolone u Excelu)
KOLONE = [('R.br.', 6), ('Šifra', 9), ('Prezime i ime', 28), ('Radno mesto', 40), ('OJ', 8), ('Pol', 6),
          ('Datum rođenja', 14), ('Datum zaposlenja', 16), ('Telefon', 18)]

PLAVA, SVETLA, LINIJA = '14395B', 'F4F7FA', 'D7DCE5'


def redovi(employees):
    return [{
        'rb': rb, 'sifra': e.employee_code, 'prezime': e.display_last_name, 'ime': e.display_first_name,
        'radno_mesto': e.position or e.job_title or '', 'oj': str(e.org_unit_code or e.department_code or ''),
        'pol': e.gender or '', 'rodjen': e.date_of_birth, 'zaposlen': e.date_of_joining,
        'telefon': e.phone_number or e.mobile_phone or '',
    } for rb, e in enumerate(employees, start=1)]


def opis_filtera(status, oj='', primalac=None, pretraga=''):
    delovi = [STATUSI.get(status, STATUSI['active'])]
    if oj:
        delovi.append(f'OJ {oj}')
    if primalac:
        delovi.append(f'Vrsta primaoca: {primalac}')
    if pretraga:
        delovi.append(f'Pretraga: „{pretraga}”')
    return ' · '.join(delovi)


def excel(employees, filteri):
    """Excel sa zaglavljem Instituta, obojenim zaglavljem kolona, filterom i zamrznutim naslovom."""
    rows = redovi(employees)
    wb, ws = create_xlsx_workbook('Zaposleni')
    poslednja = get_column_letter(len(KOLONE))
    tanka = Side(style='thin', color=LINIJA)
    okvir = Border(left=tanka, right=tanka, top=tanka, bottom=tanka)

    ws['A1'] = INSTITUT
    ws['A1'].font = Font(bold=True, size=13, color=PLAVA)
    ws['A2'] = ADRESA
    ws['A2'].font = Font(size=9, color='64748B')
    ws['A3'] = 'Spisak zaposlenih'
    ws['A3'].font = Font(bold=True, size=15, color=PLAVA)
    ws['A4'] = f'{filteri} · ukupno {len(rows)} · izrađeno {timezone.localtime():%d.%m.%Y. %H:%M}'
    ws['A4'].font = Font(size=9, italic=True, color='475569')
    for red in range(1, 5):
        ws.merge_cells(f'A{red}:{poslednja}{red}')

    zaglavlje = 6
    for kolona, (naslov, sirina) in enumerate(KOLONE, start=1):
        cell = ws.cell(row=zaglavlje, column=kolona, value=naslov)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=PLAVA)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = okvir
        ws.column_dimensions[get_column_letter(kolona)].width = sirina
    ws.row_dimensions[zaglavlje].height = 30

    zebra = PatternFill('solid', fgColor=SVETLA)
    for i, r in enumerate(rows):
        broj = zaglavlje + 1 + i
        vrednosti = [r['rb'], r['sifra'], f"{r['prezime']} {r['ime']}", r['radno_mesto'], r['oj'], r['pol'],
                     r['rodjen'], r['zaposlen'], r['telefon']]
        for kolona, vrednost in enumerate(vrednosti, start=1):
            cell = ws.cell(row=broj, column=kolona, value=vrednost)
            cell.border = okvir
            cell.alignment = Alignment(vertical='center', wrap_text=kolona == 4,
                                       horizontal='center' if kolona in (1, 2, 5, 6, 7, 8) else 'left')
            if kolona in (7, 8) and vrednost:
                cell.number_format = 'DD.MM.YYYY'
            if i % 2:
                cell.fill = zebra
        ws.cell(row=broj, column=3).font = Font(bold=True)

    ws.freeze_panes = ws.cell(row=zaglavlje + 1, column=1)
    ws.auto_filter.ref = f'A{zaglavlje}:{poslednja}{zaglavlje + max(len(rows), 1)}'
    ws.print_title_rows = f'{zaglavlje}:{zaglavlje}'
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = 'Strana &P od &N'
    return workbook_response(wb, f'spisak_zaposlenih_{timezone.localdate():%Y%m%d}.xlsx', quoted=True)
