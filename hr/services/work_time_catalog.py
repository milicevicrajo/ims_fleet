"""Local, editable work-time dictionaries initialized from the supplied workbook."""
import re
import unicodedata

from django.core.exceptions import ValidationError
from django.db import connections, transaction
from openpyxl import load_workbook

from hr.models import Employee, RecipientType, WorkTimeCategory, WorkTimeElement
from hr.sync import _resolve_hr_db_alias


CATEGORY_LABELS = {
    'redovan_rad': 'Redovan rad',
    'prekovremeni_nocni_rad': 'Prekovremeni noćni rad',
    'prekovremeni_rad': 'Prekovremeni rad',
    'prekovremeni_rad_na_drzavne_praznike': 'Prekovremeni rad na državne praznike',
    'redovan_rad_na_drzavne_praznike': 'Redovan rad na državne praznike',
    'topli_obrok': 'Topli obrok', 'regres': 'Regres', 'minuli_rad': 'Minuli rad',
    'nocni_rad': 'Noćni rad', 'drzavni_i_verski_praznik': 'Državni i verski praznik',
    'placeno_odsustvo': 'Plaćeno odsustvo', 'bolovanje': 'Bolovanje',
    'godisnji_odmor': 'Godišnji odmor',
}
AUTOMATIC_CATEGORIES = {'topli_obrok', 'regres', 'minuli_rad'}


def category_code(value):
    value = str(value or '').strip().casefold().replace('đ','dj')
    value = ''.join(c for c in unicodedata.normalize('NFD',value) if unicodedata.category(c)!='Mn')
    return re.sub(r'[^a-z0-9]+','_',value).strip('_')


@transaction.atomic
def import_work_time_catalog(path):
    book = load_workbook(path,read_only=True,data_only=True)
    rows=[]
    seen={}
    try:
        columns=None
        for number,row in enumerate(book.active.iter_rows(values_only=True),1):
            labels={str(value or '').strip():i for i,value in enumerate(row)}
            if columns is None:
                if all(name in labels for name in ('sif_prim','naz_prim','napomena','elsif','elnaz')):
                    columns=labels
                continue
            if not any(v is not None for v in row[:5]): continue
            try:
                raw_code=row[columns['sif_prim']]
                code=str(int(raw_code)).zfill(2)
                if float(raw_code)!=int(raw_code) or int(raw_code)<1:
                    raise ValueError
                name=' '.join(str(row[columns['naz_prim']] or '').split())
                category=category_code(row[columns['napomena']])
                raw_element=row[columns['elsif']]
                element=int(raw_element)
                element_name=' '.join(str(row[columns['elnaz']] or '').split())
                if not code.isdigit() or not name or category not in CATEGORY_LABELS or not element_name or element<1 or float(raw_element)!=element:
                    raise ValueError
                key=(code,element)
                values=(name,category,element_name)
                if key in seen and seen[key]!=values:
                    raise ValueError
                seen[key]=values
                rows.append((code,name,category,element,element_name))
            except (ValueError,TypeError,IndexError):
                raise ValidationError(f'Red {number}: proverite šifru primaoca, napomenu i element.')
        if not rows: raise ValidationError('Nema elemenata radne liste za uvoz.')
    finally:
        book.close()
    counts={'recipients_created':0,'categories_created':0,'elements_created':0,'existing_preserved':0}
    for code,name,category,element,element_name in rows:
        recipient,created=RecipientType.objects.get_or_create(code=code,defaults={'name':name})
        counts['recipients_created']+=created
        kind,created=WorkTimeCategory.objects.get_or_create(code=category,defaults={
            'name':CATEGORY_LABELS[category], 'employee_selectable':category not in AUTOMATIC_CATEGORIES,
            'sort_order':list(CATEGORY_LABELS).index(category)})
        counts['categories_created']+=created
        _,created=WorkTimeElement.objects.get_or_create(recipient_type=recipient,payroll_code=element,
            defaults={'category':kind,'payroll_name':element_name})
        counts['elements_created']+=created
        counts['existing_preserved']+=not created
    return counts


@transaction.atomic
def sync_employee_recipient_types(using=None):
    """Refresh only the two new source fields; preserve locally edited dictionary labels."""
    with connections[_resolve_hr_db_alias(using)].cursor() as cursor:
        cursor.execute('SELECT rasif, sif_prim, naz_prim FROM dbo.hr_employee')
        source_rows=cursor.fetchall()
    source={}
    for employee_code,code,name in source_rows:
        if employee_code is None: continue
        values=(str(code or '').strip(), ' '.join(str(name or '').split()))
        key=int(employee_code)
        if key in source and source[key]!=values:
            raise ValidationError('HR izvor ima različite vrste primaoca za istog zaposlenog. Sinhronizacija je zaustavljena.')
        source[key]=values
    counts={'updated':0,'unchanged':0,'missing_in_source':0,'recipient_types_created':0}
    for employee in Employee.objects.all():
        values=source.get(employee.employee_code)
        if values is None:
            counts['missing_in_source']+=1
            continue
        code,name=values
        if code:
            _,created=RecipientType.objects.get_or_create(code=code,defaults={'name':name or code})
            counts['recipient_types_created']+=created
        if (employee.recipient_code,employee.recipient_name)!=(code,name):
            employee.recipient_code,employee.recipient_name=code,name
            employee.save(update_fields=['recipient_code','recipient_name'])
            counts['updated']+=1
        else:counts['unchanged']+=1
    return counts
