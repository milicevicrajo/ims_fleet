"""RFZO workbook import and calendar annotations, without payroll calculations."""
import hashlib
import re
from collections import defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Q
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.utils.exceptions import InvalidFileException

from hr.models import Employee, SickLeave, SickLeaveImport

HEADERS = {
    'rfzo_id': ('Ид', 'Id', 'RFZO ID'),
    'personal_number': ('ЈМБГ', 'JMBG'),
    'source_status': ('Статус', 'Status'),
    'start_date': ('Датум почетка боловања', 'Datum početka bolovanja', 'Datum pocetka bolovanja'),
    'end_date': ('Датум закључења боловања', 'Datum zaključenja bolovanja', 'Datum zakljucenja bolovanja'),
    'source_total_days': ('Укупно дана', 'Ukupno dana'),
}


def normalize_personal_number(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if int(value) != value:
            return None
        value = str(int(value))
    value = str(value).strip().removeprefix("'")
    # A numeric Excel cell may have lost a leading zero.
    if len(value) == 12 and value.isascii() and value.isdigit():
        value = '0' + value
    return value if re.fullmatch(r'[0-9]{13}', value) else None


def _date(value, epoch):
    if value in (None, '', '-', '—'):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        parsed = from_excel(value, epoch)
        if isinstance(parsed, datetime):
            return parsed.date()
        raise ValueError
    text = str(value).strip()
    for fmt in ('%Y-%m-%d', '%d.%m.%Y', '%d.%m.%Y.', '%d/%m/%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError


def _status(value):
    original = str(value or '').strip()
    key = original.casefold()
    choices = {'aktivno':'Aktivno', 'активно':'Aktivno', 'zaključeno':'Zaključeno',
               'zakljuceno':'Zaključeno', 'закључено':'Zaključeno',
               'stornirano':'Stornirano', 'сторнирано':'Stornirano',
               'poništeno':'Poništeno', 'ponisteno':'Poništeno', 'поништено':'Poništeno'}
    if key not in choices:
        raise ValueError
    return choices[key]


def parse_rfzo_workbook(content):
    if len(content) > 10 * 1024 * 1024:
        raise ValidationError('Excel datoteka može imati najviše 10 MB.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                raise ValidationError('Excel datoteka je prevelika za ovaj uvoz.')
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, OSError, ValueError, KeyError, ParseError, InvalidFileException) as exc:
        raise ValidationError('Datoteka nije ispravan RFZO Excel (.xlsx).') from exc
    try:
        sheet = workbook.active
        mapping = None
        results = {}
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
            if row_number > 50000:
                raise ValidationError('Datoteka ima previše redova.')
            if mapping is None:
                labels = {str(v or '').strip().casefold(): i for i,v in enumerate(row)}
                candidate = {name: next((labels[a.casefold()] for a in aliases if a.casefold() in labels), None)
                             for name,aliases in HEADERS.items()}
                if all(v is not None for v in candidate.values()):
                    mapping = candidate
                elif row_number >= 20:
                    break
                continue
            if not any(v is not None and str(v).strip() for v in row):
                continue
            try:
                values = {name: row[index] if index < len(row) else None for name,index in mapping.items()}
                identity = values['rfzo_id']
                if isinstance(identity, (float,int)) and not isinstance(identity,bool):
                    if int(identity) != identity: raise ValueError
                    identity = str(int(identity))
                identity = str(identity or '').strip()
                personal_number = normalize_personal_number(values['personal_number'])
                if not identity or len(identity)>64 or personal_number is None: raise ValueError
                start, end = _date(values['start_date'],workbook.epoch),_date(values['end_date'],workbook.epoch)
                if start is None or not 1900 <= start.year <= 2100 or (end and (end<start or end.year>2100)): raise ValueError
                source_status = _status(values['source_status'])
                if source_status == 'Zaključeno' and end is None: raise ValueError
                days = values['source_total_days']
                if days not in (None,''):
                    if not re.fullmatch(r'[0-9]+(?:\.0+)?', str(days).strip()): raise ValueError
                    days = int(float(days))
                    if days > 100000: raise ValueError
                else: days=None
                parsed = dict(rfzo_id=identity,personal_number=personal_number,start_date=start,
                              end_date=end,source_status=source_status,source_total_days=days)
            except (ValueError,TypeError,OverflowError) as exc:
                raise ValidationError(f'Red {row_number}: proverite RFZO ID, JMBG, datume, status i broj dana. Nijedan red nije uvezen.') from exc
            if identity in results and results[identity] != parsed:
                raise ValidationError(f'Red {row_number}: isti RFZO ID ima različite podatke u datoteci. Nijedan red nije uvezen.')
            results[identity] = parsed
        if mapping is None:
            raise ValidationError('Nisu pronađene RFZO kolone: ID, JMBG, status, početak, zaključenje i ukupno dana.')
        if not results:
            raise ValidationError('Datoteka nema bolovanja za uvoz.')
        return list(results.values())
    except (ParseError, BadZipFile, KeyError) as exc:
        raise ValidationError('Sadržaj Excel datoteke nije ispravan. Nijedan red nije uvezen.') from exc
    finally:
        workbook.close()


@transaction.atomic
def import_rfzo_workbook(content, *, filename, source_date, user=None, dry_run=False):
    rows = parse_rfzo_workbook(content)
    identities = defaultdict(list)
    for employee_id, personal_number in Employee.objects.values_list('pk','personal_number'):
        normalized = normalize_personal_number(personal_number)
        if normalized: identities[normalized].append(employee_id)
    # Lock existing identities before updates; all validation precedes writes.
    existing = {}
    identifiers = [r['rfzo_id'] for r in rows]
    for offset in range(0,len(identifiers),1000):
        existing.update({s.rfzo_id:s for s in SickLeave.objects.select_for_update().filter(rfzo_id__in=identifiers[offset:offset+1000])})
    counts = dict(row_count=len(rows),created_count=0,updated_count=0,unchanged_count=0,unlinked_count=0)
    changes=[]
    for row in rows:
        current = existing.get(row['rfzo_id'])
        if current and current.personal_number != row['personal_number']:
            raise ValidationError('RFZO ID već postoji uz drugi JMBG. Uvoz nije izvršen; proverite izvor.')
        if current and source_date < current.source_date:
            raise ValidationError('Datoteka je starija od već evidentiranih podataka. Proverite datum RFZO izvoza.')
        employees = identities.get(row['personal_number'], [])
        row['employee_id'] = employees[0] if len(employees)==1 else None
        row['matching_note'] = '' if len(employees)==1 else ('Više zaposlenih sa istim JMBG-om' if employees else 'Zaposleni nije pronađen po JMBG-u')
        if row['employee_id'] is None: counts['unlinked_count']+=1
        row['source_date']=source_date
        changed = current is None or any(getattr(current,k)!=v for k,v in row.items())
        counts['created_count' if current is None else ('updated_count' if changed else 'unchanged_count')]+=1
        changes.append((current,row,changed))
    if dry_run:
        return counts
    batch = SickLeaveImport.objects.create(filename=Path(filename).name[:255],
        file_hash=hashlib.sha256(content).hexdigest(),source_date=source_date,created_by=user,**counts)
    for current,row,changed in changes:
        if current is None:
            SickLeave.objects.create(**row,last_import=batch)
        elif changed:
            for key,value in row.items(): setattr(current,key,value)
            current.last_import=batch
            current.save()
    return batch


def sick_leaves_by_day(employee, year, month):
    from .attendance import month_period
    start,exclusive_end,_ = month_period(year,month)
    last=exclusive_end-timedelta(days=1)
    result=defaultdict(list)
    leaves=SickLeave.objects.filter(employee=employee,start_date__lt=exclusive_end).filter(
        Q(end_date__gte=start)|Q(end_date__isnull=True,source_date__gte=start)
    ).exclude(source_status__in=['Stornirano','Poništeno'])
    for leave in leaves:
        current=max(start,leave.start_date)
        end=min(last,leave.end_date or leave.source_date)
        while current<=end:
            result[current].append(leave)
            current+=timedelta(days=1)
    return result
