import copy
import uuid
from datetime import date
from decimal import Decimal, DecimalException

from django.core import signing
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.db.models import Max, Q
from django.utils import timezone
from openpyxl import load_workbook

from core.mixins import user_has_role_permission
from core.models import OrganizationalUnit
from hr.models import (Employee, EvaluationGroup, EvaluationCriterion, EvaluationScale,
    EvaluationUnitSetup, EmployeeEvaluation, EvaluationApproval)

STAGES = ['supervisor', 'director', 'general_director']
SIGNING_SALT = 'hr.evaluation.snapshot.v1'
MONTHS = ['Januar','Februar','Mart','April','Maj','Jun','Jul','Avgust','Septembar','Oktobar','Novembar','Decembar']


def latin(value):
    text = str(value or '')
    pairs = [
        ('\u0409','Lj'),('\u040a','Nj'),('\u040f','Dž'),('\u0459','lj'),('\u045a','nj'),('\u045f','dž'),
        ('\u0410','A'),('\u0411','B'),('\u0412','V'),('\u0413','G'),('\u0414','D'),('\u0402','Đ'),('\u0415','E'),('\u0416','Ž'),('\u0417','Z'),
        ('\u0418','I'),('\u0408','J'),('\u041a','K'),('\u041b','L'),('\u041c','M'),('\u041d','N'),('\u041e','O'),('\u041f','P'),('\u0420','R'),
        ('\u0421','S'),('\u0422','T'),('\u040b','Ć'),('\u0423','U'),('\u0424','F'),('\u0425','H'),('\u0426','C'),('\u0427','Č'),('\u0428','Š'),
        ('\u0430','a'),('\u0431','b'),('\u0432','v'),('\u0433','g'),('\u0434','d'),('\u0452','đ'),('\u0435','e'),('\u0436','ž'),('\u0437','z'),
        ('\u0438','i'),('\u0458','j'),('\u043a','k'),('\u043b','l'),('\u043c','m'),('\u043d','n'),('\u043e','o'),('\u043f','p'),('\u0440','r'),
        ('\u0441','s'),('\u0442','t'),('\u045b','ć'),('\u0443','u'),('\u0444','f'),('\u0445','h'),('\u0446','c'),('\u0447','č'),('\u0448','š'),
    ]
    mapping = dict(pairs)
    return ''.join(mapping.get(char, char) for char in text)


def cyrillic(value):
    return latin(value)


def latin_data(value):
    if isinstance(value, dict):
        return {key: latin_data(item) for key, item in value.items()}
    if isinstance(value, list):
        return [latin_data(item) for item in value]
    if isinstance(value, str):
        return latin(value)
    return value


def money_decimal(value):
    try:
        result = Decimal(str(value))
    except DecimalException:
        raise ValidationError('Koeficijent mora biti broj.')
    if not result.is_finite():
        raise ValidationError('Koeficijent mora biti konačan broj.')
    try:
        return result.quantize(Decimal('0.00001'))
    except DecimalException:
        raise ValidationError('Koeficijent ima neispravnu vrednost.')


@transaction.atomic
def import_evaluation_catalog(path):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        maxima = list(book['Max_vrednosti'].iter_rows(min_row=2, max_col=4, values_only=True))
        criteria = book['kriterijumi']
        form = book['obrazac']
        rows = []
        seen = set()
        for management, criterion, points, coefficient in maxima:
            if all(value is None for value in (management, criterion, points, coefficient)):
                continue
            if management not in (0,1) or criterion not in range(1,7) or points not in range(5):
                raise ValidationError('Neispravne šifre u listu Max_vrednosti.')
            key = (management, criterion, points)
            if key in seen:
                raise ValidationError('Duplirana kombinacija grupe, merila i bodova.')
            seen.add(key)
            coefficient = money_decimal(coefficient)
            column = [1,7,13][(criterion-1) % 3]
            row = (4 if criterion <= 3 else 15) + (4-points)
            description = latin(str(criteria.cell(row, column).value or '').strip())
            title = latin(str(form.cell(18+criterion, 1).value or '').strip().split('.',1)[-1].strip())
            rows.append((management, criterion, points, coefficient, title, description))
        if len(seen) != 60:
            raise ValidationError('Očekivano je 60 kombinacija za dve grupe, šest merila i pet nivoa.')
    finally:
        book.close()
    counts = {'groups':0,'criteria':0,'scales':0,'preserved':0}
    for management, code, points, coefficient, title, description in rows:
        group, created = EvaluationGroup.objects.get_or_create(code='management' if management else 'employee',
            defaults={'name':'Menadžment' if management else 'Ostali zaposleni'})
        counts['groups'] += created
        criterion, created = EvaluationCriterion.objects.get_or_create(code=code, defaults={'name':title,'sort_order':code})
        counts['criteria'] += created
        _, created = EvaluationScale.objects.get_or_create(group=group,criterion=criterion,points=points,
            defaults={'coefficient':coefficient,'description':description})
        counts['scales'] += created
        counts['preserved'] += not created
    return counts


def can_view_all(user):
    return user_has_role_permission(user, 'hr:evaluation_view_all')


def visible_evaluations(user):
    qs = EmployeeEvaluation.objects.select_related('employee','created_by').prefetch_related('approvals')
    if can_view_all(user):
        return qs
    if not user.is_authenticated:
        return qs.none()
    if user_has_role_permission(user, 'hr:kadrovi_manage'):
        from hr.access import visible_employees
        return qs.filter(employee__in=visible_employees(user))
    match = Q(created_by=user)
    if user.employee_id:
        match |= Q(supervisor_id=user.employee_id) | Q(director_id=user.employee_id) | Q(general_director_id=user.employee_id)
    return qs.filter(match)


def editable_employees(user):
    if not user_has_role_permission(user, 'hr:evaluation_create'):
        return Employee.objects.none()
    if can_view_all(user):
        return Employee.objects.all()
    if user_has_role_permission(user, 'hr:kadrovi_manage'):
        from hr.access import visible_employees
        return visible_employees(user)
    units = EvaluationUnitSetup.objects.filter(is_active=True, supervisor_id=user.employee_id).values_list('unit_code',flat=True) if user.employee_id else []
    return Employee.objects.filter(org_unit_code__in=units)


def display_person(employee):
    return cyrillic(f'{employee.display_first_name} {employee.display_last_name}'.strip())


def build_snapshot(*, employee, group, year, month, supervisor, director, general_director):
    if not group.is_active:
        raise ValidationError('Izabrana grupa nije aktivna.')
    criteria = list(EvaluationCriterion.objects.filter(is_active=True))
    scales = list(EvaluationScale.objects.filter(group=group,is_active=True,criterion__is_active=True))
    items = []
    for criterion in criteria:
        options = [{'points':row.points,'value':str(row.coefficient),'description':latin(row.description)} for row in scales if row.criterion_id == criterion.pk]
        if not options or not any(Decimal(option['value']) == 0 for option in options):
            raise ValidationError(f'Merilo {criterion.code} nema aktivnu skalu sa neutralnom vrednošću 0.')
        neutral = next((row for row in options if row['points']==1 and Decimal(row['value'])==0), next(row for row in options if Decimal(row['value'])==0))
        items.append({'code':criterion.code,'name':latin(criterion.name),'options':options,'points':neutral['points'],'value':'0.00000','comment':''})
    if not items:
        raise ValidationError('Nema aktivnih merila za ocenjivanje.')
    unit_code = str(employee.org_unit_code or employee.department_code or '')
    unit = OrganizationalUnit.objects.filter(code=unit_code).first()
    return {
        'schema':1,'basis':'employee_unit','rule':'K4 = 1 + sum(stimulation)',
        'year':year,'month':month,'month_name':MONTHS[month-1],
        'employee':{'id':employee.pk,'code':employee.employee_code,'name':display_person(employee),
            'original_name':str(employee),'position':cyrillic(employee.position or employee.job_title),
            'title':cyrillic(employee.title),'education':cyrillic(employee.education)},
        'unit':{'code':unit_code,'name':cyrillic(unit.name) if unit else '', 'center':latin(unit.center) if unit else ''},
        'group':{'code':group.code,'name':cyrillic(group.name)},
        'reviewers':{key:{'id':person.pk,'name':display_person(person)} for key,person in zip(STAGES,[supervisor,director,general_director])},
        'items':items,'comment':'',
        'minimum':str(Decimal(1)+sum(min(Decimal(option['value']) for option in item['options']) for item in items)),
        'maximum':str(Decimal(1)+sum(max(Decimal(option['value']) for option in item['options']) for item in items)),
    }


def sign_snapshot(snapshot, user, source_id=None):
    return signing.dumps({'snapshot':snapshot,'user':user.pk,'nonce':str(uuid.uuid4()),'source_id':source_id},salt=SIGNING_SALT,compress=True)


def read_snapshot(token, user):
    try:
        data = signing.loads(token,salt=SIGNING_SALT,max_age=86400)
        if data['user'] != user.pk:
            raise signing.BadSignature
        return data
    except (signing.BadSignature, KeyError):
        raise ValidationError('Obrazac je istekao ili nije važeći. Ponovo otvorite obrazac.')


@transaction.atomic
def save_evaluation(*, envelope, values, user):
    snapshot = copy.deepcopy(envelope['snapshot'])
    employee_id = snapshot['employee']['id']
    if not editable_employees(user).filter(pk=employee_id).exists():
        raise PermissionDenied
    Employee.objects.select_for_update().get(pk=employee_id)
    existing = EmployeeEvaluation.objects.filter(request_key=envelope['nonce']).first()
    if existing:
        return existing
    source_id = envelope.get('source_id')
    if source_id and not visible_evaluations(user).filter(pk=source_id,employee_id=employee_id,year=snapshot['year'],month=snapshot['month']).exists():
        raise PermissionDenied
    total = Decimal(0)
    for item in snapshot['items']:
        choice = values[f'criterion_{item["code"]}']
        option = next((option for option in item['options'] if option['points'] == int(choice)), None)
        if option is None:
            raise ValidationError('Izabrani bodovi ne pripadaju sačuvanoj skali.')
        item.update(points=option['points'],value=option['value'],selected_description=option['description'],comment=cyrillic(values.get(f'comment_{item["code"]}','')))
        total += Decimal(option['value'])
    snapshot['comment'] = cyrillic(values.get('comment',''))
    snapshot['stimulation'] = str(total)
    snapshot['coefficient'] = str(Decimal(1)+total)
    qs = EmployeeEvaluation.objects.filter(employee_id=employee_id,year=snapshot['year'],month=snapshot['month'])
    revision = (qs.aggregate(last=Max('revision'))['last'] or 0)+1
    return EmployeeEvaluation.objects.create(employee_id=employee_id,year=snapshot['year'],month=snapshot['month'],
        revision=revision,unit_code=snapshot['unit']['code'],snapshot=snapshot,stimulation=total,
        personal_coefficient=Decimal(1)+total,request_key=envelope['nonce'],created_by=user,source_evaluation_id=source_id,
        **{key+'_id':snapshot['reviewers'][key]['id'] for key in STAGES})


def approval_state(evaluation):
    approvals = {approval.stage:approval for approval in evaluation.approvals.all()}
    latest = next((approvals[stage] for stage in reversed(STAGES) if stage in approvals),None)
    return approvals, latest.coefficient if latest else evaluation.personal_coefficient


@transaction.atomic
def approve_evaluation(*, evaluation, stage, user, document_date, coefficient=None, comment=''):
    if stage not in STAGES or not user_has_role_permission(user,'hr:evaluation_approve'):
        raise PermissionDenied
    if getattr(evaluation,stage+'_id') != user.employee_id:
        raise PermissionDenied('Saglasnost daje samo imenovani ocenjivač sa svog naloga.')
    Employee.objects.select_for_update().get(pk=evaluation.employee_id)
    evaluation = EmployeeEvaluation.objects.select_for_update().get(pk=evaluation.pk)
    if EmployeeEvaluation.objects.filter(employee=evaluation.employee,year=evaluation.year,month=evaluation.month,revision__gt=evaluation.revision).exists():
        raise ValidationError('Postoji novija verzija obrasca. Potvrdite novu verziju.')
    approvals, previous = approval_state(evaluation)
    if stage in approvals:
        raise ValidationError('Ova saglasnost je već sačuvana.')
    if any(prior not in approvals for prior in STAGES[:STAGES.index(stage)]):
        raise ValidationError('Prethodni nivo ocenjivanja još nije potvrdio obrazac.')
    value = previous if coefficient is None else money_decimal(coefficient)
    if stage == 'supervisor' and value != previous:
        raise ValidationError('Za korekciju merila sačuvajte novu verziju obrasca.')
    if value < Decimal(evaluation.snapshot['minimum']) or value > Decimal(evaluation.snapshot['maximum']):
        raise ValidationError('Koeficijent je izvan granica sačuvanog šifrarnika.')
    if value != previous and not comment.strip():
        raise ValidationError('Za korekciju koeficijenta unesite obrazloženje.')
    return EvaluationApproval.objects.create(evaluation=evaluation,stage=stage,actor=user,
        actor_name=display_person(user.employee),document_date=document_date,coefficient=value,comment=cyrillic(comment))
