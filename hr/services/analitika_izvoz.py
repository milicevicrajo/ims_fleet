"""Excel analitike zaposlenih: list po celini, sa zaglavljem Instituta i grafikonima (openpyxl)."""
from decimal import Decimal

from django.utils import timezone
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.exporting import create_xlsx_workbook, workbook_response
from .spisak_zaposlenih import ADRESA, INSTITUT, PLAVA, SVETLA, LINIJA

TANKA = Side(style="thin", color=LINIJA)
OKVIR = Border(left=TANKA, right=TANKA, top=TANKA, bottom=TANKA)


def _broj(v):
    return float(v) if isinstance(v, Decimal) else v


def _list(wb, naslov, opis, filteri, zaglavlje, redovi, *, zbir=None, prvi=False, sirine=None):
    ws = wb.active if prvi else wb.create_sheet()
    ws.title = naslov[:31]
    ws.sheet_view.showGridLines = False
    ws["A1"] = INSTITUT
    ws["A1"].font = Font(bold=True, size=12, color=PLAVA)
    ws["A2"] = f"Analitika zaposlenih · {naslov}"
    ws["A2"].font = Font(bold=True, size=14, color=PLAVA)
    ws["A3"] = f"{filteri} · {opis}" if opis else filteri
    ws["A3"].font = Font(size=9, italic=True, color="475569")
    red = 5
    for k, tekst in enumerate(zaglavlje, 1):
        c = ws.cell(red, k, tekst)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=PLAVA), OKVIR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[red].height = 30
    for i, vrednosti in enumerate(redovi):
        for k, v in enumerate(vrednosti, 1):
            c = ws.cell(red + 1 + i, k, _broj(v))
            c.border = OKVIR
            if isinstance(v, str):
                c.data_type = "s"
            elif k > 1:
                c.alignment = Alignment(horizontal="right")
                c.number_format = "0" if isinstance(v, int) else "0.0"
            if i % 2:
                c.fill = PatternFill("solid", fgColor=SVETLA)
    poslednji = red + len(redovi)
    if zbir:
        poslednji += 1
        for k, v in enumerate(zbir, 1):
            c = ws.cell(poslednji, k, _broj(v))
            c.font, c.border = Font(bold=True), OKVIR
            c.fill = PatternFill("solid", fgColor="E8F0FB")
            if not isinstance(v, str) and v is not None:
                c.number_format = "0" if isinstance(v, int) else "0.0"
    for k, sirina in enumerate(sirine or [], 1):
        ws.column_dimensions[get_column_letter(k)].width = sirina
    ws.freeze_panes = ws.cell(red + 1, 1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws, red, red + len(redovi)


def _grafikon(ws, naslov, prvi_red, poslednji_red, kolone, mesto, *, slozen=False, horizontalan=False):
    if poslednji_red <= prvi_red:
        return
    g = BarChart()
    g.type = "bar" if horizontalan else "col"
    g.title, g.height, g.width = naslov, 7.5, 16
    if slozen:
        g.grouping, g.overlap = "stacked", 100
    for k in kolone:
        g.add_data(Reference(ws, min_col=k, min_row=prvi_red, max_row=poslednji_red), titles_from_data=True)
    g.set_categories(Reference(ws, min_col=1, min_row=prvi_red + 1, max_row=poslednji_red))
    ws.add_chart(g, mesto)


def excel(a, filteri):
    wb, _ = create_xlsx_workbook()
    s = a["sazetak"]
    _list(wb, "Sažetak", f"stanje na dan {a['na_dan']:%d.%m.%Y.}", filteri, ["Pokazatelj", "Vrednost"], [
        ["Aktivnih zaposlenih", a["ukupno"]], ["Žene", s["zene"]], ["Muškarci", s["muski"]],
        ["Udeo žena %", s["udeo_zena"]], ["Prosečna starost (godine)", s["prosek"]],
        ["Medijana starosti", s["medijana"]], ["Prosečna starost žena", s["prosek_zene"]],
        ["Prosečna starost muškaraca", s["prosek_muski"]], ["Najmlađi (godine)", s["najmladji"]],
        ["Najstariji (godine)", s["najstariji"]], ["Prosečan staž u Institutu (godine)", s["prosek_staz"]],
        ["Visoka i viša sprema (VI–VIII)", s["visoka"]], ["Udeo visoke i više spreme %", s["udeo_visoka"]],
        ["60 i više godina", s["pred_penzijom"]], ["Bez ispravnog datuma rođenja", s["bez_starosti"]],
    ], prvi=True, sirine=[38, 14])

    zaglavlje = ["Grupa", "Ukupno", "Udeo %", "Žene", "Muškarci", "Udeo žena %", "Prosečna starost",
                 "Prosečna starost žena", "Prosečna starost muškaraca"]

    def redovi(lista, naziv=lambda r: r["naziv"]):
        return [[naziv(r), r["ukupno"], r["udeo"], r["zene"], r["muski"], r["udeo_zena"], r["prosek"],
                 r["prosek_zene"], r["prosek_muski"]] for r in lista]

    sirine = [40, 10, 10, 10, 11, 12, 12, 14, 16]
    ws, prvi, poslednji = _list(wb, "Starost", "rasponi od deset godina", filteri,
                                ["Godine", *zaglavlje[1:]], redovi(a["starosni"]), sirine=sirine)
    _grafikon(ws, "Starosna struktura po polu", prvi, poslednji, (4, 5), f"A{poslednji + 3}", slozen=True)
    ws, prvi, poslednji = _list(wb, "Stručna sprema", "stepen izveden iz zanimanja i škole", filteri,
                                ["Stručna sprema", *zaglavlje[1:]],
                                redovi(a["stepeni"], lambda r: r["prikaz"]), sirine=sirine)
    _grafikon(ws, "Stručna sprema po polu", prvi, poslednji, (4, 5), f"A{poslednji + 3}", slozen=True, horizontalan=True)
    _list(wb, "Starost po spremi", "broj zaposlenih", filteri,
          ["Stručna sprema", *a["rasponi_nazivi"], "Ukupno", "Prosečna starost"],
          [[r["prikaz"], *r["rasponi"], r["ukupno"], r["prosek"]] for r in a["ukrstanje"]],
          sirine=[40] + [10] * (len(a["rasponi_nazivi"]) + 2))
    ws, prvi, poslednji = _list(wb, "Staž u Institutu", "godine od datuma zaposlenja", filteri,
                                ["Godine staža", *zaglavlje[1:]], redovi(a["staz"]), sirine=sirine)
    _grafikon(ws, "Staž u Institutu", prvi, poslednji, (2,), f"A{poslednji + 3}")
    _list(wb, "Organizacione jedinice", "", filteri, ["Organizaciona jedinica", *zaglavlje[1:]], redovi(a["oj"]),
          sirine=[48] + sirine[1:])
    _list(wb, "Radni odnos", "", filteri, ["Vrsta radnog odnosa", *zaglavlje[1:]], redovi(a["odnos"]), sirine=sirine)
    _list(wb, "Razvrstavanje", "kako je svako zanimanje razvrstano u stepen spreme — za proveru", filteri,
          ["Stepen", "Prepoznato iz", "Zanimanje", "Škola", "Broj zaposlenih"],
          [[stepen, izvor or "—", zanimanje, skola, broj] for (stepen, izvor, zanimanje, skola), broj in a["razvrstavanje"]],
          sirine=[16, 14, 40, 40, 14])
    return workbook_response(wb, f"analitika_zaposlenih_{timezone.localdate():%Y%m%d}.xlsx", quoted=True)
