"""Statistika rodne ravnopravnosti: tabele po tačkama Obrasca 1, za ekran, Excel i PDF (od 07.10.2026.).

Ćelija je tekst ili par (broj, udeo %); ekran i PDF ga pišu kao u obrascu — „75 (24,1%)”, a Excel u dve
kolone (broj i %), da bi se moglo računati.
"""
from io import BytesIO
from xml.sax.saxutils import escape

from django.http import HttpResponse
from django.utils import timezone
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.exporting import create_xlsx_workbook, workbook_response
from .rodna_ravnopravnost import broj_sa_udelom
from .spisak_zaposlenih import INSTITUT, LINIJA, PLAVA, SVETLA

NASLOV = "Evidencija podataka o ostvarivanju rodne ravnopravnosti — Obrazac 1"


def _par(red, kljuc, udeo):
    return (red[kljuc], red[udeo])


def tabele(s):
    """[{broj, naslov, opis, zaglavlje, redovi, zbir}] po tačkama obrasca koje se računaju iz kadrovske baze."""
    u = s["ukupno"]
    na_dan = f"{s['na_dan']:%d.%m.%Y.}"
    rezultat = [{
        "broj": "1", "naslov": "Ukupan broj zaposlenih i radno angažovanih lica, po polu",
        "opis": f"Stanje na dan {na_dan}", "zaglavlje": ["", "Ukupno", "Žene", "Muškarci"],
        "redovi": [["Zaposleni i radno angažovani", u["ukupno"], (u["zene"], u["udeo_zena"]), (u["muski"], u["udeo_muski"])]],
    }, {
        "broj": "2", "naslov": "Broj i procenat zaposlenih i radno angažovanih lica, po polu i starosnoj dobi",
        "opis": "Ukupno — udeo u svim zaposlenima; žene i muškarci — udeo u starosnoj grupi.",
        "zaglavlje": ["Starosna dob", "Ukupno", "Žene", "Muškarci"],
        "redovi": [[r["naziv"], _par(r, "ukupno", "udeo"), _par(r, "zene", "udeo_zena"), _par(r, "muski", "udeo_muski")]
                   for r in s["starost"]],
    }, {
        "broj": "3", "naslov": "Broj i procenat zaposlenih i radno angažovanih lica, po polu i nivou kvalifikacija",
        "opis": "Udeo u svim zaposlenima. Stručna sprema je današnji podatak iz kadrovske baze.",
        "zaglavlje": ["Nivo kvalifikacija", "Ukupno", "Žene", "Muškarci"],
        "redovi": [[f"{r['rb']}. {r['naziv']}" if r["rb"] != "—" else r["naziv"], _par(r, "ukupno", "udeo"),
                    _par(r, "zene", "udeo_zena"), _par(r, "muski", "udeo_muski")] for r in s["sprema"]],
        "zbir": ["Ukupno", u["ukupno"], u["zene"], u["muski"]],
    }, {
        "broj": "4", "naslov": "Broj i procenat lica na izvršilačkim radnim mestima i na položajima, po polu",
        "opis": "Položaj je rukovodeće mesto iz kadrovske baze (šifarnik rukovodećih mesta). "
                "Žene i muškarci — udeo u ukupnom broju žena, odnosno muškaraca.",
        "zaglavlje": ["", "Ukupno", "Žene", "Muškarci"],
        "redovi": [[r["naziv"], _par(r, "ukupno", "udeo"), _par(r, "zene", "udeo_zena"), _par(r, "muski", "udeo_muski")]
                   for r in s["polozaji"]],
    }]
    for broj, kljuc, naslov in (("6a", "prijem", "Lica koja su u godini zaposlena, odnosno radno angažovana, po polu i starosnoj dobi"),
                                ("6b", "prestanak", "Lica kojima je u godini prestao radni odnos, odnosno radno angažovanje, po polu i starosnoj dobi")):
        p = s[kljuc]
        rezultat.append({
            "broj": broj, "naslov": f"{naslov} ({s['godina']}.)",
            "opis": f"Ukupno {p['ukupno']}: žene {broj_sa_udelom(p['zene'], p['udeo_zena'])}, "
                    f"muškarci {broj_sa_udelom(p['muski'], p['udeo_muski'])}. Žene i muškarci — udeo u starosnoj grupi.",
            "zaglavlje": ["Starosna dob", "Ukupno", "Žene", "Muškarci"],
            "redovi": [[r["naziv"], _par(r, "ukupno", "udeo"), _par(r, "zene", "udeo_zena"), _par(r, "muski", "udeo_muski")]
                       for r in p["starost"]],
            "zbir": ["Ukupno", p["ukupno"], p["zene"], p["muski"]],
        })
        if kljuc == "prestanak" and p["razlozi"]:
            rezultat.append({
                "broj": "6b", "naslov": "Razlozi prestanka radnog odnosa, odnosno radnog angažovanja",
                "opis": "Razlog odlaska iz kadrovske baze.", "zaglavlje": ["Razlog", "Ukupno", "Žene", "Muškarci"],
                "redovi": [[r["naziv"], (r["ukupno"], r["udeo"]), r["zene"], r["muski"]] for r in p["razlozi"]],
            })
    return rezultat


def tekst(celija):
    if isinstance(celija, tuple):
        return broj_sa_udelom(*celija)
    return "" if celija is None else str(celija)


def podnaslov(s):
    return (f"Godina {s['godina']} · stanje na dan {s['na_dan']:%d.%m.%Y.}"
            + (" (tekuća godina — današnje stanje)" if s["tekuca"] else "") + " · izvor: kadrovska baza")


# --------------------------------------------------------------------------- Excel

def excel(s):
    tanka = Side(style="thin", color=LINIJA)
    okvir = Border(left=tanka, right=tanka, top=tanka, bottom=tanka)
    wb, ws = create_xlsx_workbook("Obrazac 1")
    ws.sheet_view.showGridLines = False
    ws["A1"], ws["A2"], ws["A3"] = INSTITUT, NASLOV, podnaslov(s)
    ws["A1"].font = Font(bold=True, size=12, color=PLAVA)
    ws["A2"].font = Font(bold=True, size=14, color=PLAVA)
    ws["A3"].font = Font(size=9, italic=True, color="475569")
    red = 5
    for t in tabele(s):
        ws.cell(red, 1, f"{t['broj']}. {t['naslov']}").font = Font(bold=True, size=11, color=PLAVA)
        red += 1
        if t.get("opis"):
            ws.cell(red, 1, t["opis"]).font = Font(size=9, italic=True, color="475569")
            red += 1
        # Kolona sa parom (broj, udeo) u Excelu ima i kolonu „%”.
        kolone = [any(isinstance(r[i], tuple) for r in t["redovi"]) for i in range(len(t["zaglavlje"]))]
        k = 1
        for naziv, par in zip(t["zaglavlje"], kolone):
            for naslov_kolone in ([naziv, "%"] if par else [naziv or ""]):
                c = ws.cell(red, k, naslov_kolone)
                c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=PLAVA), okvir
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                k += 1
        for j, vrednosti in enumerate(t["redovi"] + ([t["zbir"]] if t.get("zbir") else [])):
            red += 1
            k = 1
            zbirni = t.get("zbir") is vrednosti
            for v, par in zip(vrednosti, kolone):
                delovi = (list(v) if isinstance(v, tuple) else [v, None]) if par else [v]
                for d, deo in enumerate(delovi):
                    c = ws.cell(red, k, float(deo) if d == 1 and deo is not None else deo)
                    c.border = okvir
                    if d == 1 and deo is not None:
                        c.number_format = "0.0"
                    if zbirni:
                        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="E8F0FB")
                    elif j % 2:
                        c.fill = PatternFill("solid", fgColor=SVETLA)
                    k += 1
        red += 2
    ws.cell(red, 1, "Popunjava se ručno (podaci nisu u kadrovskoj bazi)").font = Font(bold=True, size=11, color=PLAVA)
    for broj, opis in s["rucno"]:
        red += 1
        ws.cell(red, 1, f"{broj}. {opis}")
    ws.column_dimensions["A"].width = 58
    for k in range(2, 9):
        ws.column_dimensions[get_column_letter(k)].width = 11
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    return workbook_response(wb, f"rodna_ravnopravnost_{s['godina']}.xlsx")


# --------------------------------------------------------------------------- PDF

def pdf(s):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from .analitika_pdf import _Brojac, _fontovi

    obican, podebljan = _fontovi()
    plava, siva, linija = colors.HexColor("#14395B"), colors.HexColor("#64748B"), colors.HexColor("#D7DCE5")
    izradjeno = timezone.localtime()
    sirina = A4[0] - 28 * mm
    stil = ParagraphStyle("c", fontName=obican, fontSize=8.5, leading=10.5)
    stil_z = ParagraphStyle("z", parent=stil, fontName=podebljan, textColor=colors.white)
    stil_d = ParagraphStyle("d", parent=stil, alignment=2)
    stil_u = ParagraphStyle("u", parent=stil_d, fontName=podebljan)
    naslov_t = ParagraphStyle("t", fontName=podebljan, fontSize=10, leading=13, textColor=plava, spaceBefore=6)
    opis_s = ParagraphStyle("o", fontName=obican, fontSize=7.5, leading=9.5, textColor=siva)

    sadrzaj = [Paragraph(escape(NASLOV), ParagraphStyle("n", fontName=podebljan, fontSize=13, leading=16, textColor=plava)),
               Paragraph(escape(podnaslov(s)), opis_s), Spacer(1, 4 * mm)]
    for t in tabele(s):
        podaci = [[Paragraph(escape(h), stil_z) for h in t["zaglavlje"]]]
        podaci += [[Paragraph(escape(tekst(v)), stil if i == 0 else stil_d) for i, v in enumerate(r)] for r in t["redovi"]]
        if t.get("zbir"):
            podaci.append([Paragraph(escape(tekst(v)), stil_u if i else ParagraphStyle("uz", parent=stil, fontName=podebljan))
                           for i, v in enumerate(t["zbir"])])
        tabela = Table(podaci, colWidths=[sirina * 0.43] + [sirina * 0.19] * 3, repeatRows=1)
        stilovi = [("BACKGROUND", (0, 0), (-1, 0), plava), ("GRID", (0, 0), (-1, -1), 0.3, linija),
                   ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                   ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]
        stilovi += [("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F4F7FA")) for i in range(2, len(podaci), 2)]
        if t.get("zbir"):
            stilovi.append(("BACKGROUND", (0, len(podaci) - 1), (-1, len(podaci) - 1), colors.HexColor("#E8F0FB")))
        tabela.setStyle(TableStyle(stilovi))
        blok = [Paragraph(escape(f"{t['broj']}. {t['naslov']}"), naslov_t)]
        if t.get("opis"):
            blok.append(Paragraph(escape(t["opis"]), opis_s))
        blok += [Spacer(1, 1.5 * mm), tabela, Spacer(1, 3 * mm)]
        sadrzaj.append(KeepTogether(blok))
    sadrzaj.append(Paragraph("Popunjava se ručno (podaci nisu u kadrovskoj bazi)", naslov_t))
    for broj, opis in s["rucno"]:
        sadrzaj.append(Paragraph(escape(f"{broj}. {opis}"), stil))

    def podnozje(c, strana, ukupno):
        c.saveState()
        c.setFont(obican, 7)
        c.setFillColor(siva)
        c.drawString(14 * mm, 8 * mm, f"{INSTITUT} · Kadrovi · Rodna ravnopravnost {s['godina']} · izrađeno {izradjeno:%d.%m.%Y. %H:%M}")
        c.drawRightString(A4[0] - 14 * mm, 8 * mm, f"Strana {strana} od {ukupno}")
        c.restoreState()

    izlaz = BytesIO()
    SimpleDocTemplate(izlaz, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm,
                      bottomMargin=16 * mm, title=NASLOV, author=INSTITUT).build(
        sadrzaj, canvasmaker=lambda *a, **k: _Brojac(*a, crtaj=podnozje, **k))
    odgovor = HttpResponse(izlaz.getvalue(), content_type="application/pdf")
    odgovor["Content-Disposition"] = f'attachment; filename="rodna_ravnopravnost_{s["godina"]}.pdf"'
    return odgovor
