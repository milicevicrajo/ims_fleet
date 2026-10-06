"""Kadrovi: komentar na prolaze i analitika zaposlenih (od 05.10.2026.)."""
import datetime
import re
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from core.models import PermissionCode, Role
from .models import Employee, KomentarProlaza, WorkTimeSheet
from .services.analitika import analitika, stepen_spreme
from .services.attendance import ClockEvent


def zaposleni(code, *, pol="M", rodjen=datetime.date(1980, 6, 1), zaposlen=datetime.date(2010, 1, 1),
              zanimanje="", skola="", oj="41", status="Na neodredjeno vreme", aktivan=True):
    return Employee.objects.create(
        employee_code=code, first_name="Ime", last_name=f"Prezime{code}", position="Radnik", department_code=int(oj),
        org_unit_code=oj, gender=pol, date_of_birth=rodjen, date_of_joining=zaposlen, job_title=zanimanje,
        education=skola, status_name=status, is_active=aktivan)


def korisnik(ime, *kodovi, employee=None):
    user = get_user_model().objects.create_user(ime, password="x", employee=employee)
    if kodovi:
        uloga = Role.objects.create(name=f"Uloga {ime}", slug=f"uloga-{ime}")
        uloga.permissions.add(*[PermissionCode.objects.get_or_create(code=k)[0] for k in kodovi])
        user.roles.add(uloga)
    return user


@override_settings(ALLOWED_HOSTS=["testserver"])
class KomentarProlazaTests(TestCase):
    def setUp(self):
        self.employee = zaposleni(501)
        self.user = korisnik("prolazi", employee=self.employee)
        self.client.force_login(self.user)
        dogadjaji = [  # 04.05. ulaz i izlaz (u redu), 05.05. samo ulaz (problem); 06.05. bez prolaza
            ClockEvent(1, 501, "P", "I", "", 1, datetime.datetime(2026, 5, 4, 8, 0), 1),
            ClockEvent(1, 501, "P", "I", "", 2, datetime.datetime(2026, 5, 4, 16, 0), 2),
            ClockEvent(1, 501, "P", "I", "", 3, datetime.datetime(2026, 5, 5, 8, 0), 1),
        ]
        patch("hr.views.get_clock_events", return_value=dogadjaji).start()
        self.addCleanup(patch.stopall)

    def strana(self):
        return self.client.get(reverse("hr:work_time_sheet"), {"month": 5, "year": 2026})

    def posalji(self, dan, tekst):
        return self.client.post(reverse("hr:work_time_sheet"),
                                {"action": "komentar_prolaza", "month": 5, "year": 2026, "dan": dan, "tekst": tekst})

    def test_komentar_se_nudi_samo_za_problem_i_radni_dan_bez_prolaza(self):
        redovi = {r["day"]: r for r in self.strana().context["clock_attendance_rows"]}
        self.assertFalse(redovi[4]["moze_komentar"])   # u redu
        self.assertTrue(redovi[5]["moze_komentar"])    # problem
        self.assertTrue(redovi[6]["moze_komentar"])    # sreda bez prolaza
        self.assertFalse(redovi[2]["moze_komentar"])   # subota
        self.assertFalse(redovi[1]["moze_komentar"])   # 1. maj, praznik

    def test_upis_izmena_brisanje_i_prilog(self):
        odgovor = self.posalji(5, "  Zaboravio   karticu na izlazu ")
        self.assertRedirects(odgovor, reverse("hr:work_time_sheet") + "?month=5&year=2026#evidencija-prolaza",
                             fetch_redirect_response=False)
        komentar = KomentarProlaza.objects.get(employee=self.employee, datum=datetime.date(2026, 5, 5))
        self.assertEqual((komentar.tekst, komentar.created_by), ("Zaboravio karticu na izlazu", self.user))
        self.assertContains(self.strana(), "Zaboravio karticu na izlazu")
        self.posalji(5, "Teren u Obrenovcu")
        komentar.refresh_from_db()
        self.assertEqual(komentar.tekst, "Teren u Obrenovcu")
        sheet = WorkTimeSheet.objects.get(employee=self.employee, year=2026, month=5)
        prilog = self.client.get(reverse("hr:work_time_sheet_attendance_print", args=[sheet.pk]))
        self.assertContains(prilog, "Teren u Obrenovcu")
        self.posalji(5, "   ")
        self.assertFalse(KomentarProlaza.objects.exists())

    def test_odobrena_lista_je_zakljucana_i_neispravan_dan(self):
        self.strana()
        WorkTimeSheet.objects.filter(employee=self.employee).update(status=WorkTimeSheet.Status.APPROVED)
        self.posalji(5, "Kasno")
        self.assertFalse(KomentarProlaza.objects.exists())
        WorkTimeSheet.objects.filter(employee=self.employee).update(status=WorkTimeSheet.Status.DRAFT)
        self.posalji(32, "Nema tog dana")
        self.assertFalse(KomentarProlaza.objects.exists())

    def test_tudja_lista_samo_uz_dozvolu(self):
        drugi = korisnik("drugi")
        self.client.force_login(drugi)
        adresa = reverse("hr:employee_work_time_sheet", args=[self.employee.pk])
        odgovor = self.client.post(adresa, {"action": "komentar_prolaza", "month": 5, "year": 2026, "dan": 5, "tekst": "x"})
        self.assertEqual(odgovor.status_code, 403)
        self.assertFalse(KomentarProlaza.objects.exists())


class StepenSpremeTests(TestCase):
    def test_razvrstavanje_po_zanimanju_pa_skoli(self):
        for zanimanje, skola, stepen in (
                ("doktor tehn.nauka - gradj.", "", "VIII"), ("dipl.inz.gradjevinarstva", "", "VII"),
                ("strukovni master inz.gradj.", "", "VII"), ("STRUKOVNI MASTER INŽ. GRAÐ.", "", "VII"),
                ("strukovni inženjer mašinstva", "", "VI"), ("inzenjer gradjevinarstva", "", "VI"),
                ("metaloglodac-specijalista", "", "V"), ("gradjevinski tehnicar", "", "IV"),
                ("ugostiteljski tehnièar", "", "IV"), ("mat.gimn.matem.ili prir.smera", "", "IV"),
                ("vozac", "", "III"), ("lice bez zanim. i str.spreme", "", "I–II"),
                ("", "GRADJEVINSKI FAKULTET", "VII"), ("", "VISA MASINSKA SKOLA", "VI"),
                ("", "", "Nije razvrstano")):
            with self.subTest(zanimanje=zanimanje, skola=skola):
                self.assertEqual(stepen_spreme(zanimanje, skola)[0], stepen)


class AnalitikaTests(TestCase):
    def setUp(self):
        dan = datetime.date(2026, 10, 5)
        self.dan = dan
        zaposleni(1, pol="Z", rodjen=datetime.date(1996, 1, 1), zanimanje="dipl.inz.gradjevinarstva")      # 30
        zaposleni(2, pol="M", rodjen=datetime.date(1981, 12, 1), zanimanje="gradjevinski tehnicar")         # 44
        zaposleni(3, pol="M", rodjen=datetime.date(1960, 1, 1), zanimanje="doktor hem.nauka",               # 66
                  zaposlen=datetime.date(1990, 1, 1), oj="82", status="Na odredjeno vreme")
        zaposleni(4, pol="F", rodjen=datetime.date(1900, 1, 1), zanimanje="vozac")                         # bez starosti
        zaposleni(5, pol="M", rodjen=datetime.date(1970, 1, 1), aktivan=False)                             # neaktivan

    def test_pokazatelji(self):
        a = analitika(Employee.objects.filter(is_active=True), self.dan)
        s = a["sazetak"]
        self.assertEqual((a["ukupno"], s["zene"], s["muski"], s["bez_starosti"]), (4, 2, 2, 1))
        self.assertEqual((s["prosek"], s["medijana"], s["najmladji"], s["najstariji"]), (Decimal("46.7"), Decimal("44"), 30, 66))
        self.assertEqual((s["prosek_zene"], s["prosek_muski"], s["pred_penzijom"]), (Decimal("30.0"), Decimal("55.0"), 1))
        rasponi = {r["naziv"]: (r["ukupno"], r["zene"], r["muski"]) for r in a["starosni"]}
        self.assertEqual((rasponi["30–39"], rasponi["40–49"], rasponi["65 i više"]), ((1, 1, 0), (1, 0, 1), (1, 0, 1)))
        self.assertEqual({r["oznaka"]: r["ukupno"] for r in a["stepeni"]}, {"VIII": 1, "VII": 1, "IV": 1, "III": 1})
        self.assertEqual(s["visoka"], 2)
        self.assertEqual({r["naziv"]: r["ukupno"] for r in a["odnos"]}, {"Na neodredjeno vreme": 3, "Na odredjeno vreme": 1})
        ukrstanje = {r["oznaka"]: r["rasponi"] for r in a["ukrstanje"]}
        self.assertEqual(sum(ukrstanje["III"]), 0)  # vozač bez ispravnog datuma rođenja

    def test_ekran_pdf_excel_i_dozvola(self):
        bez = korisnik("bez-analitike")
        self.client.force_login(bez)
        self.assertEqual(self.client.get(reverse("hr:analitika")).status_code, 403)
        self.client.force_login(korisnik("kadrovska", "hr:analitika"))
        odgovor = self.client.get(reverse("hr:analitika"))
        self.assertEqual(odgovor.status_code, 200)
        for tekst in ("Starosna piramida", "Histogram starosti", "Starost po stručnoj spremi", "Staž u Institutu"):
            self.assertContains(odgovor, tekst)
        pdf = self.client.get(reverse("hr:analitika"), {"izvoz": "pdf", "oj": "82"})
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        strane = len(re.findall(rb"/Type /Page[^s]", pdf.content))
        self.assertGreaterEqual(strane, 2)  # zaglavlje i grafikoni na svakoj strani, koliko sadržaja treba
        excel = self.client.get(reverse("hr:analitika"), {"izvoz": "xlsx"})
        knjiga = load_workbook(BytesIO(excel.content))
        self.assertEqual(knjiga.sheetnames, ["Sažetak", "Starost", "Stručna sprema", "Starost po spremi",
                                             "Staž u Institutu", "Organizacione jedinice", "Radni odnos", "Razvrstavanje"])
        self.assertEqual(knjiga["Sažetak"]["B6"].value, 4)
