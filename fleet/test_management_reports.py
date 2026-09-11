from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.models import OrganizationalUnit
from fleet.models import JobCode, Lease, TransactionOMV, VehicleHolding
from fleet.test_vehicle_onboarding import vehicle
from fleet.support.management_reports import vehicle_insurance_rows, fuel_report_rows, group_fuel_rows, supplier_parts_rows
from nabavka.models import EufItemSnapshot, GoodsSnapshot


class ManagementReportTests(TestCase):
    day=date(2026,9,10)

    def setUp(self):
        self.user=get_user_model().objects.create_superuser('report-admin',password='test')
        self.a=OrganizationalUnit.objects.create(code='A1',name='Prvi',center='01')
        self.b=OrganizationalUnit.objects.create(code='B1',name='Drugi',center='02')

    def car(self,suffix='1',year=2018):
        v=vehicle(suffix);v.year_of_manufacture=year;v.save();return v

    def fuel(self,car,day,product='OMV EVRO DIZEL',gross='1234.56',currency='RSD'):
        return TransactionOMV.objects.create(vehicle=car,transaction_date=timezone.make_aware(datetime.combine(day,datetime.min.time())),license_plate_no='BG123-AA',card='test',issuer='test',customer='test',product_inv=product,quantity=Decimal('10'),gross_cc=Decimal(gross) if gross is not None else None,supplier_currency=currency)

    def test_casco_strictly_older_than_seven_and_not_limited_to_owned(self):
        old=self.car();self.car('2',2019);self.car('3',2099)
        Lease.objects.create(vehicle=old,partner_code='1',partner_name='Test',job_code='A',contract_number='1',current_payment_amount=1,start_date=date(2020,1,1),end_date=date(2027,1,1))
        rows,stats=vehicle_insurance_rows(self.user,{'as_of':self.day},casco=True)
        self.assertEqual([r['id'] for r in rows],[old.pk])
        self.assertEqual(rows[0]['age'],8)
        self.assertEqual(rows[0]['ownership'],'Korišćenje po ugovoru')
        self.assertEqual(stats['unknown_year'],1)

    def test_owned_report_distinguishes_evidence_and_fallback_and_excludes_contract(self):
        owned,fallback,rented=self.car(),self.car('2'),self.car('3')
        VehicleHolding.objects.create(vehicle=owned,basis='owned',start_date=date(2020,1,1))
        Lease.objects.create(vehicle=rented,partner_code='1',partner_name='Test',job_code='A',contract_number='1',current_payment_amount=1,start_date=date(2020,1,1),end_date=date(2027,1,1))
        data={'as_of':self.day,'ownership':'all'}
        rows,stats=vehicle_insurance_rows(self.user,data)
        self.assertEqual({r['id'] for r in rows},{owned.pk,fallback.pk});self.assertEqual(stats['assumed'],1)
        rows,_=vehicle_insurance_rows(self.user,{**data,'ownership':'confirmed'})
        self.assertEqual([r['id'] for r in rows],[owned.pk])

    def test_historical_fuel_scope_and_month_boundary(self):
        car=self.car()
        JobCode.objects.create(vehicle=car,organizational_unit=self.a,assigned_date=date(2026,1,1))
        JobCode.objects.create(vehicle=car,organizational_unit=self.b,assigned_date=date(2026,2,1))
        self.fuel(car,date(2026,1,31));self.fuel(car,date(2026,2,1))
        user=get_user_model().objects.create_user('restricted',allowed_center_codes='01')
        data={'date_from':date(2026,1,1),'date_to':date(2026,2,28),'fuel_type':'fuel'}
        rows=fuel_report_rows(user,data)
        self.assertEqual(len(rows),1);self.assertEqual(rows[0]['job'],'A1')
        self.assertEqual(rows[0]['month'],'2026-01')
        all_rows=fuel_report_rows(self.user,data)
        self.assertEqual({r['job'] for r in all_rows},{'A1','B1'})

    def test_fuel_filter_excludes_adblue_and_retains_unlinked_and_signed_amounts(self):
        self.fuel(None,self.day);self.fuel(None,self.day,'AdBlue');self.fuel(None,self.day,'Putarina')
        self.fuel(None,self.day,'ED MAXXMOTION','-100')
        data={'date_from':self.day,'date_to':self.day,'fuel_type':'fuel'}
        rows=fuel_report_rows(self.user,data)
        self.assertEqual(len(rows),2)
        self.assertEqual(sum(r['gross'] for r in rows),Decimal('1134.56'))
        self.assertTrue(all(r['job']=='Bez šifre' for r in rows))
        self.assertEqual(len(fuel_report_rows(self.user,{**data,'fuel_type':'all'})),4)

    def test_grouping_keeps_currency_products_and_missing_amounts(self):
        self.fuel(None,self.day,gross='0');self.fuel(None,self.day,gross=None,currency='EUR')
        data={'date_from':self.day,'date_to':self.day,'fuel_type':'fuel'}
        groups=group_fuel_rows(fuel_report_rows(self.user,data),'month')
        self.assertEqual(len(groups),2)
        by_currency={r['currency']:r for r in groups}
        self.assertEqual(by_currency['RSD']['gross'],Decimal('0'))
        self.assertIsNone(by_currency['EUR']['gross'])
        self.assertEqual(by_currency['EUR']['missing_amount'],1)

    def test_parts_uses_line_values_and_one_source_only(self):
        for key,value in [('a','50'),('b','-10')]:
            EufItemSnapshot.objects.create(source_key=key,partner_name='AUTO DEKI',partner_pib='123',document_date=self.day,item_name='Deo',quantity=1,value=Decimal(value),total=1000)
        GoodsSnapshot.objects.create(source_key='goods',partner_name='AUTO DEKI',partner_code=123,document_date=self.day,article_name='Deo',quantity=1,debit=50)
        rows=supplier_parts_rows({'source':'uf','supplier':'123','date_from':self.day,'date_to':self.day})
        self.assertEqual(len(rows),2);self.assertEqual(sum(r['value'] for r in rows),40)

    def test_export_preserves_numeric_values_and_neutralizes_invoice_formula(self):
        EufItemSnapshot.objects.create(source_key='a',partner_name='AUTO DEKI',partner_pib='123',document_date=self.day,item_name='=1+1',value=Decimal('1234.56'))
        self.client.force_login(self.user)
        response=self.client.get(reverse('supplier_parts_report'),{'source':'uf','supplier':'123','date_from':str(self.day),'date_to':str(self.day),'export':'xlsx'})
        self.assertEqual(response.status_code,200)
        book=load_workbook(BytesIO(response.content));sheet=book['Pregled']
        self.assertEqual(sheet.cell(2,11).value,1234.56)
        self.assertNotEqual(sheet.cell(2,7).data_type,'f')
        self.assertTrue(sheet.cell(2,7).value.startswith("'="))

    def test_invalid_filters_do_not_silently_show_unfiltered_data(self):
        self.client.force_login(self.user)
        response=self.client.get(reverse('fleet_fuel_report'),{'date_from':'2026-02-01','date_to':'2026-01-01','fuel_type':'all','group_by':'month'})
        self.assertEqual(response.status_code,400)
        self.assertEqual(response.context['row_count'],0)

    def test_permissions_and_default_pages(self):
        user=get_user_model().objects.create_user('no-permission',password='test')
        self.client.force_login(user)
        for name in ('casco_report','owned_insurance_report','fleet_fuel_report','supplier_parts_report'):
            self.assertEqual(self.client.get(reverse(name)).status_code,403)
        self.client.force_login(self.user)
        for name in ('casco_report','owned_insurance_report','fleet_fuel_report','supplier_parts_report'):
            response=self.client.get(reverse(name))
            self.assertEqual(response.status_code,200,(name,response.content[:500]))

    def test_vehicle_scope_uses_as_of_date_and_excludes_archived(self):
        car=self.car();old=self.car('2');old.otpis=True;old.save()
        JobCode.objects.create(vehicle=car,organizational_unit=self.a,assigned_date=date(2020,1,1))
        JobCode.objects.create(vehicle=car,organizational_unit=self.b,assigned_date=date(2027,1,1))
        user=get_user_model().objects.create_user('scoped',allowed_center_codes='01')
        rows,_=vehicle_insurance_rows(user,{'as_of':self.day,'ownership':'all'})
        self.assertEqual([r['id'] for r in rows],[car.pk])
        self.assertEqual(rows[0]['center'],'01')
