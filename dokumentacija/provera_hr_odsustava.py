"""Read-only schema/sample assessment; no names, JMBGs or medical causes in output."""
import os,sys,json
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT));os.environ.setdefault('DJANGO_SETTINGS_MODULE','ims_erp.settings.development')
import django
django.setup()
from django.db import connections
from hr.models import Employee

book=load_workbook(next(ROOT.glob('SickLeave*.xlsx')),read_only=True,data_only=True)
rows=list(book.active.iter_rows(values_only=True));book.close()
result={'headers':list(rows[0]),'rows':len(rows)-1,'status_counts':dict(Counter(str(r[3]) for r in rows[1:])),
 'id_lengths':dict(Counter(len(str(r[0])) for r in rows[1:])),
 'jmbg_lengths':dict(Counter(len(str(r[2])) for r in rows[1:])),
 'start_date_examples':list(dict.fromkeys(str(r[5]) for r in rows[1:]))[:3],
 'end_date_blank':sum(not r[6] for r in rows[1:]),'day_count_examples':list(dict.fromkeys(str(r[7]) for r in rows[1:]))[:8]}
identities=Counter((str(v or '').strip().zfill(13)) for v in Employee.objects.values_list('personal_number',flat=True))
result['jmbg_matching']={'unique':sum(identities[str(r[2]).strip().zfill(13)]==1 for r in rows[1:]),
 'missing':sum(identities[str(r[2]).strip().zfill(13)]==0 for r in rows[1:]),
 'ambiguous':sum(identities[str(r[2]).strip().zfill(13)]>1 for r in rows[1:])}
with connections['default'].cursor() as c:
    for prefix in ('[BazaLDIMS]', '[putgeo-server].[BazaLDIMS]'):
        try:
            c.execute(f'SELECT TABLE_NAME,COLUMN_NAME,DATA_TYPE FROM {prefix}.INFORMATION_SCHEMA.COLUMNS WHERE TABLE_NAME IN (%s,%s) ORDER BY TABLE_NAME,ORDINAL_POSITION',['Godmor','GodmorKor'])
            cols=list(c.fetchall())
            if cols:
                result['annual_leave_source']=prefix
                result['annual_leave_columns']=[list(r) for r in cols]
                for table in ('Godmor','GodmorKor'):
                    c.execute(f'SELECT COUNT_BIG(*) FROM {prefix}.[dbo].[{table}]')
                    result[table+'_rows']=c.fetchone()[0]
                break
        except Exception as exc:
            result[prefix+'_status']=type(exc).__name__
result['checked_at']='2026-09-11'
(ROOT/'izvestaji/provera_hr_odsustava_20260911.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
