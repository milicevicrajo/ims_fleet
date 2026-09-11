from datetime import datetime
from pathlib import Path
from shutil import copy2
from docx import Document

root = Path(__file__).resolve().parents[1]
stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
heading = 'Bolovanja — DataTable i detalj zaposlenog, 11.09.2026.'
for name in ['Kadrovi - bolovanja i godisnji odmori.docx',
             'Presek stanja - kratak pregled.docx', 'Presek stanja - primedbe i odgovori.docx']:
    path = root / name
    doc = Document(path)
    if any(p.text == heading for p in doc.paragraphs):
        continue
    copy2(path, root / 'dokumentacija/arhiva' / f'{path.stem} - pre DataTable bolovanja {stamp}.docx')
    pictures = len(doc.inline_shapes)
    doc.add_heading(heading, level=1)
    doc.add_paragraph('Dodat je DataTable na pregled bolovanja: pretraga, izbor broja redova, straničenje i sortiranje; početni redosled je od najnovijeg početka bolovanja. Datumi su prikazani kao d.m.Y, uz hronološko sortiranje. Uklonjeno je prethodno serversko straničenje, tako da DataTable obrađuje ceo skup izabran postojećim filterima.')
    doc.add_paragraph('Ime povezane osobe vodi na detalj zaposlenog uz odgovarajuću dozvolu; sopstveni profil ostaje dostupan vlasniku naloga. Nepovezano bolovanje nema link ka izmišljenom zaposlenom. Prikaz prazne tabele i ponovna inicijalizacija provereni su bez DataTables upozorenja.')
    doc.add_paragraph('Provera: šest postojećih testova HR prikaza prolazi. U Chromium-u su provereni prikaz 23 bolovanja i 23 linka, datumsko sortiranje, pretraga, prazan rezultat i zaštita od ponovne inicijalizacije. Nisu menjani modeli, migracije ili poslovni podaci.')
    doc.save(path)
    assert len(Document(path).inline_shapes) == pictures
    print(name)
