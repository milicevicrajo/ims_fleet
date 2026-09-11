"""Build the reviewable Word document from the adjacent Markdown source."""
from pathlib import Path
import re

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT

root = Path(__file__).resolve().parents[1]
source = root/'dokumentacija/Analiza aplikacije i procedure rada - IMS flota.md'
target = root/'Analiza aplikacije i procedure rada - IMS flota.docx'
lines = source.read_text(encoding='utf-8').splitlines()
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
section.left_margin, section.right_margin = Cm(1.8), Cm(1.8)
section.header_distance, section.footer_distance = Cm(.8), Cm(.8)
section.different_first_page_header_footer = True
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08
for name, size in [('Title',32),('Subtitle',17),('Heading 1',17),('Heading 2',12.5)]:
    style=doc.styles[name]
    style.font.name='Calibri';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('203F56')
    style.paragraph_format.space_before=Pt(12)
    style.paragraph_format.space_after=Pt(8)
    style.paragraph_format.keep_with_next=True
doc.styles['Heading 1'].paragraph_format.page_break_before=True
doc.core_properties.title='IMS flota — Analiza aplikacije i procedure rada'
doc.core_properties.subject='Stanje, odgovornosti, procedure, izveštaji i plan razvoja'
doc.core_properties.author='IMS — radni dokument'
doc.core_properties.keywords='flota, garaža, procedure, servisi, nabavka, troškovi'

def field(paragraph, instruction):
    run=paragraph.add_run()
    begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');run._r.append(begin)
    text=OxmlElement('w:instrText');text.set(qn('xml:space'),'preserve');text.text=instruction;run._r.append(text)
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');run._r.append(end)

def hyperlink(paragraph,label,url):
    rel=paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
    run=OxmlElement('w:r');props=OxmlElement('w:rPr')
    color=OxmlElement('w:color');color.set(qn('w:val'),'206889');props.append(color)
    run.append(props);text=OxmlElement('w:t');text.text=label;run.append(text);link.append(run);paragraph._p.append(link)

def inline(paragraph,text):
    for part in re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\([^\)]+\))',text):
        if part.startswith('**') and part.endswith('**'):
            paragraph.add_run(part[2:-2]).bold=True
        elif (match:=re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)',part)):
            hyperlink(paragraph,*match.groups())
        else:paragraph.add_run(part)

def table(rows):
    count=len(rows[0]);t=doc.add_table(rows=1,cols=count)
    t.style='Table Grid';t.autofit=False
    widths=[8.7,8.7] if count==2 else [4.3,6.2,6.9]
    if rows[0][0]=='Uloga iz baze':widths=[3.4,3.0,11.0]
    if rows[0][0]=='Izvor':widths=[6.2,11.2]
    for col,width in zip(t.columns,widths):col.width=Cm(width)
    for index,values in enumerate(rows):
        cells=t.rows[0].cells if index==0 else t.add_row().cells
        trpr=t.rows[index]._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
        if index==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for col,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Cm(widths[col])
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'203F56' if index==0 else ('F0F5F8' if index%2 else 'FFFFFF'));pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ('top','bottom','left','right'):
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.0
            inline(p,value)
            for run in p.runs:
                run.font.size=Pt(9.2)
                if index==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

header=section.header.paragraphs[0]
header.text='IMS FLOTA  |  ANALIZA I PROCEDURE RADA'
header.runs[0].font.size=Pt(8);header.runs[0].font.color.rgb=RGBColor.from_string('647684')
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('10.09.2026.  •  Verzija 1.0  |  Strana ');field(footer,'PAGE');footer.add_run(' / ');field(footer,'NUMPAGES')
for r in footer.runs:r.font.size=Pt(8)

doc.add_paragraph('IMS FLOTA',style='Title')
doc.add_paragraph('Analiza aplikacije\ni procedure rada',style='Subtitle')
doc.add_paragraph('Vozilo • Zaduženja • Putovanja • Garaža • Nabavka • Troškovi')
doc.add_paragraph('Presek 10.09.2026.\nVerzija 1.0 — dokument za razmatranje uprave, garaže, Nabavke i finansija.')
p=doc.add_paragraph();inline(p,'**Cilj:** kontrola troškova, odobravanje intervencija, potvrđivanje stvarnih radova i planiranje malog i velikog servisa.')
p=doc.add_paragraph();inline(p,'**POSTOJI** — provereno u aplikaciji/bazi.\n**DOGOVORENO** — potvrđeno sa korisnikom.\n**PREDLOG** — postupak ili dorada koja tek treba da se usvoji i sprovede.')
doc.add_paragraph('Dokument se zasniva na čitanju postojeće aplikacije i baze. Opis predloženog postupka nije tvrdnja da ga aplikacija već sprovodi. U ovoj analizi nisu menjani poslovni zapisi, prava niti aplikacioni tokovi.')
doc.add_page_break()
p=doc.add_paragraph('Sadržaj',style='Title')
field(doc.add_paragraph(),'TOC \\o "1-2" \\h \\z \\u')

start=next(i for i,line in enumerate(lines) if line.startswith('# 1.'))
i=start
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].lstrip().startswith('|'):
            cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',c) for c in cells):rows.append(cells)
            i+=1
        table(rows);continue
    if line.startswith('# '):doc.add_paragraph(line[2:],style='Heading 1')
    elif line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 2')
    elif line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet');inline(p,line[2:])
    elif re.match(r'^\d+\. ',line):
        p=doc.add_paragraph();p.paragraph_format.left_indent=Cm(.5);p.paragraph_format.first_line_indent=Cm(-.5);inline(p,line)
    else:
        p=doc.add_paragraph();inline(p,line)
    i+=1
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
doc.save(target)
check=Document(target)
assert len([p for p in check.paragraphs if p.style.name=='Heading 1'])==15
assert len(check.tables)>=12
assert '186.900,00' in '\n'.join(p.text for p in check.paragraphs)
print(f'Generated: {target.name}; {len(check.tables)} tables; {len(source.read_text(encoding="utf-8").split())} source words.')
