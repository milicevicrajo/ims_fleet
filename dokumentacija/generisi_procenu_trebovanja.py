"""Generate the assessment Word and record the conclusion in existing status docs."""
from pathlib import Path
from datetime import datetime
from shutil import copy2
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
source = root/'dokumentacija/Procena prenosa trebovanja u Nabavku.md'
target = root/'Procena prenosa trebovanja u Nabavku.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width,sec.page_height = Cm(21),Cm(29.7)
sec.left_margin=sec.right_margin=Cm(1.8)
sec.top_margin=sec.bottom_margin=Cm(1.8)
sec.header_distance=sec.footer_distance=Cm(.8)
normal=doc.styles['Normal']
normal.font.name='Calibri';normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.05
for name,size in [('Title',24),('Heading 1',15),('Heading 2',12)]:
    st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size)
    st.font.color.rgb=RGBColor.from_string('203F56')
    st.paragraph_format.keep_with_next=True
    st.paragraph_format.space_before=Pt(10)
    st.paragraph_format.space_after=Pt(6)
doc.core_properties.title='Procena prenosa trebovanja u Nabavku'
doc.core_properties.subject='Requisition: podaci, arhitektura, migracija, procena rada'
doc.core_properties.author='IMS — radni dokument'
sec.header.paragraphs[0].text='IMS FLOTA  /  NABAVKA — PROCENA PRENOSA TREBOVANJA'
sec.header.paragraphs[0].style=doc.styles['Caption']
footer=sec.footer.paragraphs[0]
footer.add_run('11.09.2026.  •  Procena, bez primene migracije  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def inline(p,text):
    for i,part in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
        p.add_run(part).bold=bool(i%2)

lines=source.read_text(encoding='utf-8').splitlines()
i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:
        i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',c) for c in cells):rows.append(cells)
            i+=1
        table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Table Grid'
        table.autofit=False
        widths=[5.1,6.0,6.3] if len(rows[0])==3 else [8.3,9.1]
        for col,w in zip(table.columns,widths):col.width=Cm(w)
        for n,vals in enumerate(rows):
            row=table.rows[0] if n==0 else table.add_row()
            props=row._tr.get_or_add_trPr()
            props.append(OxmlElement('w:cantSplit'))
            if n==0:props.append(OxmlElement('w:tblHeader'))
            for cell,val,w in zip(row.cells,vals,widths):
                cell.width=Cm(w);p=cell.paragraphs[0];inline(p,val)
                p.paragraph_format.space_after=Pt(4)
                p.paragraph_format.space_before=Pt(3)
                for run in p.runs:
                    run.font.size=Pt(9.5)
                    if n==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
                if n==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'203F56');cell._tc.get_or_add_tcPr().append(shade)
        doc.add_paragraph()
        continue
    if line.startswith('# '):inline(doc.add_paragraph(style='Title'),line[2:])
    elif line.startswith('## '):inline(doc.add_paragraph(style='Heading 1'),line[3:])
    elif line.startswith('- '):inline(doc.add_paragraph(style='List Bullet'),line[2:])
    else:inline(doc.add_paragraph(),line)
    i+=1
doc.save(target)
assert len(Document(target).tables)==6
print(target.name)

stamp=datetime.now().strftime('%Y%m%d_%H%M%S')
marker='Procena prenosa modela Requisition — 11.09.2026.'
for filename in ['Presek stanja - kratak pregled.docx','Presek stanja - primedbe i odgovori.docx','Garaza IMS - tok rada i povezivanje.docx']:
    path=root/filename
    state=Document(path)
    if any(p.text==marker for p in state.paragraphs):continue
    copy2(path,root/'dokumentacija/arhiva'/f'{path.stem} - pre procene trebovanja {stamp}.docx')
    images=len(state.inline_shapes)
    state.add_heading(marker,level=1)
    state.add_paragraph('Baza je ponovo dostupna. Proverene definicije izvora: trebovanja Flote koriste sif_dok=40, a roba Nabavke sif_dok=10. Postojeći GoodsSnapshot nije kopija trebovanja. Predlog je preneti Requisition u Nabavku kao poseban model izdatog materijala, uz početno zadržavanje iste SQL tabele i ID-jeva; nije potrebno prepisivati dokumente ili odmah uvoditi opšti magacin.')
    state.add_paragraph('Za očuvanje: 2.980 stavki, 2.978 veza sa vozilima, 2.209 kategorija, 1.833 kilometraže i 13 napomena. Sadašnji izvor nema 2.209 istorijskih stavki iz 2024–2025, ali ima 164 nove koje lokalno nisu preuzete. Prenos i naknadna sinhronizacija moraju se kontrolisati odvojeno. Osam postojećih vrednosti razlikuje se od zaokruženog proizvoda količine i cene; sačuvati originalne iznose.')
    state.add_paragraph('Procena preporučenog obuhvata: 24–40 radnih sati (približno 3–5 radnih dana jednog programera), uz uslove i provere navedene u „Procena prenosa trebovanja u Nabavku.docx“. Obuhvata model, migraciono stanje, prikaze/obračune, uvoz, prava i proveru. Ne obuhvata pun tok odobravanja garaže ili prenos servisa. Status: završena procena; migracije i poslovni upisi nisu izvršeni.')
    state.save(path)
    assert len(Document(path).inline_shapes)==images
    print(f'Dopunjeno: {filename}')
