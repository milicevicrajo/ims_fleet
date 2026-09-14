from datetime import datetime
from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
heading = 'Kadrovi — godišnji odmori i usklađen izgled, 11.09.2026.'
sections = [
    ('Radna lista', 'Uklonjen je slobodan unos napomene. U redu ostaje padajući izbor vrste rada ili odsustva, koji ulazi i u štampu. Raniji slobodan tekst nije obrisan iz baze, ali se više ne prikazuje niti menja kroz ovu formu. Sistemske informacije o bolovanju i putnim nalozima uz prolaske ostaju dostupne.'),
    ('Tabele godišnjih odmora', 'Uvedeni su AnnualLeaveAllowance za dodelu dana iz Godmor i AnnualLeaveDecision za rešenja iz GodmorKor. Posebna evidencija AnnualLeaveSync čuva vreme, obuhvat, pokretača i rezultate uspešne sinhronizacije. Primenjene su migracije hr 0008 i 0009.'),
    ('Povezivanje i osvežavanje', 'Izvor je IMS, sif_pred=1, na [putgeo-server].[BazaLDIMS]. Dodela se prepoznaje po preduzeću, godini i šifri zaposlenog rasif; rešenje dodatno po tačnom početnom datumu/vremenu iz izvornog ključa. Korigovan završni datum ili broj dana ažurira postojeći zapis. Promenjen početak predstavlja novi izvorni ključ; prethodni zapis se čuva kao povučen.'),
    ('Učitani podaci', 'Preuzeta je istorija 2003–2026: 7.643 dodele i 13.838 rešenja. Za 2026. postoje 328 dodela i 293 rešenja; svi su povezani sa zaposlenima. Ponovna provera 2026. nije pronašla nove ni izmenjene zapise.'),
    ('Istorijski nepovezani zapisi', 'U celoj istoriji 3.076 dodela i 5.252 rešenja nema odgovarajuću osobu u sadašnjoj lokalnoj evidenciji zaposlenih. Sačuvani su pod izvornom šifrom, bez kreiranja izmišljenih zaposlenih. Među njima je i jedno rešenje sa izvornom šifrom 0, koje se nikada automatski ne povezuje. Ostale šifre mogu se povezati narednom sinhronizacijom kada zaposleni postoji lokalno.'),
    ('Način rada', 'U Kadrovima je novi pregled Godišnji odmori, sa izborom godine, pretragom i opcijom prikaza povučenih zapisa. Dugme Sinhronizuj učitava izabranu godinu, odnosno celu istoriju kada su izabrane Sve godine. Dostupna je i komanda manage.py sync_annual_leave, sa opcijama --year 2026 i --dry-run. Sinhronizacija se pokreće ručno; periodični raspored nije podešen.'),
    ('Pregled kod zaposlenog', 'Dodato je dugme na linku ka zaposlenom i kartica Godišnji odmori na njegovom detalju. Kartica prikazuje dodelu dana i rešenja kroz godine, sa datumima d.m.Y i hronološkim sortiranjem. Zaposleni u svom profilu vidi samo svoju evidenciju. Globalni pregled i sinhronizacija dostupni su superuseru i ulozi Uprava, kroz zasebne dozvole.'),
    ('Pouzdanost sinhronizacije', 'Ponovljeni uvoz ne duplira zapise. Zapisi koji nestanu iz izvora označavaju se kao povučeni i ostaju u istoriji. Uvoz jedne godine ne menja druge godine. Neispravan ili dupliran ključ, nedostupan izvor ili greška pri upisu vraća celu lokalnu transakciju. Potpuno prazan izvor ne briše postojeće podatke. Istovremena pokretanja su zaključana na nivou baze.'),
    ('Usklađen izgled Kadrova', 'Spisak zaposlenih, bolovanja, godišnji odmori i šifrarnici koriste stil modula Mobilni: zajednički gradijentni baner, svetle kartice filtera, ujednačene tabele, oznake statusa i dugmad sa ikonama. Spisak zaposlenih ima kombinovanje statusa, OJ, vrste primaoca i pretrage. Filteri se primenjuju pri promeni izbora ili završetku kucanja, kao u Mobilnim. Usklađene su i forme uvoza bolovanja i izmene šifrarnika.'),
    ('Preostali saldo — naredna faza', 'Odobreni dani iz rešenja nisu automatski stvarno iskorišćeni dani. Saldo još nije obračunat. Za njega treba usaglasiti izvor ranijeg korišćenja ili početno stanje, potvrđivanje radne liste, godinu prava iz koje se dani troše i pravilo brojanja radnih dana. Trenutni prikaz jasno govori da su u pitanju dodela i rešenja.'),
    ('Provera', 'Prošlo je 90 automatizovanih testova HR-a i dozvola. U stvarnoj bazi potvrđeni su uvoz i ponavljanje sinhronizacije. U Chromium-u provereni su baneri, tabele, dugmad, filteri, sortiranje datuma, odsustvo slobodne napomene i prazan rezultat pretrage, bez JavaScript grešaka. Za snimke prikaza korišćeni su izmišljeni podaci.'),
]

doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Cm(1.6)
section.left_margin = section.right_margin = Cm(1.8)
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10)
style.paragraph_format.space_after = Pt(6)
for name in ('Title','Heading 1','Heading 2'):
    doc.styles[name].font.color.rgb = RGBColor.from_string('14395B')
doc.add_heading('Kadrovi — godišnji odmori i izgled',0)
doc.add_paragraph('Presek stanja · 11.09.2026. | Završene izmene')
for index,(title,text) in enumerate(sections):
    if index == 6:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.add_run(title+'. ').bold = True
    p.add_run(text)
path = ROOT/'Kadrovi - godisnji odmori i izgled.docx'
if path.exists():
    copy2(path,ROOT/'dokumentacija/arhiva'/f'{path.stem} - prethodno {stamp}.docx')
doc.save(path)
print(path.name)

for name in ['Kadrovi - radne liste i sifrarnici.docx','Kadrovi - bolovanja i godisnji odmori.docx',
             'Presek stanja - kratak pregled.docx','Presek stanja - primedbe i odgovori.docx']:
    path = ROOT/name
    doc = Document(path)
    if any(p.text == heading for p in doc.paragraphs):
        continue
    copy2(path,ROOT/'dokumentacija/arhiva'/f'{path.stem} - pre godisnjih i izgleda {stamp}.docx')
    pictures = len(doc.inline_shapes)
    doc.add_heading(heading,1)
    doc.add_paragraph('Ova dopuna zamenjuje raniji predlog godišnjih odmora i raniji opis slobodne napomene. Tabele i sinhronizacija su sada uvedeni, a slobodna napomena uklonjena iz unosa i štampe radne liste.')
    for index in [0,1,3,4,5,6,8,9,10]:
        title,text = sections[index]
        p = doc.add_paragraph()
        p.add_run(title+'. ').bold = True
        p.add_run(text)
    doc.add_paragraph('Detalji su objedinjeni u novom dokumentu „Kadrovi - godisnji odmori i izgled.docx“.')
    doc.save(path)
    assert len(Document(path).inline_shapes) == pictures
    print(path.name)
