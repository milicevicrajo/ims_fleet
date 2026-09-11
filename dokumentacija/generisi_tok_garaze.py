"""Focused garage workflow proposal; no application or business-data changes."""
from pathlib import Path
from datetime import datetime
from shutil import copy2
from docx import Document
from docx.shared import Cm, Pt, RGBColor

root = Path(__file__).resolve().parents[1]
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(1.7)
sec.left_margin = sec.right_margin = Cm(1.9)
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
for name in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.color.rgb = RGBColor.from_string('203F56')
doc.core_properties.title = 'Garaža IMS — tok rada i povezivanje'
doc.core_properties.subject = 'Usmeren predlog procedure: nalog, odobrenje, nabavka, materijal, izvršenje'

def p(text):
    doc.add_paragraph(text)

def h(text):
    doc.add_heading(text, level=2)

doc.add_heading('Garaža IMS — tok rada i povezivanje', level=0)
p('10.09.2026. | Predlog za dogovor, još nije implementiran')
p('Dopuna obuhvata: korisnik je potvrdio da ostali moduli odgovaraju potrebama i rade. Sadašnji razvoj treba usmeriti na logičan tok garaže i njegove veze sa postojećom Nabavkom. Šira analiza ostaje referentni dokument; nije nalog za preuređivanje ostalih modula.')
h('Osnovna odluka')
p('Jedna intervencija ima jedan centralni nalog garaže. Prijava započinje taj predmet; nalog se dopunjava dijagnozom, obuhvatom i načinom izvršenja. Trebovanja, zahtevi za nabavku/uslugu, fakture i knjiženja prate nalog. Svaki dokument zadržava svoj broj i svoju svrhu. Nabavka dela, izdavanje dela i izvršenje rada nisu tri popravke.')
h('1. Prijava kvara ili potreba za redovnim servisom')
p('Zaposleni prijavljuje problem Veljoviću, a garaža evidentira vozilo, datum, kilometražu, prijavioca i opis. Za mali/veliki servis razlog otvaranja je plan održavanja; ne mora postojati kvar. Početni zapis ima status Prijavljeno. Ako nema potrebe za radom, prijava se zatvara uz obrazloženje.')
h('2. Procena i priprema naloga')
p('Garaža upisuje nalaz, predložene radove, vrstu intervencije, delove i predlog izvršenja: u IMS, van IMS ili kombinovano. Proverava postojeću servisnu istoriju i stanje materijala. Nalog dobija broj kada se pripremi za odobrenje; nije potrebno ponovo prepisivati prijavu u drugu nepovezanu formu.')
h('3. Odobrenje uprave')
p('Uprava vidi šta se predlaže, zašto, gde bi se radilo, procenu troška i eventualno upozorenje za ponavljanje servisa. Odobrava konkretan obuhvat ili vraća nalog na dopunu. Nabavka može prethodno pribaviti ponudu radi procene; odobrenje treba da prethodi naručivanju i izvršenju dogovorenog posla. Ako je potrebna plaćena dijagnostika, prvo se odobrava dijagnostika, pa dodatni radovi kada budu poznati.')

doc.add_page_break()
doc.add_heading('Izvršenje i završetak', level=1)
h('4A. Popravka u IMS')
p('Ako potrebni delovi postoje na stanju: uz nalog se povezuje trebovanje/izdavanje iz magacina i garaža izvodi rad. Ne otvara se zahtev za kupovinu samo zato što postoji popravka. Ako delovi nedostaju: iz naloga se otvara postojeći zahtev za nabavku, sa prenetim vozilom i vezom ka nalogu. Nabavka sprovodi kupovinu; prijem i izdavanje prate stvarni tok robe.')
p('Ako deo ide kroz magacin, faktura dokazuje kupovinu, a izdavanje raspoređuje materijal na konkretan rad. Ako postoji dozvoljen direktan prijem/utrošak bez magacina, koristi se odgovarajući stvarni dokument — ne pravi se izmišljeno trebovanje. Potvrđuju se stvarno utrošene količine, uz povrat viška kada ga ima.')
h('4B. Popravka van IMS')
p('Iz istog naloga otvara se postojeći zahtev za uslugu. Nabavka vodi ponudu, izbor/naručivanje i fakturu prema postojećem postupku. Garaža prati predaju vozila, izvršenje kod dobavljača i preuzimanje. Spoljni servis može obezbediti i rad i delove; tada IMS ne pravi sopstveno trebovanje za njegove delove. Ako IMS daje svoje delove, njihovo izdavanje se vezuje za isti nalog.')
h('4C. Kombinovana intervencija')
p('Primer: IMS demontira sklop, spoljni servis ga popravi, IMS ga ugradi. Ostaje isti nalog, sa jasno navedenim unutrašnjim i spoljnim radom. Uz njega može biti zahtev za delove i zahtev za uslugu. Mesto izvršenja nije razlog za otvaranje dve nepovezane istorije iste popravke.')
h('5. Garaža potvrđuje stvarno izvršenje')
p('Veljović ili ovlašćeno lice potvrđuje datum završetka, kilometražu, stvarno obavljene radove, izvođača i da li je vozilo vraćeno u upotrebu. Za rad van IMS potvrđuje prijem i obuhvat na osnovu provere i servisne dokumentacije. Promena obuhvata/troška u odnosu na odobreno vraća se upravi na dopunu odobrenja.')
p('Mali/veliki servis postaje potvrđen servisni događaj na osnovu ove potvrde. Sam zahtev, faktura ili trebovanje ne pomeraju automatski datum sledećeg servisa. Datum izvršenja čuva se odvojeno od datuma dokumenta; postojeći prikaz datuma fakture/trebovanja ostaje smislen kao dokumentacioni podatak.')
h('6. Dokumentacija i trošak se dopunjavaju')
p('Faktura može stići posle vraćanja vozila u upotrebu. Rad tada ima status Izvršeno, uz oznaku da nedostaje dokumentacija. Nabavka/finansije dopunjavaju dokumente i troškove, a nalog se potom zatvara. Jedan isti iznos ne sabira se ponovo kao faktura, knjiženje i trebovanje.')
p('Predlog glavnih statusa: Prijavljeno → U pripremi → Čeka odobrenje → Odobreno → U radu → Izvršeno → Zatvoreno. Vraćeno na dopunu i Otkazano čuvaju obrazloženje. Čeka delove, čeka spoljni servis i čeka fakturu mogu biti oznake razloga čekanja, bez umnožavanja glavnih statusa.')

doc.add_page_break()
doc.add_heading('Šta ostaje u Floti i šta se dopunjuje', level=1)
h('Trebovanja — zadržati kao evidenciju materijala')
p('Postojeći Requisition sadrži stavke, artikle, količine, cene i datum, i već ima opcionu vezu ka nalogu garaže. Potreban je za istoriju materijala i troškove, ali nije samostalni nalog za popravku. Predlog naziva pregleda: „Trebovanja / utrošeni materijal“, uz jasno tumačenje izvorne vrste dokumenta. Materijal za opšte potrebe garaže ne treba prisilno vezivati za izmišljeno vozilo ili intervenciju.')
h('Servisi van IMS — zadržati izvor, razjasniti naziv')
p('Postojeći ServiceTransaction je evidencija knjiženih stavki sa kontima, partnerom, brojem i vezanim dokumentom. Nije model potvrđenog rada spoljnog servisa. Trenutni pregled prikazuje sve zapise te tabele; naziv „Popravke van IMS“ zato ne objašnjava njen pun obuhvat. Predlog naziva: „Knjižene usluge i troškovi“. Poseban pregled stvarnih popravki van IMS treba zasnovati na nalozima i potvrđenim radovima, sa filterom mesta izvršenja.')
p('Za korisnika je cilj ukidanje odvojenog unosa trebovanja i popravke van IMS u Floti: dokument se jednom preuzima/formira u svom izvornom postupku, a Flota ga prikazuje preko veze. Postojeće tabele i istoriju treba sačuvati tokom prelaska, dok se ne potvrdi da novi izvori pokrivaju sve potrebne podatke i obračune. To nije zahtev da stare tabele zauvek ostanu aktivne. Ne unosi se ista popravka još jednom samo radi njenog prikaza u Floti.')
h('Najmanja korisna tehnička dorada')
p('Proširiti postojeći Kvar kao predmet intervencije: status, prijavilac, nalaz/obuhvat, način izvršenja, broj i datum naloga, odobravalac i vreme odobrenja, potvrđivač i datum/km izvršenja. Dodati istoriju važnih promena. Prijava i izdati nalog mogu biti faze istog zapisa, uz očuvan izvorni opis prijave i trag izmena odobrenog obuhvata.')
p('Iskoristiti postojeće veze ProcurementCase.garage_order i Requisition.kvar. Za knjižene stavke dopuniti opcionu vezu ka nalogu ili raspodelu kada dokument pokriva više radova. Veza fakture preko zahteva ostaje; samostalna faktura može se naknadno povezati bez izmišljanja prethodnog zahteva. Širi model raspodele uvoditi samo tamo gde jedan dokument stvarno pokriva više naloga/vozila.')
p('Iz naloga ponuditi radnje „Pripremi nalog“, „Pošalji na odobrenje“, „Zahtev za delove“, „Zahtev za uslugu“, „Poveži trebovanje/dokument“ i „Potvrdi izvršenje“, prema statusu i ulozi. Generičke liste delova ostaju izmenjiv predlog; samo otvaranje naloga ne treba da upisuje stvarni utrošak.')
p('Migracija treba da bude dopunska: sačuvati istoriju, dodati polja/veze i istoriju statusa. Stari zapisi ne dobijaju izmišljeno odobrenje ili potvrdu izvršenja; označavaju se kao istorijski/nepotvrđeni. Prva dorada je tok naloga i povezivanje dokumenata; servisni intervali zatim koriste potvrđena izvršenja.')
h('Jedino operativno razjašnjenje za povezivanje materijala')
p('Pri razradi treba potvrditi gde se danas izdaje zvanično trebovanje i ko ga potvrđuje. Ako nastaje u postojećem magacinskom/ERP sistemu, Flota ga preuzima i povezuje. Ako taj postupak još ne postoji, tek tada treba dogovoriti formiranje dokumenta u aplikaciji. Ne treba praviti drugi magacin samo radi povezivanja naloga.')

doc.add_page_break()
doc.add_heading('Dopuna: Nabavka bez dvostrukog unosa', level=1)
p('Dodatni pregled koda, 10.09.2026. Korisnik je pojasnio nameru direktorke: ukinuti zasebne evidencije za ručno unošenje istih dokumenata u Floti i podatke dobijati povezivanjem sa Nabavkom. To je realan cilj. Uklanjanje zasebnog ekrana za unos i fizičko brisanje istorijske tabele nisu isti korak.')
h('Šta Nabavka već ima')
p('ProcurementCase već ima vozilo, tip intervencije i vezu ka nalogu garaže. ProcurementItem omogućava povezivanje sa EUF fakturom, UF fakturom ili robnom stavkom. Postoje uvezene fakture, UF stavke i robne stavke sa datumom, količinom i cenom. Povezivanje izvora na zahtev/stavku već radi kroz posebne radnje. Postoji i štampa „Trebovanje materijala“ iz stavki zahteva.')
p('Štampa koristi količine sa zahteva, broj zahteva i tekući datum štampe. Taj prikaz sam ne evidentira stvarno izdavanje robe, izdate količine, elektronsku potvrdu izdavaoca/primaoca ili povrat. Potpisan papir može biti prateći dokaz, ali u aplikaciji još nema zasebnog događaja potvrde izdavanja. Traženu i stvarno izdatu količinu treba razlikovati.')
h('Šta još nije povezano u jednu evidenciju')
p('Trebovanja Flote uvoze se iz dbo.fleet_trebovanja; roba Nabavke iz dbo.nbv_roba. Requisition ima vezu sa vozilom/nalogom i vrednost nabavke. GoodsSnapshot ima podatke o robi i vrsti dokumenta, ali nema direktnu vezu sa vozilom/nalogom niti isto eksplicitno polje vrednosti trebovanja. Ne može se pretpostaviti da svaki robni red predstavlja izdavanje za popravku ili da cena puta količina uvek zamenjuje izvornu vrednost troška.')
p('Postojeći tab Servisi čita povezane EUF/UF dokumente i Requisition, ali trenutno ne uključuje ProcurementItem.goods_item u dokumentacioni pregled. Troškovi vozila i dalje koriste Requisition i ServiceTransaction. Zbog toga samo povezivanje robe u Nabavci ili uklanjanje stare tabele još ne prebacuje servisni pregled i troškove na novi izvor.')
p('Veza sa dokumentom trenutno nije raspodela njegovih količina i iznosa. Jedna stavka zahteva ima jedan izbor izvora, a postoji i mogućnost povezivanja istog izvora sa više stavki zahteva. To je korisno za pronalaženje dokumentacije, ali zbir cele fakture ne sme da se obračuna za svaku povezanu stavku. Za delimična izdavanja ili dokument koji pokriva više naloga potrebne su zasebne veze sa stvarnim količinama/iznosima.')
h('Kako izgleda željeni unos')
p('Primer spoljne popravke: garaža formira i dobija odobrenje naloga, iz njega otvara zahtev za uslugu; Nabavka obrađuje zahtev i povezuje već uvezenu fakturu. Veza zahtev–nalog već određuje vozilo. Garaža unosi samo šta je izvršeno, datum i kilometražu. Flota prikazuje isti dokument i pripadajući iznos; ne otvara se dodatna forma „Popravka van IMS“ radi prepisivanja broja fakture, dobavljača i cene.')
p('Primer internog rada: za delove sa stanja povezuje se postojeće izdavanje sa nalogom, bez zahteva za novu kupovinu. Ako delovi nedostaju, zahtev za nabavku nastaje iz naloga; posle kupovine povezuje se stvarno izdavanje za taj rad. Primer: kupljeno 20 filtera za magacin, izdat jedan za konkretan auto — tom radu pripada izdat jedan filter, ne puna faktura za 20. Kupovina i izdavanje ostaju dva različita događaja, ali se nijedan ne prepisuje u Floti.')
h('Dorada koja omogućava ukidanje paralelnih tabela')
p('Predlog je zajednička veza naloga i izvornog dokumenta/stavke: izvor i identitet dokumenta, nalog, uloga dokumenta (ponuda, kupovina, izdavanje, povrat, usluga, knjiženje), pripadajuća količina/iznos, ko je povezao i kada. Koristiti postojeće veze gde su dovoljne; dodatnu tabelu veza/raspodela uvesti za direktno povezivanje i delimična izdavanja, umesto nove kopije svih podataka dokumenta. Potreban je stabilan ključ izvora: sadašnji ključ robnog snimka računa se iz svih vrednosti reda, pa izmena izvorne cene/količine može stvoriti novi identitet snimka.')
p('Nalog prikazuje objedinjene dokumente, materijal i troškove iz tih veza. Jedan dokument može imati više poslovnih veza, ali pojedinačni iznos troška mora imati jasno poreklo i kontrolisanu raspodelu. Faktura i odgovarajuće knjiženje predstavljaju dva prikaza istog troška, dok faktura za magacin i kasnije izdavanje imaju različitu svrhu. Obračun mora poštovati te razlike.')
p('Redosled prelaska: potvrditi obuhvat izvora i vrste robnih dokumenata; dopuniti veze i prikaze uz nalog; preneti postojeće veze vozila, kategorije, kilometražu i napomene; usaglasiti stare i nove iznose na reprezentativnim dokumentima; prebaciti čitanje i ugasiti odvojeni unos; tek potom arhivirati ili ukloniti zamenjene tabele i uvoze. Ako deo istorije nije pokriven Nabavkom, ostaje dostupan kao istorijski izvor bez obaveze dvostrukog budućeg unosa.')
h('Granica provere')
p('Modeli, forme povezivanja, štampa, izvori uvoza i obračuni pregledani su u kodu. Novo poređenje stvarnih redova nije izvršeno: pokušaj povezivanja sa SQL serverom završio se greškom dostupnosti/istekom vremena povezivanja. Potpuno preklapanje dbo.fleet_trebovanja i dbo.nbv_roba, kao ni pokrivenost starih knjiženja fakturama Nabavke, nije potvrđeno. U ovom koraku nisu menjani aplikacija, tabele ili poslovni podaci.')

h('Preciziranje: odluka o Django modelima u fleet aplikaciji')
p('Korisnik pita za modele, ne samo za ekrane. Ciljni model ne treba da sadrži dve paralelne kopije istog dokumenta. Requisition može da se ukine nakon objedinjavanja njegovih podataka u zajedničkom izvoru robnih dokumenata i prenošenja veza ka vozilu/nalogu. Postojeći GoodsSnapshot još nije potvrđena potpuna zamena: treba obezbediti izdavanje, vrednost utroška i očuvanje istorijskih dopuna.')
p('Service i ServiceTransaction su različiti modeli. Service je jednostavna evidencija vozila, tipa, datuma, troška, dobavljača i opisa. Predlog je ukinuti ga kao paralelnu servisnu evidenciju kada prošireni nalog garaže preuzme izvršene radove, a povezani dokumenti trošak. Nabavka sama ne zamenjuje potvrdu stvarnog servisa. ServiceTransaction sadrži knjižene troškove i nije automatski zamenljiv fakturom; može se ukloniti tek kada zajednička evidencija sačuva potrebna knjiženja, obuhvat i postojeće obračune. ServiceType je zaseban šifarnik i ne briše se usput.')

target = root / 'Garaza IMS - tok rada i povezivanje.docx'
doc.save(target)
assert len(Document(target).paragraphs) > 35
print(target.name)

marker = 'Usmeren obuhvat — tok garaže, 10.09.2026.'
stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
for name in ['Presek stanja - kratak pregled.docx', 'Presek stanja - primedbe i odgovori.docx']:
    path = root / name
    state = Document(path)
    if any(p.text == marker for p in state.paragraphs):
        continue
    copy2(path, root / 'dokumentacija/arhiva' / f'{path.stem} - pre toka garaze {stamp}.docx')
    images = len(state.inline_shapes)
    state.add_heading(marker, level=1)
    state.add_paragraph('Korisnik je precizirao da ostali moduli odgovaraju potrebama i rade. Fokus predloga je garaža: prijava/potreba za servisom → priprema naloga → odobrenje uprave → interni/spoljni rad, uz Nabavku samo kada je potrebna → potvrda izvršenja → dopuna dokumentacije i zatvaranje.')
    state.add_paragraph('Trebovanja ostaju izvor podataka o materijalu; postojeći „Servisi van IMS“ ostaju izvor knjiženih troškova. Predlog je jasnije imenovanje pregleda i povezivanje sa centralnim nalogom, bez brisanja tabela i bez dvostrukog unosa popravke. Postojeće veze sa Nabavkom i trebovanjem treba iskoristiti. Detalji su u „Garaza IMS - tok rada i povezivanje.docx“. Ovo je dokumentovan predlog; aplikacija nije menjana u ovom koraku.')
    state.save(path)
    assert len(Document(path).inline_shapes) == images
    print(f'Dopunjeno: {name}')

link_marker = 'Dopuna — povezivanje sa Nabavkom bez duplog unosa'
for name in ['Presek stanja - kratak pregled.docx', 'Presek stanja - primedbe i odgovori.docx']:
    path = root / name
    state = Document(path)
    if any(p.text == link_marker for p in state.paragraphs):
        continue
    copy2(path, root / 'dokumentacija/arhiva' / f'{path.stem} - pre objedinjavanja izvora {stamp}.docx')
    images = len(state.inline_shapes)
    state.add_heading(link_marker, level=1)
    state.add_paragraph('Po pojašnjenju zahteva direktorke cilj je ukidanje posebnog unosa trebovanja i popravki van IMS u Floti. Nabavka već podržava povezivanje EUF/UF faktura i robnih stavki, kao i štampu trebovanja iz zahteva. Nalog garaže treba da koristi te dokumente preko veza, uz potvrdu stvarnog rada. Stare tabele čuvaju se tokom prelaska; to nije odluka da trajno ostanu aktivne.')
    state.add_paragraph('Pre fizičkog uklanjanja tabela treba potvrditi da roba Nabavke pokriva izdavanja materijala, dopuniti vezu izdatih količina sa nalogom i prebaciti obračune. Štampa zahtevanog materijala trenutno nije elektronska potvrda izdavanja. U kodu su različiti izvori dbo.fleet_trebovanja i dbo.nbv_roba; novo poređenje njihovih redova nije završeno jer SQL server nije dostupan. Predlog i granice provere dopisani su u „Garaza IMS - tok rada i povezivanje.docx“. Nisu menjani aplikacija ili poslovni podaci.')
    state.save(path)
    assert len(Document(path).inline_shapes) == images
    print(f'Dopunjeno povezivanje: {name}')
