from datetime import datetime
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
heading = 'Kadrovi — radne liste i šifrarnici, 11.09.2026.'

implemented = [
    ('Radna lista zaposlenog', 'Superuser iz detalja zaposlenog otvara njegovu mesečnu radnu listu, bira period, čuva podatke i štampa. Korisnik bez tog ovlašćenja preko nove adrese ne može da pristupi tuđoj listi.'),
    ('Dva nova podatka iz HR-a', 'View dbo.hr_employee već sadrži sif_prim i naz_prim. Model Employee proširen je poljima recipient_code i recipient_name: šifra i izvorni naziv vrste primaoca. Podaci se povezuju po postojećoj šifri zaposlenog rasif. To nije radno mesto niti status zaposlenja.'),
    ('Šifrarnici u Kadrovima', 'Elementi RL imaju tri kartice: elementi radne liste, vrste rada/odsustva i vrste primaoca. Dostupni su dodavanje, izmena naziva, podešavanje veza i deaktiviranje. Pregled elemenata ima pretragu i sortiranje. Pristup imaju superuser i postojeća uloga Uprava kroz zasebne dozvole.'),
    ('Izbor u radnoj listi', 'U svakom redu bira se vrsta rada ili odsustva za vrstu primaoca zaposlenog, a ispod ostaje dodatna slobodna napomena. Obe informacije ulaze u štampu. Ranije napomene ostaju sačuvane. Neaktivne stavke nisu ponuđene za novi unos, a ranije sačuvana veza ostaje dostupna.'),
    ('Obračunske šifre', 'Jedna napomena, npr. Bolovanje, može imati više šifara elsif. Zaposleni zato bira vrstu odsustva; ovom izmenom se ne određuje automatski tačna obračunska šifra, procenat naknade ili iznos. Topli obrok, regres i minuli rad inicijalno nisu ponuđeni kao izbor vrste rada/odsustva.'),
    ('Uvoz i provera', 'Iz datoteke elementi_rl (2).xlsx učitano je 27 elemenata i 13 vrsta rada/odsustva, od kojih je 10 za izbor zaposlenog. HR izvor i Excel ukupno daju 3 vrste primaoca. Nova polja osvežena su za 368 zaposlenih; 2 postojeće osobe nisu pronađene u izvoru i nisu automatski povezane. Ponovljeni uvoz čuva lokalne izmene šifrarnika. Prošlo je 76 automatizovanih testova HR-a i dozvola.'),
]

annual = [
    ('Preporučeno rešenje', 'Uvesti dve lokalne tabele kao sinhronizovanu kopiju: AnnualLeaveAllowance za dodeljene dane iz Godmor i AnnualLeaveDecision za rešenja iz GodmorKor. Dok postojeći sistem izdaje dodelu i rešenja, on ostaje njihov jedini izvor. U ovoj aplikaciji se isti podatak ne unosi ponovo.'),
    ('Šta se povezuje', 'Dodela: preduzeće sif_pred, godina god, zaposleni rasif, ukupan broj dana i izvorne komponente. Rešenje: ista veza sa zaposlenim i godinom, period od/do, izvorni broj dana i komentar. Sačuvati vreme sinhronizacije. Pre upisa usaglasiti stabilan ključ rešenja, jer samo ime zaposlenog ili broj reda nisu pouzdan identifikator.'),
    ('Stvarno korišćenje', 'Rešenje prikazuje odobreni period; iskorišćenost se utvrđuje iz radne liste. Predlog je da nove potvrđene liste koriste kategoriju Godišnji odmor, uz godinu prava iz koje se dani troše. Jedan datum se ne sme sabrati više puta zbog više redova ili šifara posla. Ne pretvarati sate u dane automatskim deljenjem sa 8 bez usaglašenog rasporeda rada.'),
    ('Istorija i stanje', 'Pre početka računanja salda potreban je pouzdan izvor ranijeg korišćenja ili potvrđeno početno stanje na datum prelaska. U suprotnom bi stari godišnji odmori izgledali kao neiskorišćeni. Predloženi prikaz: dodeljeno, stvarno iskorišćeno, odobreno a još nekorišćeno i preostalo, odvojeno po godini prava.'),
    ('Utvrđeno u bazi', 'Za 2026. Godmor ima 328 zapisa o dodeli, a GodmorKor 293 zapisa o rešenjima; svih 328 šifara iz dodele odgovara postojećim zaposlenima. Ovi brojevi nisu broj iskorišćenih dana. Značenje polja ostalo nije potvrđeno i nije korišćeno kao gotov saldo.'),
    ('Status', 'Sinhronizacija godišnjih odmora i računanje salda još nisu uvedeni. Ovo je predlog naredne faze; sada je dodat izbor Godišnji odmor u radnoj listi kao potrebna osnova.'),
]

questions = [
    'Da li prihvatamo da zaposleni bira samo vrstu rada/odsustva, a obračun određuje tačnu šifru kada jedna vrsta ima više elemenata? To je sadašnji način rada.',
    'Iz koje radne liste preuzimamo ranije stvarno korišćenje godišnjeg odmora, odnosno ko potvrđuje početno stanje na datum prelaska?',
    'Ko potvrđuje novu radnu listu i kako se označava godina prava kada se koristi preneti odmor? To je potrebno pre računanja konačnog salda.',
]

doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Cm(1.5)
section.left_margin = section.right_margin = Cm(1.8)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(5)
for style in ['Title','Heading 1','Heading 2']:
    doc.styles[style].font.color.rgb = RGBColor.from_string('203F56')
doc.add_heading('Kadrovi — radne liste i šifrarnici',0)
doc.add_paragraph('Presek stanja · 11.09.2026. | Urađeno i predlog godišnjih odmora')
for title, text in implemented:
    p = doc.add_paragraph()
    p.add_run(title + '. ').bold = True
    p.add_run(text)
doc.add_paragraph('Primenjene migracije: fleet 0076 i hr 0006–0007. Postojeća HR sinhronizacija ubuduće ažurira i dva nova polja. Izvorni view nije menjan. Učitavanje šifrarnika: manage.py import_work_time_catalog "elementi_rl (2).xlsx" --sync-recipients. Postojeća dodela i rešenja godišnjeg odmora nisu menjani.').runs[0].font.size = Pt(9)
doc.add_page_break()
doc.add_heading('Godišnji odmori — tok bez duplog unosa',1)
for title, text in annual:
    p = doc.add_paragraph()
    p.add_run(title + '. ').bold = True
    p.add_run(text)
doc.add_heading('Pitanja za dogovor',2)
for question in questions:
    doc.add_paragraph(question,style='List Bullet')
path = ROOT/'Kadrovi - radne liste i sifrarnici.docx'
if path.exists():
    copy2(path,ROOT/'dokumentacija/arhiva'/f'{path.stem} - prethodno {stamp}.docx')
doc.save(path)
print(path.name)

for name in ['Kadrovi - bolovanja i godisnji odmori.docx','Presek stanja - kratak pregled.docx','Presek stanja - primedbe i odgovori.docx']:
    path = ROOT/name
    doc = Document(path)
    if any(p.text == heading for p in doc.paragraphs):
        continue
    copy2(path,ROOT/'dokumentacija/arhiva'/f'{path.stem} - pre sifrarnika RL {stamp}.docx')
    pictures = len(doc.inline_shapes)
    doc.add_heading(heading,1)
    for title,text in implemented:
        p = doc.add_paragraph()
        p.add_run(title+'. ').bold = True
        p.add_run(text)
    doc.add_paragraph('Godišnji odmori — predlog, još nije implementiran: lokalna sinhronizacija Godmor za dodelu i GodmorKor za rešenja; stvarno korišćenje iz potvrđenih radnih lista, po godini prava. Potrebno je usaglasiti izvor ranijeg korišćenja i početno stanje, odobravanje i prenos dana. Rešenja se ne tretiraju kao već iskorišćeni dani. Detalji i pitanja su u dokumentu „Kadrovi - radne liste i sifrarnici.docx“.')
    doc.save(path)
    assert len(Document(path).inline_shapes) == pictures
    print(path.name)
