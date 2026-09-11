from datetime import datetime
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.shared import Cm, Pt, RGBColor

root = Path(__file__).resolve().parents[1]
target = root / 'Kadrovi - bolovanja i godisnji odmori.docx'
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Cm(1.8)
section.left_margin = section.right_margin = Cm(1.9)
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(7)
for name in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.color.rgb = RGBColor.from_string('203F56')

doc.add_heading('Kadrovi — bolovanja i godišnji odmori', level=0)
doc.add_paragraph('Presek 11.09.2026. | Bolovanja: implementirano. Godišnji odmori: provereni izvori i predlog daljeg povezivanja.')
doc.add_heading('1. Šta je završeno za bolovanja', level=1)
doc.add_paragraph('Dodati su HR modeli SickLeave i SickLeaveImport, lista bolovanja i forma za uvoz RFZO Excel datoteke. Bolovanje se prepoznaje po RFZO ID-ju, a zaposleni se povezuje po JMBG-u iz Employee.personal_number. Ponovni uvoz ažurira postojeći zapis, bez stvaranja još jednog istog bolovanja.')
doc.add_paragraph('Primenjene su migracije hr.0004_sickleaveimport_sickleave_and_more i hr.0005_sick_leave_permissions. Dostavljeni fajl „SickLeave_20260910_070012 - Bolovanja 08 2026 (1).xlsx“ uvezen je sa datumom RFZO izvoza 10.09.2026: 23 bolovanja, sva jednoznačno povezana; 4 statusa Aktivno i 19 Zaključeno. Ponovna provera daje 0 novih, 0 promenjenih i 23 nepromenjena zapisa.')
doc.add_paragraph('Za svaki datum obuhvaćen bolovanjem, u donjoj tabeli Evidencija prolaza na radnoj listi automatski se prikazuje napomena „Bolovanje“, sa RFZO ID-jem i periodom. Oznaka se prikazuje i ako nema prolazaka. Ako postoji putni nalog ili problem sa prolaskom, i on ostaje vidljiv. Uvoz ne popunjava sate, ne menja ručne napomene redova radne liste i ne proglašava bolovanje odrađenim vremenom.')
doc.add_paragraph('Kada izvor prolazaka nije dostupan, bolovanja i putni nalozi se i dalje prikazuju iz lokalne aplikacije. Nepreuzeti sati označeni su crticom, umesto lažnog prikaza nule. Za bolovanje bez završnog datuma prikaz je ograničen poslednjim datumom RFZO izvoza; ne produžava se proizvoljno u budućnost.')
doc.add_heading('2. Kako se koristi', level=1)
for text in [
    'Kadrovi → Bolovanja → Uvezi RFZO Excel.',
    'Izabrati .xlsx datoteku i datum njenog RFZO izvoza. Datum nije isto što i mesec za koji želimo radnu listu.',
    'Posle uvoza proveriti broj novih, ažuriranih i nepovezanih zapisa.',
    'Na radnoj listi za odgovarajući mesec odsustvo je odmah vidljivo uz svaki obuhvaćeni dan, uključujući interval koji prelazi iz jednog meseca u drugi.'
]:
    doc.add_paragraph(text, style='List Number')
doc.add_paragraph('Pregled i uvoz imaju zasebne dozvole hr:sick_leave_list i hr:sick_leave_import. Dodeljene su postojećoj ulozi Uprava; drugim ovlašćenim kadrovskim korisnicima mogu se dodeliti kroz postojeće upravljanje ulogama. Zaposleni na sopstvenoj radnoj listi vidi samo svoja bolovanja. Lista bolovanja ne prikazuje pun JMBG; kod nepovezanog zapisa prikazuje se maskirana oznaka.')
doc.add_heading('3. Pravila povezivanja i ispravki', level=1)
doc.add_paragraph('Ne postojeće ili višestruko podudaranje JMBG-a ne rešava se izborom osobe po imenu. Bolovanje ostaje sačuvano bez veze i označeno za proveru. Posle sređivanja kadrovskog JMBG-a ponavlja se uvoz. Vodeća nula izgubljena u numeričkoj Excel ćeliji može se obnoviti za dvanaestocifreni zapis; neispravne vrednosti odbijaju se.')
doc.add_paragraph('Neispravni datumi, nepoznat status, konflikt istog RFZO ID-ja u jednoj datoteci ili pokušaj povezivanja postojećeg RFZO ID-ja sa drugim JMBG-om zaustavljaju ceo uvoz, bez delimičnog upisa. Stariji izvoz ne prepisuje već noviji zapis. Izostanak bolovanja iz sledećeg fajla ne briše postojeću istoriju. Stornirani/poništeni zapisi ostaju evidentirani, ali se ne prikazuju kao odsustvo.')
doc.add_paragraph('Preuzimaju se podaci potrebni za evidenciju odsustva: RFZO ID, JMBG/veza sa zaposlenim, status, datumi i izvorni broj dana. Broj RFZO dana čuva se kao izvorni podatak; nije automatski pretvoren u radne dane ili sate. Kolone uzroka, lekara i zdravstvene ustanove nisu potrebne za ovu napomenu i nisu prenete u novu evidenciju.')

doc.add_page_break()
doc.add_heading('4. Provera godišnjih odmora', level=1)
doc.add_paragraph('Obe tabele dostupne su preko povezanog servera: [putgeo-server].[BazaLDIMS].[dbo].[Godmor] i [putgeo-server].[BazaLDIMS].[dbo].[GodmorKor]. Pregled je bio isključivo čitanje. Nijedan podatak o godišnjem odmoru nije uvezen ili promenjen u ovom koraku.')
table=doc.add_table(rows=1,cols=3)
table.style='Table Grid'
for c,t in zip(table.rows[0].cells,['Evidencija','Potvrđena struktura','Predložena namena']):c.text=t
for row in [
    ('Godmor','sif_pred, god, rasif, dana1–dana6, dana, ostalo, ranaz','Dodeljeni dani po radniku i godini. Ukupno 7.643 reda; za 2026. ima 328, svi povezivi sa zaposlenima po rasif.'),
    ('GodmorKor','sif_pred, god, rasif, od_datuma, do_datuma, dana, komentar','Rešenja/planirani periodi odmora. Ukupno 13.838 redova; za 2026. ima 293.'),
    ('Radna lista','Izvor i oznaka godišnjeg odmora treba da budu potvrđeni','Stvarno iskorišćeni dani; ne izjednačavati ih sa brojem dana iz rešenja.')
]:
    for c,t in zip(table.add_row().cells,row):c.text=t
doc.add_paragraph('Veza ovih tabela je rasif → Employee.employee_code, uz sif_pred i godinu radi identiteta izvora. U proveri je sif_pred=1 u obe tabele. Bolovanja RFZO koriste JMBG, a godišnji odmori već imaju šifru zaposlenog, pa nije potrebno povezivati ih preko imena.')
doc.add_heading('5. Predlog prikaza i modela za godišnje odmore', level=1)
doc.add_paragraph('Predlog su dve povezane evidencije u HR-u: dodela godišnjeg odmora po godini (npr. AnnualLeaveAllowance) i rešenja sa periodima (AnnualLeaveDecision). Podaci se čitaju/sinhronizuju iz postojećih izvora. Ne treba prepisivati postojeća rešenja u novu ručnu formu.')
doc.add_paragraph('Na profilu zaposlenog može se prikazati: dodeljeno za godinu, rešenja sa periodima, iskorišćeno iz potvrđene radne liste i preostalo po odgovarajućoj godini prava. Dani iz rešenja predstavljaju odobren/planski period; stvarno korišćenje može odstupati. Potrebno je očuvati kojoj godini prava pripadaju korišćeni dani, naročito kod prenetog odmora.')
doc.add_paragraph('Polje Godmor.ostalo ne treba automatski proglasiti stvarnim preostalim danima, jer je korisnik naveo da se stvarno korišćenje računa iz radne liste. Značenje komponenti dana1–dana6 i eventualnog prenosa prava treba potvrditi pre računanja. Isto važi za izuzetke i delimične dane; ne uvoditi proizvoljnu formulu sati / 8.')
doc.add_heading('6. Otvoreno pitanje za nastavak godišnjih odmora', level=1)
doc.add_paragraph('Potrebno je potvrditi koja radna lista je izvor stvarnog korišćenja: postojeća evidencija u BazaLDIMS ili nova radna lista aplikacije. Ako je BazaLDIMS, treba navesti tabelu/polje ili šifru elementa koja označava godišnji odmor. U šemi postoje Zarada/Zarada1 sa poljima elsif, sati i rekol, ali njihova veza sa godišnjim odmorom nije potvrđena; nisu korišćene za obračun niti su čitani iznosi zarada.')
doc.add_paragraph('Nova WorkTimeSheetLine trenutno sadrži sate po datumima, OJ, uslove rada i slobodnu napomenu. Nema strukturiranu vrstu odsustva niti godinu prava na godišnji odmor. Ako ona treba da bude izvor, potrebno je dodati tu oznaku pre pouzdanog obračuna korišćenja. Slobodan tekst „godišnji“ nije pouzdana osnova za sabiranje.')
doc.add_heading('7. Provera i dokazni trag', level=1)
doc.add_paragraph('Prošlo je 54 HR testova: parser RFZO fajla, datumi, JMBG, ponovno učitavanje, ažuriranja, transakcijsko odbijanje pogrešnog fajla, nepovezani zaposleni, dozvole, granice meseca, zajednički prikaz sa putnim nalogom i prikaz kada izvor prolazaka nije dostupan. Lista bolovanja i forma za uvoz dodatno su renderovane nad stvarnom bazom sa HTTP 200. Modeli i migracije su usaglašeni.')
doc.add_paragraph('Zbirni dokazi bez JMBG-a, imena i uzroka bolovanja: izvestaji/provera_hr_odsustava_20260911.json i izvestaji/provera_hr_bolovanja_prikaz_20260911.json. Uvoz je dostupan i komandom manage.py import_rfzo_sick_leave DATOTEKA --source-date YYYY-MM-DD, uz --dry-run za proveru bez upisa.')
doc.save(target)
assert len(Document(target).tables)==1

stamp=datetime.now().strftime('%Y%m%d_%H%M%S')
marker='Kadrovi — RFZO bolovanja i provera godišnjih odmora, 11.09.2026.'
for filename in ['Presek stanja - kratak pregled.docx','Presek stanja - primedbe i odgovori.docx']:
    path=root/filename
    state=Document(path)
    if any(p.text==marker for p in state.paragraphs):continue
    copy2(path,root/'dokumentacija/arhiva'/f'{path.stem} - pre HR odsustava {stamp}.docx')
    images=len(state.inline_shapes)
    state.add_heading(marker,level=1)
    state.add_paragraph('Rešeno: HR modeli bolovanja i istorije uvoza, lista i RFZO Excel uvoz sa povezivanjem po JMBG-u. Primenjene migracije 0004/0005. Uvezen dostavljeni primer: 23 bolovanja, sva povezana, ponovna provera bez duplikata. U radnoj listi se po datumima automatski prikazuje „Bolovanje“ u napomeni uz prolaske, zajedno sa putnim nalozima; sati se ne menjaju. Uvedene su posebne dozvole pregleda/uvoza za Upravu i mogućnost dodele ovlašćenim korisnicima.')
    state.add_paragraph('Provereno: Godmor ima dodelu dana, GodmorKor rešenja, sa rasif vezom ka zaposlenom. Za 2026. ima 328 dodela i 293 rešenja. Stvarno korišćenje ostaje iz radne liste. Otvoreno: koja tabela/šifra u postojećoj evidenciji označava korišćenje godišnjeg odmora, ili da li se taj podatak vodi u novoj radnoj listi. Godišnji odmori nisu automatski obračunati iz rešenja. Detalji su u „Kadrovi - bolovanja i godisnji odmori.docx“.')
    state.save(path)
    assert len(Document(path).inline_shapes)==images
print(target.name)
