"""Document the delivered evaluation feature, preserving existing Word content."""
from datetime import datetime
from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
STAMP = datetime.now().strftime('%Y%m%d_%H%M%S')
HEADING = 'Kadrovi — ocenjivanje zaposlenih, 11.09.2026.'

sections = [
    ('Šta je uvedeno', 'U Kadrovima je dodat pregled Ocenjivanje zaposlenih. Obuhvata mesečne obrasce, podesiv šifrarnik, automatski obračun ličnog koeficijenta K4, istoriju verzija, saglasnosti i izlaz za štampu/PDF na ćirilici. Elementi radne liste ostaju na dnu menija, a Godišnji odmori odmah ispod Spiska zaposlenih.'),
    ('Početni podaci iz Excela', 'Iz datoteke Obrazac za ocenjivanje.xlsx učitane su 2 grupe (menadžment i ostali zaposleni), 6 merila i 60 kombinacija bodova i koeficijenata. Brojčane vrednosti preuzete su iz lista Max_vrednosti, nazivi merila iz obrasca, a opisi nivoa iz lista kriterijumi. Ponovljeni uvoz ne prepisuje ručno podešene vrednosti.'),
    ('Šifrarnik koji se podešava', 'Posebno se uređuju grupe, merila i njihov redosled, bodovi 0–4 sa koeficijentima i opisima, ocenjivači po OJ i pripadnost zaposlenog grupi. Stavke mogu da se deaktiviraju. Postojeći obrasci ostaju vezani za vrednosti koje su sačuvane uz njih.'),
    ('Priprema obrasca', 'Izabere se zaposleni i mesec rezultata. Aplikacija preuzima grupu i ocenjivače iz podešavanja, kada postoje. Prikazuju se šifra zaposlenog, ime, radno mesto, zvanje, stručna sprema i OJ. Nedostajuća grupa i ocenjivači biraju se pre pripreme obrasca. Početna verzija ocenjuje zaposlenog po njegovoj OJ.'),
    ('Ocene i računica', 'Za svako merilo bira se broj bodova uz pripadajuću stimulaciju i može se uneti komentar. K4 = 1 + zbir stimulacija. Neutralni izbor u početnoj skali je 1 bod i stimulacija 0; 0 bodova predstavlja destimulaciju. Početne granice K4 su 0,79000–1,50000 za menadžment i 0,88000–1,30000 za ostale zaposlene. Iznosi se prikazuju sa decimalnim zarezom, a obračun pri čuvanju ponovo proverava server.'),
    ('Snapshot i korekcije', 'Čuvaju se podaci zaposlenog i OJ, grupa, imenovani ocenjivači, nazivi merila, raspoložive skale i opisi, izabrani bodovi i vrednosti, komentari, zbir i K4. Kasnija promena zaposlenog ili šifrarnika ne menja sačuvani obrazac. Korekcija pravi novu verziju; prethodna ostaje u istoriji. Nova verzija izvedena iz stare koristi njen sačuvani šifrarnik. Za primenu novog šifrarnika priprema se novi obrazac za isti period.'),
    ('Saglasnosti', 'Redosled je neposredni rukovodilac → direktor centra → generalni direktor. Svaku potvrdu daje imenovana osoba sa svog naloga; čuvaju se osoba, stvarno vreme potvrde, datum na obrascu, koeficijent i komentar. Direktor i generalni direktor mogu korigovati K4 unutar granica sačuvane skale, uz obavezno obrazloženje. Za promenu ocena rukovodilac formira novu verziju. Nova verzija zahteva nove saglasnosti.'),
    ('Mesečni pregled i štampa', 'Pregled ima filtere meseca, godine i OJ, predloženi K4, koeficijente potvrđene na svakom nivou, poslednju vrednost i status. Podrazumevano se prikazuje poslednja verzija, a istorija se uključuje posebno. Direktor i generalni direktor mogu zbirno potvrditi označene obrasce za koje su imenovani, uz poštovanje redosleda. Svaki obrazac ima zasebnu ćiriličnu štampu; provereni primer staje na jednu A4 stranu. Obimni komentari mogu produžiti izlaz.'),
    ('Ko ima pristup', 'Superuser i Uprava upravljaju šifrarnikom i imaju pregled svih obrazaca. Uvedena je uloga Ocenjivač. Ona omogućava kreiranje za OJ u kojima je korisnik podešen kao rukovodilac i potvrđivanje obrazaca u kojima je imenovan. Korisnički nalog mora biti povezan sa odgovarajućim zaposlenim. Saglasnosti se ne daju u ime druge osobe, uključujući superusera. Zaposleni nema automatski uvid u sve ocene.'),
    ('Šta je potrebno podesiti pre korišćenja', 'U šifrarniku odrediti grupu zaposlenima i ocenjivače po OJ. Zatim odgovarajućim korisnicima dodeliti ulogu Ocenjivač i proveriti vezu naloga sa zaposlenim. Ove lične dodele nisu proizvoljno popunjene iz naziva radnog mesta ili vrste primaoca zarade.'),
    ('Otvoreno: pravilo bodovanja', 'Koji list je merodavan za destimulaciju? Max_vrednosti sadrži skalu 0–4 sa manjim negativnim koeficijentima, dok list kriterijumi opisuje i negativne bodove i drugačije granice. Početna implementacija koristi brojčanu tabelu Max_vrednosti. Potrebno je potvrditi to pravilo i po potrebi uskladiti tekstove i vrednosti kroz šifrarnik pre zvaničnog ocenjivanja.'),
    ('Otvoreno: osnov ocenjivanja i obračun zarade', 'Da li ostaje ocenjivanje po OJ zaposlenog ili treba više poslova/šifara posla sa ponderisanjem po satima iz radne liste? Ko određuje pripadnost menadžmentu i ko su ocenjivači po OJ? Trenutno je sve podesivo ručno i koristi se jedna OJ. Zbirna tabela koeficijenata čuva se u aplikaciji; automatski upis/izvoz za obračun zarade zahteva dogovor o odredišnoj tabeli, periodu zaključavanja i formatu prenosa.'),
    ('Tehnička realizacija i provera', 'Dodato je pet modela za šifrarnike i povezivanje (EvaluationGroup, EvaluationCriterion, EvaluationScale, EvaluationUnitSetup, EvaluationEmployeeSetup), EmployeeEvaluation za verzije obrasca i EvaluationApproval za saglasnosti. Primenjene su migracije hr 0010 i 0011. Prošlo je 105 automatizovanih testova Kadrova i dozvola. U Chromium-u provereni su unos, obračun, detalj, šifrarnik, prazna i popunjena tabela i PDF, bez JavaScript grešaka. Za probne snimke korišćeni su izmišljeni podaci.'),
]


def paragraph(doc, title, text):
    p = doc.add_paragraph()
    p.add_run(title + '. ').bold = True
    p.add_run(text)


doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Cm(1.6)
section.left_margin = section.right_margin = Cm(1.8)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
for name in ['Title','Heading 1','Heading 2']:
    doc.styles[name].font.color.rgb = RGBColor.from_string('14395B')
doc.add_heading('Kadrovi — ocenjivanje zaposlenih',0)
doc.add_paragraph('Presek stanja · 11.09.2026. | Obrazac, šifrarnik i postupak rada')
for index, (title, text) in enumerate(sections):
    if index == 6:
        doc.add_page_break()
    paragraph(doc,title,text)
path = ROOT/'Kadrovi - ocenjivanje zaposlenih.docx'
if path.exists():
    copy2(path,ROOT/'dokumentacija/arhiva'/f'{path.stem} - prethodno {STAMP}.docx')
doc.save(path)
print(path.name)

for name in ['Presek stanja - kratak pregled.docx','Presek stanja - primedbe i odgovori.docx']:
    path = ROOT/name
    doc = Document(path)
    if any(p.text == HEADING for p in doc.paragraphs):
        continue
    copy2(path,ROOT/'dokumentacija/arhiva'/f'{path.stem} - pre ocenjivanja {STAMP}.docx')
    pictures = len(doc.inline_shapes)
    doc.add_heading(HEADING,1)
    for index in [0,1,5,6,10,11]:
        paragraph(doc,*sections[index])
    doc.add_paragraph('Postupak, pristup, podešavanja i provere opisani su u dokumentu „Kadrovi - ocenjivanje zaposlenih.docx“.')
    doc.save(path)
    assert len(Document(path).inline_shapes) == pictures
    print(path.name)
