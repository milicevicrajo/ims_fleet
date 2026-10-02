"""Arhiva, faza 1: delovodni broj, podbroj, storno, zaključenje; Lista kategorija; ekrani; obuhvat po registru."""
import datetime
import shutil
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.db import IntegrityError, transaction
from django.test import TestCase, override_settings
from django.urls import reverse

from arhiva.access import predmeti
from arhiva.models import EvidencionaKnjiga, Kategorija, Predmet, VerzijaListe
from arhiva.services import delovodnik, lista
from core.models import PermissionCode, Role, RolePermission
from organizacija.models import DodelaUloge, OrgNode, OrgNodeVersion

DAN = datetime.date(2026, 10, 1)


def cvor(nivo, sifra, naziv, roditelj=None):
    node = OrgNode.objects.create(company=1, level=nivo)
    OrgNodeVersion.objects.create(node=node, parent=roditelj, segment=sifra[-2:], full_code=sifra, name=naziv,
                                  valid_from=datetime.date(2020, 1, 1))
    return node


def uloga(slug, *kodovi):
    role = Role.objects.create(name=slug, slug=slug)
    for kod in kodovi:
        RolePermission.objects.create(role=role, permission=PermissionCode.objects.get_or_create(code=kod)[0])
    return role


class RegistarMixin:
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.c43 = cvor(OrgNode.LEVEL_CENTER, "43", "Centar za puteve i geotehniku")
        cls.oj431 = cvor(OrgNode.LEVEL_UNIT, "431", "Laboratorija za geotehniku", cls.c43)
        cls.c41 = cvor(OrgNode.LEVEL_CENTER, "41", "Centar za metale")
        cls.posao = cvor(OrgNode.LEVEL_JOB, "430111", "Stručni nadzor", cls.oj431)
        cls.korisnik = get_user_model().objects.create_user("pisarnica", password="x")

    def zavedi(self, **polja):
        podaci = dict(datum=DAN, naslov="Dopis o ispitivanju", glavna_oj=self.c43, zaveo=self.korisnik,
                      korespondent="Putevi Srbije")
        podaci.update(polja)
        return delovodnik.zavedi(**podaci)


class DelovodnikTests(RegistarMixin, TestCase):
    def test_broj_je_centar_i_redni_broj_u_godini(self):
        prvi, drugi = self.zavedi(), self.zavedi(glavna_oj=self.oj431)
        self.assertEqual((prvi.delovodni_broj, drugi.delovodni_broj), ("43-1", "43-2"))  # OJ daje broj svog centra
        self.assertEqual(self.zavedi(glavna_oj=self.c41).delovodni_broj, "41-3")  # jedan niz za ceo Institut
        self.assertEqual(self.zavedi(datum=datetime.date(2027, 1, 4)).delovodni_broj, "43-1")  # nova godina
        self.assertEqual(prvi.akti.get().delovodni_broj, "43-1")  # osnovni akt nosi broj predmeta

    def test_sifra_posla_se_ne_zavodi(self):
        with self.assertRaises(ValidationError):
            self.zavedi(glavna_oj=self.posao)
        self.assertFalse(Predmet.objects.exists())

    def test_podbroj(self):
        predmet = self.zavedi()
        drugi = delovodnik.dodaj_podbroj(predmet, datum=DAN, opis="Odgovor", zaveo=self.korisnik)
        treci = delovodnik.dodaj_podbroj(predmet, datum=DAN, opis="Dopuna", zaveo=self.korisnik)
        self.assertEqual((drugi.delovodni_broj, treci.delovodni_broj), ("43-1/2", "43-1/3"))
        self.assertEqual(drugi.smer, predmet.smer)

    def test_storno_cuva_broj(self):
        predmet = self.zavedi()
        with self.assertRaises(ValidationError):
            delovodnik.storniraj(predmet, razlog=" ", ko=self.korisnik)
        delovodnik.storniraj(predmet, razlog="Pogrešno zavedeno", ko=self.korisnik)
        predmet.refresh_from_db()
        self.assertEqual((predmet.status, predmet.delovodni_broj), (Predmet.Status.STORNIRAN, "43-1"))
        self.assertEqual(self.zavedi().delovodni_broj, "43-2")  # broj se ne koristi ponovo
        with self.assertRaises(ValidationError):
            delovodnik.dodaj_podbroj(predmet, datum=DAN, opis="x", zaveo=self.korisnik)

    def test_zakljucena_i_istorijska_knjiga_bez_upisa(self):
        self.zavedi()
        knjiga = delovodnik.zakljuci_knjigu(EvidencionaKnjiga.objects.get(godina=2026), ko=self.korisnik)
        self.assertEqual(knjiga.broj_akata_pri_zakljucenju, 1)
        with self.assertRaises(ValidationError):
            self.zavedi()
        EvidencionaKnjiga.objects.create(godina=2024, istorijska=True)
        with self.assertRaises(ValidationError):
            self.zavedi(datum=datetime.date(2024, 5, 5))

    def test_isti_osnovni_broj_odbija_baza(self):
        predmet = self.zavedi()
        with self.assertRaises(IntegrityError), transaction.atomic():
            Predmet.objects.create(knjiga=predmet.knjiga, osnovni_broj=1, delovodni_broj="43-1x", datum_zavodjenja=DAN,
                                   naslov="dupli", glavna_oj=self.c43, oznaka_centra="43")

    @override_settings(ARHIVA_FORMAT_BROJA="{broj}/{centar}", ARHIVA_FORMAT_PODBROJA="{osnovni}-{podbroj}")
    def test_format_iz_podesavanja(self):
        predmet = self.zavedi()
        self.assertEqual(predmet.delovodni_broj, "1/43")
        self.assertEqual(delovodnik.dodaj_podbroj(predmet, datum=DAN, opis="x", zaveo=self.korisnik).delovodni_broj, "1/43-2")

    def test_prazan_predmet(self):
        with self.assertRaises(ValidationError):
            self.zavedi(naslov="  ")


def napravi_docx(putanja, redovi):
    """Minimalan .docx sa jednom tabelom (kao Lista kategorija)."""
    def celija(tekst):
        return f"<w:tc><w:p><w:r><w:t xml:space=\"preserve\">{escape(tekst)}</w:t></w:r></w:p></w:tc>"
    telo = "".join("<w:tr>" + "".join(celija(c) for c in red) + "</w:tr>" for red in redovi)
    xml = ('<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/'
           f'wordprocessingml/2006/main"><w:body><w:p/><w:tbl>{telo}</w:tbl></w:body></w:document>')
    with zipfile.ZipFile(putanja, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", xml)


LISTA = [
    ["Redni broj", "Klasifikacione oznake", "Kategorija dokumentarnog materijala", "Rok čuvanja"],
    ["", "20", "Osnivanje i organizacija", ""],
    ["1.", "", "Statut", "Trajno"],
    ["2.", "", "Investicioni krediti", "Trajno"],
    ["3.", "", "Investicioni krediti", "Trajno"],
    ["4.", "", "Projekti dogradnje", "Trajno operativno"],
    ["11.", "", "Investicioni  krediti", "Trajno"],
    ["", "81", "Finansijsko poslovanje", ""],
    ["5.", "", "Knjiga izlaznih računa, Izlazni računi (za obveznike PDV)", "10 godina"],
    ["6.", "", "Izlazni račun", "5 godine"],
    ["7.", "", "Ugovori o zakupu", "1 godina od isteka zakupa"],
    ["8.", "", "Personalni dosijei", "70 godina osim Trajno kod istaknutih ličnosti"],
    ["9.", "", "Zapisi sistema menadžmenta", "3 godine nakon prestanka važenja"],
    ["12.", "", "Akta lica za BZR", "6 godina"],
    ["13.", "", "Akta lica za BZR", "Trajno"],
    ["", "426", "Sertifikacija osoba", ""],
    ["10.", "", "Zapisi sistema menadžmenta", "Trajno"],
]


class ListaTests(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.dir = Path(tempfile.mkdtemp())
        cls.docx = cls.dir / "lista.docx"
        napravi_docx(cls.docx, LISTA)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, ignore_errors=True)
        super().tearDownClass()

    def test_citanje_i_rokovi(self):
        stavke = {s.redni_broj: s for s in lista.procitaj(self.docx)}
        self.assertEqual(len(stavke), 13)
        self.assertEqual((stavke[1].oznaka, stavke[1].grupa), ("20", "Osnivanje i organizacija"))
        self.assertTrue(stavke[1].trajno and not stavke[1].operativno)
        self.assertTrue(stavke[4].trajno and stavke[4].operativno)
        self.assertEqual((stavke[5].rok_godina, stavke[5].pocetak_roka), (10, Kategorija.PocetakRoka.ZAVRSETAK_PREDMETA))
        self.assertEqual((stavke[6].rok_godina, stavke[6].pravopis), (5, True))  # „5 godine”
        self.assertEqual(stavke[7].pocetak_roka, Kategorija.PocetakRoka.PRESTANAK_ZAKUPA)
        self.assertTrue(stavke[8].nejasan_rok)
        self.assertEqual(stavke[9].pocetak_roka, Kategorija.PocetakRoka.PRESTANAK_VAZENJA)

    def test_sporne_stavke_po_vaznosti(self):
        sporne = lista.sporne_stavke(lista.procitaj(self.docx))
        po_vrsti = {s.vrsta: s for s in reversed(sporne)}
        razlicit = [s.redni_brojevi for s in sporne if s.vrsta == "isti_naziv_razlicit_rok"]
        self.assertEqual(sporne[0].vrsta, "isti_naziv_razlicit_rok")  # najvažnije prvo
        self.assertEqual(razlicit, [[9, 10], [12, 13]])  # i između grupa i u istoj grupi
        self.assertEqual(po_vrsti["protivrecan_rok"].redni_brojevi, [5, 6])  # izlazni računi 10 i 5 godina
        self.assertEqual(po_vrsti["duplikat"].redni_brojevi, [2, 3, 11])  # tri ista naziva su jedan red
        self.assertEqual(po_vrsti["neobicna_oznaka"].redni_brojevi, [10])
        self.assertEqual(po_vrsti["nejasan_rok"].redni_brojevi, [8])
        self.assertEqual(po_vrsti["trajno_operativno"].redni_brojevi, [4])
        self.assertEqual(po_vrsti["pravopis_roka"].redni_brojevi, [6])

    def test_uvoz_i_izvestaj(self):
        stavke = lista.procitaj(self.docx)
        verzija = lista.uvezi(stavke, naziv="Lista 2023", datum_donosenja=datetime.date(2023, 3, 10), aktivna=True)
        self.assertEqual((verzija.kategorije.count(), verzija.grupe.count()), (13, 3))
        for rb in (3, 11):
            dup = Kategorija.objects.get(verzija=verzija, redni_broj=rb)
            self.assertEqual((dup.duplikat_od.redni_broj, dup.aktivna), (2, False))  # duplikat se ne nudi
        razlicit = Kategorija.objects.get(verzija=verzija, redni_broj=13)
        self.assertEqual((razlicit.duplikat_od, razlicit.aktivna), (None, True))  # različit rok: odlučuje arhivista
        self.assertEqual(Kategorija.objects.get(verzija=verzija, redni_broj=10).klasifikaciona_oznaka, "426")
        lista.uvezi(stavke, naziv="Lista 2024", aktivna=True)
        self.assertEqual(list(VerzijaListe.objects.filter(aktivna=True).values_list("naziv", flat=True)), ["Lista 2024"])
        izvestaj = self.dir / "sporne.xlsx"
        lista.izvestaj_xlsx(stavke, lista.sporne_stavke(stavke), izvestaj)
        from openpyxl import load_workbook
        knjiga = load_workbook(izvestaj)
        self.assertEqual(knjiga.sheetnames, ["Sporne stavke", "Sve kategorije"])
        self.assertEqual(knjiga["Sporne stavke"]["A3"].value, "Isti naziv, različit rok")
        self.assertEqual(knjiga["Sve kategorije"].max_row, 14)

    def test_komanda_samo_izvestaj_ne_upisuje(self):
        call_command("import_lista_kategorija", str(self.docx), "--samo-izvestaj", "--izvestaj", str(self.dir / "r.xlsx"),
                     stdout=open(self.dir / "izlaz.txt", "w", encoding="utf-8"))
        self.assertFalse(VerzijaListe.objects.exists())
        self.assertTrue((self.dir / "r.xlsx").exists())


class EkraniTests(RegistarMixin, TestCase):
    KODOVI = ("arhiva:delovodnik", "arhiva:pisarnica", "arhiva:predmet_detail", "arhiva:akt_add",
              "arhiva:predmet_storniraj", "arhiva:kategorije")

    def setUp(self):
        self.korisnik.roles.add(uloga("arhiva-pisarnica", *self.KODOVI))
        self.client.force_login(self.korisnik)

    def upis(self, **polja):
        podaci = {"smer": "ulazni", "datum_zavodjenja": "2026-10-01", "naslov": "Zahtev za ispitivanje",
                  "korespondent": "Putevi Srbije", "mesto": "Beograd", "glavna_oj": str(self.oj431.pk)}
        podaci.update(polja)
        return self.client.post(reverse("arhiva:pisarnica"), podaci, follow=True)

    def test_pisarnica_zavodi_i_pamti_oj(self):
        self.assertEqual(self.client.get(reverse("arhiva:pisarnica")).status_code, 200)
        odgovor = self.upis()
        self.assertContains(odgovor, "Zavedeno pod brojem 43-1")
        self.assertEqual(odgovor.context["forma"].initial["glavna_oj"], str(self.oj431.pk))  # sledeći upis, ista OJ
        predmet = Predmet.objects.get()
        self.assertEqual((predmet.korespondent, predmet.zaveo), ("Putevi Srbije", self.korisnik))

    def test_ulazna_posta_trazi_posiljaoca(self):
        odgovor = self.upis(korespondent="")
        self.assertEqual(odgovor.status_code, 400)
        self.assertFalse(Predmet.objects.exists())

    def test_delovodnik_detalj_podbroj_storno(self):
        self.upis()
        predmet = Predmet.objects.get()
        spisak = self.client.get(reverse("arhiva:delovodnik"), {"godina": 2026})
        self.assertContains(spisak, "43-1")
        self.assertContains(spisak, "431 · Laboratorija za geotehniku")
        self.assertContains(self.client.get(reverse("arhiva:predmet_detail", args=[predmet.pk])), "Novi akt u predmetu")
        self.client.post(reverse("arhiva:akt_add", args=[predmet.pk]), {"datum": "2026-10-02", "opis": "Odgovor", "smer": "izlazni"})
        self.assertEqual(list(predmet.akti.values_list("delovodni_broj", flat=True)), ["43-1", "43-1/2"])
        self.client.post(reverse("arhiva:predmet_storniraj", args=[predmet.pk]), {"razlog": "Pogrešna OJ"})
        predmet.refresh_from_db()
        self.assertTrue(predmet.storniran)
        self.assertContains(self.client.get(reverse("arhiva:predmet_detail", args=[predmet.pk])), "Pogrešna OJ")

    def test_kategorije(self):
        verzija = VerzijaListe.objects.create(naziv="Lista 2023", aktivna=True)
        grupa = verzija.grupe.create(redosled=1, klasifikaciona_oznaka="82", naziv="Pravni poslovi")
        Kategorija.objects.create(verzija=verzija, grupa=grupa, redni_broj=26, naziv="Izveštaji o radu", rok_tekst="Trajno",
                                  trajno=True)
        self.assertContains(self.client.get(reverse("arhiva:kategorije")), "Izveštaji o radu")
        self.assertContains(self.client.get(reverse("arhiva:pisarnica")), "26. Izveštaji o radu · Trajno")

    def test_bez_dozvole(self):
        drugi = get_user_model().objects.create_user("bez-arhive", password="x")
        self.client.force_login(drugi)
        self.assertEqual(self.client.get(reverse("arhiva:delovodnik")).status_code, 403)
        self.assertEqual(self.upis().status_code, 403)

    def test_meni(self):
        odgovor = self.client.get(reverse("arhiva:delovodnik"))
        self.assertContains(odgovor, reverse("arhiva:pisarnica"))
        self.assertContains(odgovor, "switch-app/arhiva/")


@override_settings(PRAVA_PO_REGISTRU={"arhiva": True})
class ObuhvatTests(RegistarMixin, TestCase):
    def setUp(self):
        self.p431 = self.zavedi(glavna_oj=self.oj431)
        self.p41 = self.zavedi(glavna_oj=self.c41)
        self.uloga = uloga("arhiva-referent", "arhiva:delovodnik", "arhiva:predmet_detail")
        self.radnik = get_user_model().objects.create_user("referent-43", password="x")
        self.radnik.roles.add(self.uloga)

    def vidi(self):
        return set(predmeti(Predmet.objects.all(), get_user_model().objects.get(pk=self.radnik.pk)))

    def dodeli(self, **obuhvat):
        DodelaUloge.objects.create(korisnik=self.radnik, uloga=self.uloga, vazi_od=datetime.date(2026, 1, 1),
                                   status=DodelaUloge.STATUS_AKTIVNA, **obuhvat)

    def test_bez_dodele_samo_svoje(self):
        self.assertEqual(self.vidi(), set())
        Predmet.objects.filter(pk=self.p41.pk).update(referent=self.radnik)
        self.assertEqual(self.vidi(), {self.p41})  # zaduženi referent uvek vidi svoj predmet

    def test_centar_vidi_svoje_oj_i_dodatne(self):
        self.dodeli(cvor_id=self.c43.pk)
        self.assertEqual(self.vidi(), {self.p431})  # OJ 431 pripada centru 43
        self.p41.dodatne_oj.add(self.oj431)
        self.assertEqual(self.vidi(), {self.p431, self.p41})
        self.client.force_login(self.radnik)
        self.assertEqual(self.client.get(reverse("arhiva:predmet_detail", args=[self.p41.pk])).status_code, 200)

    def test_tudji_predmet_404(self):
        self.dodeli(cvor_id=self.c43.pk)
        self.client.force_login(self.radnik)
        self.assertEqual(self.client.get(reverse("arhiva:predmet_detail", args=[self.p41.pk])).status_code, 404)

    def test_cela_firma(self):
        self.dodeli(cela_firma=True)
        self.assertEqual(self.vidi(), {self.p431, self.p41})
