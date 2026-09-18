from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import hashlib
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from django import forms
from django.contrib import messages
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.shortcuts import render, redirect
from django.utils import timezone
from django.views.decorators.http import require_GET, require_http_methods
from core.mixins import user_has_role_permission as has
from core.exporting import workbook_response
from .access import check_access
from .forms import AmountField
from .models import FinancePartnerIdentity, CollectionNotice, CollectionState, CollectionFileImport
from .services.source import payload
from .views_operations import writable, page_context, audit


@require_GET
def export_table(request):
    from .views import TABLES, LABELS, rows_for, current_snapshot
    check_access(request.user)
    if not has(request.user, 'potrazivanja:export'): raise PermissionDenied
    kind=request.GET.get('kind',request.GET.get('view','partners'))
    if kind not in TABLES: raise PermissionDenied
    snapshot=current_snapshot(request)
    rows=rows_for(request,snapshot,kind)
    query=request.GET.get('search[value]',request.GET.get('q','')).casefold()
    if query: rows=[r for r in rows if any(query in str(c['display']).casefold() for c in r)]
    columns=[(i,label) for i,label in enumerate(TABLES[kind]) if label!='Radnje']
    workbook=Workbook();sheet=workbook.active;sheet.title='Potrazivanja'
    sheet.append([LABELS[kind]])
    sheet.append([f'Stanje na dan {snapshot.as_of_date:%d.%m.%Y.}' if snapshot else 'Operativna evidencija'])
    sheet.append([f'Centar: {request.GET.get("center") or "svi"}; posao: {request.GET.get("job") or "svi"}'])
    sheet.append([label for _,label in columns])
    for row in rows:
        values=[]
        for i,_ in columns:
            item=row[i]
            value=Decimal(item['sort']) if item['kind']=='money' and item['sort']!='' else date.fromisoformat(item['sort']) if item['kind']=='date' and item['sort'] else item['display']
            values.append(value)
        sheet.append(values)
    for row in sheet:
        for cell in row:
            if isinstance(cell.value,str):cell.data_type='s' # Imported text must never become an Excel formula.
            if isinstance(cell.value,Decimal):cell.number_format='#,##0.00;[Red]-#,##0.00'
            if isinstance(cell.value,date):cell.number_format='dd.mm.yyyy'
    for cell in sheet[4]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='175E7C')
    for i in range(1,len(columns)+1):sheet.column_dimensions[sheet.cell(4,i).column_letter].width=25 if i>1 else 45
    sheet.freeze_panes='A5';sheet.auto_filter.ref=f'A4:{sheet.cell(sheet.max_row,len(columns)).coordinate}'
    return workbook_response(workbook,f'potrazivanja-{kind}.xlsx')


class ImportForm(forms.Form):
    file=forms.FileField(label='Excel dokument (.xlsx)')


def parse_notices(upload):
    if not upload.name.lower().endswith('.xlsx') or upload.size>5*1024*1024:
        raise ValidationError('Izaberite .xlsx dokument do 5 MB.')
    content=upload.read()
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist())>30*1024*1024:raise ValidationError('Dokument je prevelik nakon raspakivanja.')
        book=load_workbook(BytesIO(content),read_only=True,data_only=False)
    except (BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ValidationError('Excel dokument nije ispravan.') from exc
    aliases={'sifra':'sif_par','šifra':'sif_par','sifra partnera':'sif_par','šifra partnera':'sif_par','naziv':'naz_par','naziv partnera':'naz_par','godina':'god','broj':'br_opomene','broj opomene':'br_opomene','napomena':'napomene'}
    try:
        rows=book.active.iter_rows(values_only=True)
        header=next(rows,())
        names=[aliases.get(str(v or '').strip().lower(),str(v or '').strip().lower()) for v in header]
        if not {'sif_par','datum','iznos'}.issubset(names):raise ValidationError('Obavezne kolone: sif_par, datum, iznos. Ostale: naz_par, god, br_opomene, fakture, napomene.')
        if len([v for v in names if v]) != len(set(v for v in names if v)):raise ValidationError('Zaglavlje sadrži duplirane kolone.')
        result=[];errors=[];partners={}
        for number,values in enumerate(rows,2):
            if number>5001:raise ValidationError('Najviše 5.000 redova po uvozu.')
            if not any(v is not None for v in values):continue
            raw=dict(zip(names,values))
            try:
                code=Decimal(str(raw.get('sif_par')))
                if not code.is_finite() or code!=int(code):raise ValidationError('Šifra partnera mora biti ceo broj.')
                code=int(code)
                if code not in partners:partners[code]=FinancePartnerIdentity.objects.filter(company=1,partner_group=1,partner_code=code).first()
                partner=partners[code]
                if not partner:raise ValidationError('Šifra partnera nije pronađena u lokalnom šifarniku.')
                issued=raw.get('datum')
                if isinstance(issued,datetime):issued=issued.date()
                elif not isinstance(issued,date):issued=forms.DateField(input_formats=['%Y-%m-%d','%d.%m.%Y','%d.%m.%Y.']).clean(str(issued or ''))
                amount=AmountField(max_digits=18,decimal_places=2).clean(raw.get('iznos'))
                if raw.get('god') is not None and int(raw['god'])!=issued.year:raise ValidationError('Godina se ne slaže sa datumom.')
                values={'identity_id':partner.pk,'partner_name_snapshot':partner.source_name,'kind':'reminder','year':issued.year,
                        'issued_on':issued,'total_amount':amount,'currency':'RSD','number':str(raw.get('br_opomene') or ''),
                        'original_invoice_text':str(raw.get('fakture') or ''),'note':str(raw.get('napomene') or '')}
                obj=CollectionNotice(**values);obj.full_clean()
                result.append(payload(values))
            except (ValidationError,ValueError,TypeError,ArithmeticError) as exc:
                errors.append(f'Red {number}: {"; ".join(exc.messages) if isinstance(exc,ValidationError) else "Neispravna vrednost."}')
        if errors:raise ValidationError(errors)
        if not result:raise ValidationError('Dokument nema redove za uvoz.')
        return {'fingerprint':hashlib.sha256(content).hexdigest(),'filename':upload.name[:255],'rows':result,'prepared':timezone.now().isoformat()}
    finally:book.close()


@require_http_methods(['GET','POST'])
def import_notices(request):
    writable(request,'notice','create')
    if not has(request.user,'potrazivanja:notice_import'):raise PermissionDenied
    pending=request.session.get('potrazivanja_notice_import')
    form=ImportForm(request.POST or None,request.FILES or None)
    if request.method=='POST' and request.POST.get('confirm'):
        if not pending or pending['fingerprint']!=request.POST.get('confirm'):
            messages.error(request,'Pregled uvoza je istekao. Izaberite dokument ponovo.')
        elif (timezone.now()-datetime.fromisoformat(pending['prepared'])).total_seconds()>3600:
            request.session.pop('potrazivanja_notice_import',None);pending=None
            messages.error(request,'Pregled je stariji od sat vremena. Učitajte dokument ponovo.')
        else:
            with transaction.atomic():
                CollectionState.objects.select_for_update().get(company=1)
                if CollectionFileImport.objects.filter(fingerprint=pending['fingerprint']).exists():
                    messages.info(request,'Ovaj dokument je već uvezen; zapisi nisu duplirani.')
                else:
                    ids=[]
                    for data in pending['rows']:
                        obj=CollectionNotice(**data,created_by=request.user,updated_by=request.user)
                        obj.full_clean();obj.save();audit(request.user,obj,'xlsx_import',{});ids.append(obj.pk)
                    CollectionFileImport.objects.create(fingerprint=pending['fingerprint'],filename=pending['filename'],row_count=len(ids),notice_ids=ids,created_by=request.user)
                    messages.success(request,f'Uvezeno je {len(ids)} opomena u Potraživanja.')
            request.session.pop('potrazivanja_notice_import',None)
            return redirect('/potrazivanja/?view=notices&type=reminder')
    elif request.method=='POST':
        request.session.pop('potrazivanja_notice_import',None);pending=None
        if form.is_valid():
            try:
                pending=parse_notices(form.cleaned_data['file'])
                request.session['potrazivanja_notice_import']=pending
            except ValidationError as exc:form.add_error('file',exc)
    return render(request,'potrazivanja/notice_import.html',page_context(request,form=form,pending=pending,title='Uvoz opomena',section='notices'))
