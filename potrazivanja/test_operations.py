from copy import deepcopy
from unittest.mock import patch
from decimal import Decimal
from django.test import TestCase
from django.contrib.auth import get_user_model
from django.urls import reverse
from django.utils import timezone
from .models import (CollectionState, FinancePartnerIdentity, CollectionContact, ContactPoint, CollectionActivity,
                     CollectionNotice, CollectionLegalCase, CollectionLegalEvent, CollectionProfile, CollectionAudit)
from .services.sync import sync_collections
from .services.contacts import normalize_contacts, parse_contact
from .services.independence import initialize_independent_operations
from .test_sync import fixture, OBSERVED


class IndependentOperationsTests(TestCase):
    def setUp(self):
        self.data=fixture()
        self.data['postupak']=[{'id':15,'tip':'tuzeni','sifra_partnera':42,'naziv_partnera':'Kupac','broj_predmeta':'P 1','osnovni_dug':'50.25','valuta':'RSD'}]
        self.data['promena_postupka']=[{'id':31,'postupak_id':15,'datum':'2026-09-01','promena':'Rociste'}]
        with patch('potrazivanja.services.sync.extract',return_value=(deepcopy(self.data),OBSERVED)):
            self.result=initialize_independent_operations(progress=None)
        self.partner=FinancePartnerIdentity.objects.get(partner_code=42)
        self.admin=get_user_model().objects.create_user(username='ops-admin',is_superuser=True)
        self.client.force_login(self.admin)

    def url(self,kind,pk=None):
        return reverse('potrazivanja:record_update',args=[kind,pk]) if pk else reverse('potrazivanja:record_create',args=[kind])

    def contact_data(self):
        return {'identity':self.partner.pk,'first_name':'Ana','last_name':'Petrovic','position':'Finansije','note':'Nova',
                'items-TOTAL_FORMS':'2','items-INITIAL_FORMS':'0','items-MAX_NUM_FORMS':'50',
                'items-0-kind':'mobile','items-0-value':'064/123-4567','items-0-label':'Sluzbeni',
                'items-1-kind':'email','items-1-value':'ana@example.test','items-1-label':'Poslovni'}

    def test_initialization_is_once_and_legal_records_are_local(self):
        self.assertTrue(CollectionState.objects.get().operations_independent_at)
        case=CollectionLegalCase.objects.get(legacy_id=15)
        self.assertEqual(case.osnovni_dug,Decimal('50.25'))
        self.assertEqual(case.events.get().legacy_id,31)
        with patch('potrazivanja.services.sync.extract') as extract:
            self.assertEqual(initialize_independent_operations(),{'already_initialized':True})
        extract.assert_not_called()
        with self.assertRaises(ValueError):sync_collections(include_legacy=True)

    def test_financial_sync_does_not_reimport_or_overwrite_operations(self):
        contact=CollectionContact.objects.get(import_parent=None);contact.note='Lokalno';contact.save()
        self.data['kontakti'][0]['napomena']='Promenjeno u staroj'
        with patch('potrazivanja.services.sync.extract',return_value=(deepcopy(self.data),OBSERVED)) as extract:
            run=sync_collections()
        self.assertFalse(extract.call_args.kwargs['include_legacy'])
        self.assertNotIn('kontakti',run.source_counts)
        contact.refresh_from_db();self.assertEqual(contact.note,'Lokalno')
        self.assertEqual(CollectionLegalCase.objects.count(),1)

    def test_structured_contact_and_audit_and_archive(self):
        response=self.client.post(self.url('contact'),self.contact_data())
        self.assertEqual(response.status_code,302)
        contact=CollectionContact.objects.get(first_name='Ana')
        self.assertEqual(contact.phone,'0641234567');self.assertEqual(contact.points.count(),2)
        self.assertTrue(CollectionAudit.objects.filter(entity_id=contact.pk,action='create').exists())
        url=reverse('potrazivanja:record_archive',args=['contact',contact.pk])
        self.assertEqual(self.client.post(url,{'version':'stale','archived':'1'}).status_code,409)
        self.assertEqual(self.client.post(url,{'version':contact.updated_at.isoformat(),'archived':'1'}).status_code,302)
        contact.refresh_from_db();self.assertFalse(contact.active);self.assertEqual(contact.points.count(),2)

    def test_invalid_combined_email_is_rejected(self):
        data=self.contact_data();data['items-1-value']='ana@example.test b@example.test'
        response=self.client.post(self.url('contact'),data)
        self.assertEqual(response.status_code,200)
        self.assertFalse(CollectionContact.objects.filter(first_name='Ana').exists())

    def test_call_with_followup_and_cross_partner_contact_validation(self):
        contact=CollectionContact.objects.get(import_parent=None)
        data={'identity':self.partner.pk,'kind':'phone','contact':contact.pk,'occurred_at':'2026-09-18T11:30','text':'Poziv','outcome':'Dogovor','next_action_date':'2026-09-21'}
        self.assertEqual(self.client.post(self.url('activity'),data).status_code,302)
        self.assertTrue(CollectionActivity.objects.filter(outcome='Dogovor').exists())
        contact.identity=FinancePartnerIdentity.objects.get(partner_code=43);contact.save()
        self.assertEqual(self.client.post(self.url('activity'),data).status_code,200)
        self.assertEqual(CollectionActivity.objects.filter(outcome='Dogovor').count(),1)

    def notice_data(self):
        return {'identity':self.partner.pk,'kind':'letter','year':'2026','number':'P-1','issued_on':'2026-09-18','currency':'RSD',
                'total_amount':'1.200,50','body':'Molimo odgovor','items-TOTAL_FORMS':'1','items-INITIAL_FORMS':'0','items-MAX_NUM_FORMS':'500',
                'items-0-original_reference':'IF-15','items-0-job_code':'436111','items-0-due_date_snapshot':'2026-09-01','items-0-amount_snapshot':'1.200,50'}

    def test_notice_totals_and_print(self):
        data=self.notice_data();data['total_amount']='100'
        self.assertEqual(self.client.post(self.url('notice'),data).status_code,200)
        self.assertFalse(CollectionNotice.objects.filter(number='P-1').exists())
        data=self.notice_data();self.assertEqual(self.client.post(self.url('notice'),data).status_code,302)
        notice=CollectionNotice.objects.get(number='P-1')
        self.assertEqual(notice.total_amount,Decimal('1200.50'));self.assertEqual(notice.items.get().line_number,1)
        response=self.client.get(reverse('potrazivanja:notice_print',args=[notice.pk]))
        self.assertContains(response,'IF-15');self.assertContains(response,'Molimo odgovor')

    def test_review_and_edit_conflicts(self):
        profile=CollectionProfile.objects.get(identity=self.partner)
        url=reverse('potrazivanja:review_toggle',args=[self.partner.pk])
        self.assertEqual(self.client.post(url,{'value':'0','version':profile.updated_at.isoformat()}).status_code,200)
        self.assertEqual(self.client.post(url,{'value':'1','version':profile.updated_at.isoformat()}).status_code,409)
        profile.refresh_from_db();self.assertFalse(profile.needs_review)
        data=self.contact_data();contact=CollectionContact.objects.get(import_parent=None)
        data.update(version='stale',**{'items-TOTAL_FORMS':'0'})
        self.assertEqual(self.client.post(self.url('contact',contact.pk),data).status_code,200)
        contact.refresh_from_db();self.assertNotEqual(contact.first_name,'Ana')

    def test_readonly_user_cannot_write(self):
        user=get_user_model().objects.create_user(username='no-writes');self.client.force_login(user)
        for kind in ('contact','activity','notice','legal'):
            self.assertEqual(self.client.post(self.url(kind),{}).status_code,403)

    def test_all_forms_and_legal_history_render(self):
        for kind in ('contact','activity','notice','legal'):
            self.assertEqual(self.client.get(self.url(kind),{'partner':self.partner.pk}).status_code,200)
        case=CollectionLegalCase.objects.get()
        self.assertEqual(self.client.get(reverse('potrazivanja:legal_detail',args=[case.pk])).status_code,200)
        self.assertEqual(self.client.post(reverse('potrazivanja:legal_event_create',args=[case.pk]),{'date':'2026-09-18','text':'Nova promena'}).status_code,302)
        self.assertEqual(case.events.count(),2)

    def spreadsheet(self, rows):
        from io import BytesIO
        from openpyxl import Workbook
        from django.core.files.uploadedfile import SimpleUploadedFile
        book=Workbook();sheet=book.active
        sheet.append(['sif_par','datum','iznos','br_opomene','fakture'])
        for row in rows:sheet.append(row)
        output=BytesIO();book.save(output)
        return SimpleUploadedFile('opomene.xlsx',output.getvalue())

    def test_excel_import_preview_atomic_rejection_and_repeat(self):
        from .models import CollectionFileImport
        url=reverse('potrazivanja:notice_import')
        count=CollectionNotice.objects.count()
        response=self.client.post(url,{'file':self.spreadsheet([[42,'2026-09-18','100.25','X-1','IF-1'],[99999,'2026-09-18','50','X-2','IF-2']])})
        self.assertEqual(response.status_code,200);self.assertEqual(CollectionNotice.objects.count(),count)
        self.assertNotIn('potrazivanja_notice_import',self.client.session)
        response=self.client.post(url,{'file':self.spreadsheet([[42,'2026-09-18','100.25','X-1','IF-1']])})
        self.assertEqual(response.status_code,200);self.assertEqual(CollectionNotice.objects.count(),count)
        pending=dict(self.client.session['potrazivanja_notice_import'])
        self.assertEqual(self.client.post(url,{'confirm':pending['fingerprint']}).status_code,302)
        self.assertEqual(CollectionNotice.objects.count(),count+1)
        session=self.client.session;session['potrazivanja_notice_import']=pending;session.save()
        self.client.post(url,{'confirm':pending['fingerprint']})
        self.assertEqual(CollectionNotice.objects.count(),count+1);self.assertEqual(CollectionFileImport.objects.count(),1)

    def test_excel_export_preserves_numbers_and_treats_names_as_text(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.partner.source_name='=1+2';self.partner.save()
        response=self.client.get(reverse('potrazivanja:export'))
        self.assertEqual(response.status_code,200)
        book=load_workbook(BytesIO(response.content));sheet=book.active
        self.assertEqual(sheet['A5'].value,'=1+2');self.assertEqual(sheet['A5'].data_type,'s')
        self.assertEqual(sheet['L5'].value,100.01)

    def test_review_list_includes_marked_partners_without_open_balances(self):
        partner=FinancePartnerIdentity.objects.get(partner_code=43)
        CollectionProfile.objects.create(identity=partner,needs_review=True)
        response=self.client.get(reverse('potrazivanja:table_data'),{'kind':'review'})
        self.assertEqual(response.json()['recordsTotal'],2)

    def test_separate_tabs_never_mix_activity_or_notice_types(self):
        CollectionNotice.objects.create(identity=self.partner,kind='letter',number='P-2')
        CollectionNotice.objects.create(identity=self.partner,kind='legacy_claim',number='T-2')
        from .views import SPLIT_OPERATIONS
        for kind,(base,record_type,title) in SPLIT_OPERATIONS.items():
            model=CollectionActivity if base=='activities' else CollectionNotice
            with self.subTest(kind=kind):
                # A conflicting query-string filter must not change a dedicated tab.
                response=self.client.get(reverse('potrazivanja:table_data'),{'kind':kind,'partner':self.partner.pk,'type':'wrong'})
                self.assertEqual(response.status_code,200)
                data=response.json()
                expected=model.objects.filter(identity=self.partner,kind=record_type)
                self.assertEqual(data['recordsTotal'],expected.count())
                self.assertTrue(data['data'])
                self.assertEqual({row[1]['display'] for row in data['data']},{expected.first().get_kind_display()})


class ContactCleanupTests(TestCase):
    def test_unparsed_email_text_is_flagged(self):
        parsed=parse_contact('011/7358533','neispravna adresa')
        self.assertEqual(parsed['unresolved'],['neispravna adresa'])

    def test_separate_people_and_multiple_emails_without_guessing_association(self):
        parsed=parse_contact('011/7358533, 065/3601247 Tamara, 063/292769 Marija','a@example.test b@example.test')
        self.assertEqual([p['first_name'] for p in parsed['people']],['Tamara','Marija'])
        self.assertEqual(parsed['phones'],['0117358533']);self.assertEqual(len(parsed['emails']),2)

    def test_ambiguous_extensions_preserved_and_normalization_idempotent(self):
        raw='011/414-6541/6540, 063/685-499 Zoran'
        contact=CollectionContact.objects.create(name=raw,email='a@example.test')
        stats=normalize_contacts();contact.refresh_from_db()
        self.assertEqual(stats['people_created'],1);self.assertEqual(contact.original_data['name'],raw)
        self.assertTrue(contact.needs_review);self.assertEqual(contact.points.count(),1)
        self.assertEqual(normalize_contacts()['source_contacts'],0)
