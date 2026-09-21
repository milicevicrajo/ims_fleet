from datetime import date
from decimal import Decimal as D
from io import BytesIO
from unittest.mock import MagicMock, patch

from django.contrib.auth import get_user_model
from django.db import DatabaseError
from django.test import TestCase, SimpleTestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import PermissionCode, Role
from .models import FinanceJob, SyncRun
from .services.cash_flow import cash_flow_many
from .services.job_overview import enrich_jobs
from .tests import job, save_entry


def shared_result(company, codes, start, end):
    return {code: {"available": True, "cost": D("10"), "complete": True, "note": "test"} for code in codes}


def cash_result(company, codes, start, end):
    return {code: {"available": True, "inflow": D("120"), "outflow": D("150"), "net": D("-30"),
                   "complete": True, "note": "test"} for code in codes}


class JobOverviewTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser("overview-admin", "test@example.invalid", "test")
        self.client.force_login(self.user)
        FinanceJob.objects.create(**job())
        FinanceJob.objects.create(**job("420001", "42", name="Other job"))
        FinanceJob.objects.create(**job("410002", "41", name="Empty job"))
        save_entry()
        save_entry(number=2, account="51200", debit=D("40"), credit=D("0"))
        save_entry(number=3, job_code="420001", center="42", credit=D("999"))
        SyncRun.objects.create(company=1, year_from=2025, year_to=2026, status="success")
        self.params = {"group": "job", "date_from": "2026-02-01", "date_to": "2026-02-28", "center": "41"}
        self.shared = patch("finansije.services.job_overview.shared_cost_many", side_effect=shared_result).start()
        self.cash = patch("finansije.services.job_overview.cash_flow_many", side_effect=cash_result).start()
        self.addCleanup(patch.stopall)

    def test_page_has_only_period_center_no_hero_totals_and_100_rows_without_source_reads(self):
        response = self.client.get(reverse("finansije:report"), self.params)
        self.assertContains(response, 'data-page-length="100"')
        self.assertNotContains(response, 'class="finance-hero-metrics"')
        self.assertNotContains(response, 'class="finance-hero-period"')
        for name in ("job", "kind", "unit", "account", "include_empty", "account_exact"):
            self.assertNotContains(response, f'name="{name}"')
        self.assertContains(response, "od 01.02.2026. do 28.02.2026.")
        self.assertEqual(len(response.context["rows"]), 2)
        self.shared.assert_not_called()
        self.cash.assert_not_called()

    def test_ajax_has_eight_signed_metrics_and_does_not_apply_removed_filters(self):
        response = self.client.get(reverse("finansije:jobs_data"), dict(self.params, kind="expense", account="59900", job="420001"))
        rows = response.json()["data"]
        self.assertEqual({row[0]["sort"] for row in rows}, {"410001", "410002"})
        row = next(row for row in rows if row[0]["sort"] == "410001")
        self.assertEqual([D(c["sort"]) for c in row[3:11]], list(map(D, (100, -40, 60, -10, 50, 120, -150, -30))))
        self.assertEqual(len(row), 11)
        self.assertIn("finance-job-detail-button", row[0]["display"])
        self.assertIn("/finansije/posao/?job=410001", row[0]["display"])
        self.shared.assert_called_once_with(1, {"410001", "410002"}, date(2026, 2, 1), date(2026, 2, 28))
        self.assertIn("no-store", response.headers["Cache-Control"])

    def test_endpoint_rechecks_scope_and_rejects_invalid_filters_before_sources(self):
        user = get_user_model().objects.create_user("limited-overview", allowed_center_codes="41")
        role = Role.objects.create(name="Finance overview", slug="finance-overview")
        role.permissions.add(PermissionCode.objects.create(code="finansije:dashboard"))
        user.roles.add(role)
        self.client.force_login(user)
        response = self.client.get(reverse("finansije:jobs_data"), dict(self.params, center="42"))
        self.assertEqual(response.status_code, 400)
        self.shared.assert_not_called()
        self.assertEqual(self.client.get(reverse("finansije:jobs_data"), dict(self.params, date_from="bad")).status_code, 400)
        rows = self.client.get(reverse("finansije:jobs_data"), dict(self.params, center="")).json()["data"]
        self.assertEqual({row[0]["sort"] for row in rows}, {"410001", "410002"})
        user.roles.clear()
        self.assertEqual(self.client.get(reverse("finansije:jobs_data"), self.params).status_code, 403)

    def test_missing_source_preserves_known_pnl_and_shared_and_does_not_leak_error(self):
        self.cash.side_effect = DatabaseError("private database detail")
        with self.assertLogs("finansije.services.job_overview", level="ERROR"):
            response = self.client.get(reverse("finansije:jobs_data"), self.params)
        row = response.json()["data"][0]
        self.assertEqual(D(row[3]["sort"]), D("100"))
        self.assertEqual(D(row[7]["sort"]), D("50"))
        self.assertEqual(row[8]["sort"], "")
        self.assertNotIn("private database detail", response.content.decode())

    def test_year_export_keeps_signed_values_and_year_detail_link(self):
        response = self.client.get(reverse("finansije:export"), dict(self.params, date_from="2026-01-01", date_to="2026-12-31"))
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content), read_only=True).active
        records = list(sheet.values)
        self.assertEqual(records[2][3:11], ("Prihodi", "Rashodi", "Rezultat bez ZT", "Zajednički troškovi", "Rezultat P − R − ZT", "Priliv", "Odliv", "Neto gotovina"))
        self.assertEqual(records[3][3:11], (100, -40, 60, -10, 50, 120, -150, -30))
        self.assertTrue(records[3][-1].endswith("year=2026&month="))

    def test_excel_contains_native_table_filters_numeric_formats_and_frozen_identifiers(self):
        FinanceJob.objects.filter(code='410001').update(name='=SUM(A1:A2)')
        response = self.client.get(reverse('finansije:export'), self.params)
        sheet = load_workbook(BytesIO(response.content)).active
        table = sheet.tables['SifrePosla']
        self.assertEqual(table.ref, f'A3:M{sheet.max_row}')
        self.assertEqual(table.autoFilter.ref, table.ref)
        self.assertTrue(table.tableStyleInfo.showRowStripes)
        self.assertEqual([column.name for column in table.tableColumns], [c.value for c in sheet[3]])
        self.assertEqual(sheet.freeze_panes, 'D4')
        self.assertEqual(sheet['D4'].data_type, 'n')
        self.assertIn('#,##0.00', sheet['D4'].number_format)
        self.assertEqual(sheet['A4'].data_type, 's')
        self.assertEqual(sheet['B4'].data_type, 's')
        self.assertFalse(any(c.data_type == 'f' for row in sheet for c in row))
        self.assertEqual(sheet['D3'].fill.fgColor.rgb, '00235B83')

    def test_empty_excel_still_has_filterable_table_without_invented_zero_values(self):
        with patch('finansije.views.grouped_report', return_value=({}, [])):
            response = self.client.get(reverse('finansije:export'), self.params)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.tables['SifrePosla'].ref, 'A3:M4')
        self.assertTrue(all(c.value is None for c in sheet[4]))


class JobOverviewCalculationTests(SimpleTestCase):
    def row(self):
        return {"code": "a", "revenue": D("100"), "expense": D("40"), "result": D("60")}

    def test_multiple_years_sum_separate_allocations_and_cash_results(self):
        with patch("finansije.services.job_overview.shared_cost_many", side_effect=shared_result) as shared, \
             patch("finansije.services.job_overview.cash_flow_many", side_effect=cash_result):
            rows = enrich_jobs([self.row()], date(2025, 12, 10), date(2026, 2, 5), {2025, 2026})
        self.assertEqual(shared.call_args_list[0].args[2:], (date(2025, 12, 10), date(2025, 12, 31)))
        self.assertEqual(shared.call_args_list[1].args[2:], (date(2026, 1, 1), date(2026, 2, 5)))
        metrics = rows[0]["metrics"]
        self.assertEqual(metrics["shared_cost"]["value"], D("-20"))
        self.assertEqual(metrics["after_shared"]["value"], D("40"))
        self.assertEqual(metrics["net_cash"]["value"], D("-60"))

    def test_unsynchronized_period_is_not_zero_and_does_not_read_sources(self):
        with patch("finansije.services.job_overview.shared_cost_many") as shared, \
             patch("finansije.services.job_overview.cash_flow_many") as cash:
            rows = enrich_jobs([self.row()], date(2026, 1, 1), date(2026, 2, 28), set())
        shared.assert_not_called()
        cash.assert_not_called()
        self.assertTrue(all(metric["value"] is None for metric in rows[0]["metrics"].values()))

    def test_batch_cash_keeps_blank_job_mapping_vat_and_requested_jobs_only(self):
        cursor = MagicMock()
        cursor.fetchall.side_effect = [
            [(None, "20400", date(2026, 8, 1), "BANK", D("0"), D("25")),
             ("111111", "20400", date(2026, 8, 1), "BANK", D("0"), D("75")),
             ("secret", "20400", date(2026, 8, 1), "BANK", D("0"), D("999"))],
            [("20400", ""), ("47900", "")], [("BANK", "NT")], [],
            [(None, D("2")), ("111111", D("3"))], [("111111", 8)],
        ]
        cursor.fetchone.return_value = (8, 8)
        connection = MagicMock()
        connection.cursor.return_value.__enter__.return_value = cursor
        with patch("finansije.services.cash_flow.connections", {"server_db": connection}):
            results = cash_flow_many(1, {"111111", "empty"}, date(2026, 8, 1), date(2026, 8, 31))
        self.assertEqual(set(results), {"111111", "empty"})
        self.assertEqual(results["111111"]["inflow"], D("100"))
        self.assertEqual(results["111111"]["outflow"], D("5"))
        self.assertFalse(results["empty"]["available"])
        self.assertEqual(cursor.execute.call_count, 7)
