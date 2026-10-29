from datetime import date
from decimal import Decimal
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import SimpleTestCase, TestCase
from django.urls import reverse

from core.models import PermissionCode, Role
from .models import BankaNbs, BankaPartner, BankePreuzimanje, PovezanoLiceBanke
from .services import banke_nbs
from .services.banks import fetch_banks
from .tests import save_entry

SPISAK = """<table class="responsive-table"><tr><td>Banca Intesa</td><td>
<a href="/WebApp/FinInstBank/LegalEntityFinInstitution/Details?id=621521dc-3644-4151-9b87-92378344c7f9">Детаљи</a></td></tr>
<tr><td>NLB</td><td><a href="/WebApp/FinInstBank/LegalEntityFinInstitution/Details?id=db0743dc-3932-497a-b89a-e80ef13786cf">Детаљи</a></td></tr></table>"""


def detalji(naziv, mb, pib, adresa="Милентија Поповића 7б, Београд", predsednik="Дарко Поповић"):
    return f"""<main><h1>Основни подаци банке</h1>
<input id="Name" type="text" value="{naziv}"/><input id="City" type="text" value="Београд"/>
<input id="Address" type="text" value="{adresa}"/><input id="PostalCode" type="text" value="11000"/>
<input id="NationalCode" type="text" value="{mb}"/><input id="TaxCode" type="text" value="{pib}"/>
<input id="Phone" type="text" value="011/2011-200"/><input id="Fax" type="text" value=""/>
<label for="Url">Веб сајт</label><div class="form-control"><a href="https://x">www.bancaintesa.rs</a></div>
<label for="Email">И-мејл</label><div class="form-control"><a href="mailto:k">kabinet@bancaintesa.rs</a></div>
<input id="SwiftCode" type="text" value="DBDBRSBG"/><input id="WorkLicence" type="text" value="Решење НБЈ O.бр. 274"/>
<input id="WorkLicenceDate" type="text" value="19.9.1991."/><input id="EmployeeNumber" type="text" value="2884"/>
<textarea id="Description">Промена назива.

* Подаци се ажурирају квартално
</textarea>
<div class="myColapsiblePanel"><div class="myColapsiblePanelCaptionContent">Прилози</div>
<table><tbody><tr><td><a href="/WebApp/FileSystem/X?ID=1">BDO DOO BEOGRAD</a></td></tr></tbody></table></div>
<div class="myColapsiblePanel"><div class="myColapsiblePanelCaptionContent">Највећи акционари (преко 5%)</div>
<table><thead><tr><th>Назив</th><th>%</th></tr></thead><tbody><tr><td>Intesa Holding International S.A.</td><td>100,00 %</td></tr></tbody></table></div>
<div class="myColapsiblePanel"><div class="myColapsiblePanelCaptionContent">Чланови управног одбора</div>
<table><thead><tr><th>Назив</th><th>Функција</th></tr></thead><tbody><tr><td>Драгиња Ђурић</td><td>Председник управног одбора</td></tr></tbody></table></div>
<div class="myColapsiblePanel"><div class="myColapsiblePanelCaptionContent">Чланови извршног одбора</div>
<table><thead><tr><th>Назив</th><th>Функција</th></tr></thead><tbody><tr><td>{predsednik}</td><td>Председник извршног одбора</td></tr></tbody></table></div>
</main>"""


class Odgovor:
    def __init__(self, text, status=200):
        self.text, self.status_code = text, status

    def raise_for_status(self):
        if self.status_code >= 400:
            raise RuntimeError(f"HTTP {self.status_code}")


class LaznaSesija:
    def __init__(self, strane):
        self.strane = strane

    def get(self, url, timeout=None):
        if url.endswith("/Index"):
            return Odgovor(self.strane["spisak"])
        return Odgovor(*self.strane[url.rsplit("id=", 1)[1]])


def strane(intesa_adresa="Милентија Поповића 7б, Београд", nlb=True, nlb_status=200):
    rezultat = {"spisak": SPISAK if nlb else SPISAK.split("<tr><td>NLB")[0] + "</table>",
                "621521dc-3644-4151-9b87-92378344c7f9": (detalji("Banca Intesa akcionarsko društvo Beograd", "7759231", "100001159", intesa_adresa),)}
    if nlb:
        rezultat["db0743dc-3932-497a-b89a-e80ef13786cf"] = (detalji("NLB Banka AD Beograd", "7737068", "100001931"), nlb_status)
    return rezultat


class CitanjeTests(SimpleTestCase):
    def test_latinica_including_capital_digraphs(self):
        self.assertEqual(banke_nbs.latinica("Љубиша ЉУБИША Џон ЊЕГОШ Ђорђе"), "Ljubiša LJUBIŠA Džon NJEGOŠ Đorđe")

    def test_details_page_is_read_and_transliterated(self):
        podaci = banke_nbs.procitaj_detalje("x", detalji("Banca Intesa", "7759231", "100001159"))
        self.assertEqual((podaci["grad"], podaci["adresa"], podaci["resenje"]), ("Beograd", "Milentija Popovića 7b, Beograd", "Rešenje NBJ O.br. 274"))
        self.assertEqual((podaci["web"], podaci["email"], podaci["swift"]), ("www.bancaintesa.rs", "kabinet@bancaintesa.rs", "DBDBRSBG"))
        self.assertEqual((podaci["datum_resenja"], podaci["broj_zaposlenih"], podaci["napomena"]), (date(1991, 9, 19), 2884, "Promena naziva."))
        self.assertEqual(podaci["akcionari"], [{"naziv": "Intesa Holding International S.A.", "procenat": "100,00 %"}])
        self.assertEqual(podaci["upravni_odbor"], [{"ime": "Draginja Đurić", "funkcija": "Predsednik upravnog odbora"}])
        self.assertEqual(podaci["prilozi"][0]["url"], "https://webappcenter.nbs.rs/WebApp/FileSystem/X?ID=1")

    def test_changed_page_layout_is_an_error(self):
        with self.assertRaises(ValueError):
            banke_nbs.procitaj_detalje("x", "<main>Nema podataka</main>")
        with self.assertRaises(ValueError):
            banke_nbs.preuzmi_sa_sajta(LaznaSesija({"spisak": "<table></table>"}), pauza=0)


class OsvezavanjeTests(TestCase):
    def setUp(self):
        cache.clear()

    def test_first_run_creates_banks_and_second_records_changes_and_missing_banks(self):
        prvo = banke_nbs.osvezi(http=LaznaSesija(strane()), pauza=0)
        self.assertEqual((prvo.status, prvo.banaka, prvo.novih), ("success", 2, 2))
        drugo = banke_nbs.osvezi(http=LaznaSesija(strane(intesa_adresa="Нова 1, Београд", nlb=False)), pauza=0)
        self.assertEqual((drugo.promenjenih, drugo.nestalih), (1, 1))
        self.assertIn({"banka": "Banca Intesa akcionarsko društvo Beograd", "polje": "Adresa",
                       "staro": "Milentija Popovića 7b, Beograd", "novo": "Nova 1, Beograd"}, drugo.promene)
        nlb = BankaNbs.objects.get(pib="100001931")
        self.assertFalse(nlb.aktivna)
        self.assertIsNotNone(nlb.nestala)

    def test_failed_page_changes_nothing(self):
        banke_nbs.osvezi(http=LaznaSesija(strane()), pauza=0)
        run = banke_nbs.osvezi(http=LaznaSesija(strane(intesa_adresa="Нова 1, Београд", nlb_status=500)), pauza=0)
        self.assertEqual(run.status, "failed")
        self.assertEqual(BankaNbs.objects.get(pib="100001159").adresa, "Milentija Popovića 7b, Beograd")
        self.assertTrue(BankaNbs.objects.get(pib="100001931").aktivna)

    def test_only_one_run_at_a_time(self):
        cache.add(f"{banke_nbs.ZAKLJUCAVANJE}:nbs", "1", 60)
        with self.assertRaises(banke_nbs.PreuzimanjeUToku):
            banke_nbs.osvezi(http=LaznaSesija(strane()), pauza=0)
        self.assertFalse(BankePreuzimanje.objects.exists())

    def test_partner_codes_are_matched_by_pib_then_mb_and_grouped(self):
        banke_nbs.osvezi(http=LaznaSesija(strane()), pauza=0)
        banks = banke_nbs.razvrstaj([dict(code=11, name="Banca Intesa", pib="100001159", mb="7759231"),
                                     dict(code=14, name="Intesa - bolovanja", pib="", mb="07759231"),
                                     dict(code=1, name="Privredna banka", pib="", mb="")])
        self.assertEqual([b["grupa"].naziv if b["grupa"] else None for b in banks], ["Intesa Sanpaolo", "Intesa Sanpaolo", None])
        self.assertEqual([b["van_nbs"] for b in banks], [False, False, True])

    def test_director_list_of_groups_is_loaded(self):
        self.assertEqual(PovezanoLiceBanke.objects.count(), 38)
        self.assertEqual(PovezanoLiceBanke.objects.get(pib="103156991").grupa.naziv, "NLB")

    def test_monthly_schedule(self):
        from core.management.commands.sync_celery_periodic_tasks import EXPECTED_PERIODIC_TASKS

        zadatak = next(s for s in EXPECTED_PERIODIC_TASKS if s["task"] == "finansije.tasks.osvezi_banke_nbs_task")
        self.assertEqual(zadatak["day_of_month"], "1")


BANKS = [dict(code=11, name="Banca Intesa akcionarsko drustvo Beograd", pib="100001159", mb="7759231", fields=[], zr="908-1",
              naz_par="Banca Intesa akcionarsko drustvo Beograd", ulica_par="Stara 1", mesto_par="Beograd"),
         dict(code=1, name="Privredna banka", pib="", mb="", fields=[], zr=""),
         dict(code=10, name="Uprava za trezor", pib="", mb="", fields=[], zr="")]


@patch("finansije.bank_views.service.fetch_banks", return_value=BANKS)
class BankeNbsViewTests(TestCase):
    def setUp(self):
        cache.clear()
        banke_nbs.osvezi(http=LaznaSesija(strane()), pauza=0)
        self.user = get_user_model().objects.create_user("banke-nbs", password="test", allowed_center_codes="41")
        self.role = Role.objects.create(name="Banke NBS test", slug="banke-nbs-test")
        self.user.roles.add(self.role)
        self.client.force_login(self.user)

    def grant(self, *codes):
        for code in codes:
            self.role.permissions.add(PermissionCode.objects.get_or_create(code=code)[0])

    def test_list_groups_banks_and_marks_historical(self, source):
        self.grant("finansije:bank_list", "finansije:bank_detail")
        save_entry(number=1, account="24100", partner_group=11, partner_code=10, debit=Decimal("5"), credit=Decimal("0"))
        response = self.client.get(reverse("finansije:bank_list"), {"date_from": "2026-01-01", "date_to": "2026-06-30"})
        grupe = response.context["bank_groups"]
        self.assertEqual([(g["naziv"], [b["code"] for b in g["banke"]]) for g in grupe], [("Intesa Sanpaolo", [11])])
        van = {b["code"]: b["istorijska"] for b in response.context["outside_nbs"]}
        self.assertEqual(van, {1: True, 10: False})
        self.assertContains(response, "Istorijska")
        self.assertContains(response, "Nije banka sa spiska NBS")
        self.assertNotContains(response, "Osveži podatke banaka")

    def test_nbs_and_group_tabs(self, source):
        self.grant("finansije:bank_detail")
        nbs = self.client.get(reverse("finansije:bank_detail", args=[11]), {"tab": "nbs"})
        self.assertContains(nbs, "Draginja Đurić")
        self.assertContains(nbs, "razlikuje se")  # adresa u šifarniku je „Stara 1”
        with patch("finansije.bank_views.banke_nbs.lica_grupe", return_value=[]) as lica:
            grupa = self.client.get(reverse("finansije:bank_detail", args=[11]), {"tab": "group"})
        self.assertContains(grupa, "Intesa Sanpaolo")
        self.assertEqual(lica.call_args.args[1].naziv, "Intesa Sanpaolo")
        self.assertContains(self.client.get(reverse("finansije:bank_detail", args=[1]), {"tab": "nbs"}), "nije povezana")

    def test_refresh_runs_partner_list_then_nbs_and_requires_whole_company(self, source):
        url = reverse("finansije:bank_refresh")
        with patch("finansije.bank_views.banke_nbs.osvezi") as osvezi, patch("finansije.bank_views.banke_nbs.sinhronizuj_sifarnik") as sifarnik:
            self.grant("finansije:bank_refresh")
            self.assertEqual(self.client.post(url).status_code, 403)
            self.grant("finansije:view_all", "finansije:bank_list")
            self.assertContains(self.client.get(reverse("finansije:bank_list")), "Osveži podatke banaka")
            osvezi.return_value = BankePreuzimanje(izvor="nbs", status="success", banaka=19)
            sifarnik.return_value = BankePreuzimanje(izvor="sifarnik", status="success", banaka=23)
            self.assertRedirects(self.client.post(url, {"next": "https://evil.example/"}), reverse("finansije:bank_list"),
                                 fetch_redirect_response=False)
        osvezi.assert_called_once()
        sifarnik.assert_called_once()
        self.assertEqual(self.client.get(url).status_code, 405)


def red_sifarnika(sifra, naziv, pib="", mb="", ulica=""):
    from .services.banks import PARTNER_FIELDS

    red = dict.fromkeys(PARTNER_FIELDS, "")
    red.update(sif_par=str(sifra), naz_par=naziv, pib=pib, mb=mb, ulica_par=ulica)
    return red


class SifarnikTests(TestCase):
    def setUp(self):
        cache.clear()

    def sinhronizuj(self, redovi):
        with patch("finansije.services.banks.source_banks", return_value=redovi):
            return banke_nbs.sinhronizuj_sifarnik()

    def test_local_copy_records_changes_and_deleted_codes_and_keeps_our_labels(self):
        prvo = self.sinhronizuj([red_sifarnika(16, "Aikbank", "100618836", ulica="Bulevra 59a"), red_sifarnika(1, "Privredna banka")])
        self.assertEqual((prvo.status, prvo.izvor, prvo.novih), ("success", "sifarnik", 2))
        BankaPartner.objects.filter(sifra=1).update(status="istorijska", napomena="ugašena")
        drugo = self.sinhronizuj([red_sifarnika(16, "Aikbank", "100618836", ulica="Bulevar 59a")])
        self.assertEqual((drugo.promenjenih, drugo.nestalih), (1, 1))
        self.assertIn({"banka": "16 Aikbank", "polje": "Ulica", "staro": "Bulevra 59a", "novo": "Bulevar 59a"}, drugo.promene)
        privredna = BankaPartner.objects.get(sifra=1)
        self.assertEqual((privredna.u_izvoru, privredna.status, privredna.napomena), (False, "istorijska", "ugašena"))
        banks = {b["code"]: b for b in fetch_banks(1)}
        self.assertEqual(banks[16]["ulica_par"], "Bulevar 59a")
        self.assertFalse(banks[1]["u_izvoru"])

    def test_empty_source_changes_nothing(self):
        self.sinhronizuj([red_sifarnika(16, "Aikbank")])
        run = self.sinhronizuj([])
        self.assertEqual(run.status, "failed")
        self.assertTrue(BankaPartner.objects.get(sifra=16).u_izvoru)

    def test_source_is_read_until_local_copy_exists(self):
        with patch("finansije.services.banks.source_banks", return_value=[red_sifarnika(3, "NLB")]) as izvor:
            self.assertEqual([b["code"] for b in fetch_banks(1)], [3])
            self.sinhronizuj([red_sifarnika(3, "NLB")])
            izvor.reset_mock()
            self.assertEqual([b["code"] for b in fetch_banks(1)], [3])
            izvor.assert_not_called()

    def test_nightly_schedule(self):
        from core.management.commands.sync_celery_periodic_tasks import EXPECTED_PERIODIC_TASKS

        self.assertTrue(any(s["task"] == "finansije.tasks.sync_banke_sifarnik_task" for s in EXPECTED_PERIODIC_TASKS))


class OznakeTests(TestCase):
    def setUp(self):
        cache.clear()
        with patch("finansije.services.banks.source_banks", return_value=[
                red_sifarnika(16, "Aikbank", "100618836"), red_sifarnika(15, "Aikbank", "100618836"),
                red_sifarnika(10, "Uprava za trezor")]):
            banke_nbs.sinhronizuj_sifarnik()
        self.user = get_user_model().objects.create_user("banke-oznake", password="test", allowed_center_codes="41")
        self.role = Role.objects.create(name="Banke oznake", slug="banke-oznake")
        self.user.roles.add(self.role)
        self.client.force_login(self.user)

    def grant(self, *codes):
        for code in codes:
            self.role.permissions.add(PermissionCode.objects.get_or_create(code=code)[0])

    def test_labels_require_whole_company_and_drive_the_list(self):
        url = reverse("finansije:bank_partner_edit", args=[15])
        glavna = BankaPartner.objects.get(sifra=16)
        self.grant("finansije:bank_partner_edit", "finansije:bank_list", "finansije:bank_detail")
        self.assertEqual(self.client.post(url, {"glavna": glavna.pk}).status_code, 403)
        self.grant("finansije:view_all")
        self.assertContains(self.client.get(reverse("finansije:bank_detail", args=[15])), "Sačuvaj oznake")
        self.client.post(url, {"glavna": glavna.pk, "namena": "ostalo", "status": "aktivna"})
        self.client.post(reverse("finansije:bank_partner_edit", args=[10]), {"status": "nije_banka", "napomena": "Trezor"})
        zapis = BankaPartner.objects.get(sifra=15)
        self.assertEqual((zapis.glavna, zapis.namena, zapis.izmenio), (glavna, "ostalo", self.user))
        lista = self.client.get(reverse("finansije:bank_list"), {"date_from": "2026-01-01", "date_to": "2026-06-30"})
        grupa = lista.context["bank_groups"][0]
        self.assertEqual([b["code"] for b in grupa["banke"]], [15])  # 16 nema ručni status ni NBS, pa je van spiska
        trezor = next(b for b in lista.context["outside_nbs"] if b["code"] == 10)
        self.assertEqual((trezor["nije_banka"], trezor["istorijska"]), (True, False))
        self.assertContains(lista, "Nije banka")

    def test_main_code_cannot_become_additional(self):
        self.grant("finansije:bank_partner_edit", "finansije:view_all")
        glavna, dodatna = BankaPartner.objects.get(sifra=16), BankaPartner.objects.get(sifra=15)
        dodatna.glavna = glavna
        dodatna.save()
        self.client.post(reverse("finansije:bank_partner_edit", args=[16]), {"glavna": BankaPartner.objects.get(sifra=10).pk})
        self.assertIsNone(BankaPartner.objects.get(sifra=16).glavna)


@patch("finansije.bank_views.service.fetch_banks", return_value=BANKS)
class BankeUSrbijiTests(TestCase):
    def setUp(self):
        from potrazivanja.models import FinancePartnerIdentity

        cache.clear()
        banke_nbs.osvezi(http=LaznaSesija(strane()), pauza=0)
        FinancePartnerIdentity.objects.create(company=1, partner_group=1, partner_code=17647, source_name="INTESA LEASING DOO", source_tax_id="103023875")
        save_entry(number=1, account="43500", partner_group=1, partner_code=17647, debit=Decimal("0"), credit=Decimal("250"))
        self.user = get_user_model().objects.create_user("banke-srbija", password="test", allowed_center_codes="41")
        self.role = Role.objects.create(name="Banke Srbija", slug="banke-srbija")
        self.user.roles.add(self.role)
        self.client.force_login(self.user)

    def grant(self, *codes):
        for code in codes:
            self.role.permissions.add(PermissionCode.objects.get_or_create(code=code)[0])

    def test_derived_fields_like_director_overview(self, source):
        intesa = BankaNbs.objects.get(pib="100001159")
        self.assertEqual(intesa.najveci_akcionar, {"naziv": "Intesa Holding International S.A.", "procenat": "100,00 %"})
        self.assertEqual((intesa.predsednik_uo, intesa.predsednik_io, intesa.revizor), ("Draginja Đurić", "Darko Popović", "BDO DOO BEOGRAD"))

    def test_page_is_opened_from_banks_and_shows_all_banks_and_groups(self, source):
        self.grant("finansije:bank_list", "finansije:bank_market", "finansije:bank_detail")
        self.assertContains(self.client.get(reverse("finansije:bank_list")), reverse("finansije:bank_market"))
        response = self.client.get(reverse("finansije:bank_market"), {"date_from": "2026-01-01", "date_to": "2026-06-30"})
        saradnja = {r["banka"].pib: [b["code"] for b in r["nase"]] for r in response.context["banke"]}
        self.assertEqual(saradnja, {"100001159": [11], "100001931": []})
        self.assertEqual(len(response.context["grupe"]), 19)
        self.assertEqual(response.context["grupe"][0]["grupa"].naziv, "Intesa Sanpaolo")  # sa nama prvo
        lizing = next(m for m in response.context["grupe"][0]["lica"] if m["lice"].pib == "103023875")
        self.assertEqual([p["sifra"] for p in lizing["partneri"]], [17647])
        self.assertIsNone(lizing["promet"])  # promet samo uz pristup celoj firmi
        self.assertContains(response, "ne radimo")

    def test_turnover_and_excel_with_whole_company_access(self, source):
        self.grant("finansije:bank_market", "finansije:view_all")
        period = {"date_from": "2026-01-01", "date_to": "2026-06-30"}
        response = self.client.get(reverse("finansije:bank_market"), period)
        lizing = next(m for m in response.context["grupe"][0]["lica"] if m["lice"].pib == "103023875")
        self.assertEqual((lizing["promet"]["stavki"], lizing["promet"]["potrazuje"]), (1, Decimal("250")))
        excel = self.client.get(reverse("finansije:bank_market"), {**period, "izvoz": "xlsx"})
        from io import BytesIO
        from openpyxl import load_workbook
        vrednosti = [c for r in load_workbook(BytesIO(excel.content)).active.iter_rows(values_only=True) for c in r if c is not None]
        for tekst in ("Banke – registar NBS", "Grupe i povezana lica", "Draginja Đurić – Predsednik upravnog odbora", "ne radimo", "Ukupno / prosek"):
            self.assertIn(tekst, vrednosti)
