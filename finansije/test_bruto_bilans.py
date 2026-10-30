from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import PermissionCode, Role
from .services.bruto_bilans import agregiraj, izvrsi
from .tests import save_entry


def red(knt, naziv, dug0="0", pot0="0", dug1="0", pot1="0", tip=1):
    dugu, potu = Decimal(dug0) + Decimal(dug1), Decimal(pot0) + Decimal(pot1)
    return {"k1": knt[:1], "k2": knt[:2], "k3": knt[:3], "knt": knt + " ", "naz_knt": naziv + "   ",
            "naz_k1": f"Klasa {knt[:1]}", "naz_k2": f"Grupa {knt[:2]}", "naz_k3": f"Sintetika {knt[:3]}",
            "dug0": Decimal(dug0), "pot0": Decimal(pot0), "dug1": Decimal(dug1), "pot1": Decimal(pot1),
            "dugu": dugu, "potu": potu, "saldo": dugu - potu, "tip": tip,
            "naz_tip": "VANBILANSNA AKTIVA I PASIVA" if tip == 2 else "BILANSNA AKTIVA I PASIVA"}


REDOVI = [red("24100", "Tekući račun", dug0="100", dug1="50", pot1="30"), red("24400", "Devizni račun", dug1="20"),
          red("43500", "Dobavljači", pot0="100", pot1="40"), red("89310", "Vanbilansno", pot0="7", tip=2)]


class AgregacijaTests(SimpleTestCase):
    def test_konto_level_keeps_accounts_and_separates_off_balance(self):
        bilansna, vanbilansna = agregiraj(REDOVI, "konto")
        self.assertEqual([s["konto"] for s in bilansna["stavke"]], ["24100", "24400", "43500"])
        self.assertEqual(bilansna["stavke"][0]["naziv"], "Tekući račun")
        self.assertEqual(bilansna["zbir"]["dugu"], Decimal("170"))
        self.assertEqual(bilansna["zbir"]["potu"], Decimal("170"))
        self.assertEqual(vanbilansna["naziv"], "VANBILANSNA AKTIVA I PASIVA")
        self.assertEqual((vanbilansna["zbir"]["saldo_d"], vanbilansna["zbir"]["saldo_p"]), (Decimal("0"), Decimal("7")))
        self.assertEqual((bilansna["zbir"]["saldo_d"], bilansna["zbir"]["saldo_p"]), (Decimal("140"), Decimal("140")))

    def test_synthetic_level_sums_accounts(self):
        bilansna = agregiraj(REDOVI, "k3")[0]
        self.assertEqual([(s["konto"], s["naziv"]) for s in bilansna["stavke"]], [("241", "Sintetika 241"), ("244", "Sintetika 244"), ("435", "Sintetika 435")])
        self.assertEqual(agregiraj(REDOVI, "k2")[0]["stavke"][0]["dug1"], Decimal("70"))
        self.assertEqual(agregiraj(REDOVI, "k1")[0]["stavke"][0]["saldo_d"], Decimal("140"))

    def test_saldo_is_split_per_account_before_summing(self):
        # 24100 je dugovni (100), 24190 potražni (30): grupa 241 nema saldo 70, nego 100 duguje i 30 potražuje.
        grupa = agregiraj([red("24100", "Tekući", dug1="100"), red("24190", "Prelazni", pot1="30")], "k3")[0]["stavke"][0]
        self.assertEqual((grupa["saldo_d"], grupa["saldo_p"]), (Decimal("100"), Decimal("30")))


class IzracunTests(TestCase):
    """Isto kao procedura SPFINizv12: POC je početno stanje, ostalo promet; filteri po periodu, OJ i kontima."""

    def setUp(self):
        for i, (konto, vrsta, dug, pot, dan, oj) in enumerate((
                ("52000", "POC", "100", "0", date(2026, 1, 1), 41), ("52000", "LD", "50", "0", date(2026, 3, 1), 41),
                ("52000", "ON", "-20", "0", date(2026, 4, 30), 20), ("43500", "UF", "0", "30", date(2026, 3, 1), 41),
                ("89310", "POC", "0", "7", date(2026, 1, 1), 41), ("52000", "LD", "999", "0", date(2027, 1, 5), 41)), 1):
            save_entry(number=i, account=konto, journal_type=vrsta, debit=Decimal(dug), credit=Decimal(pot), booking_date=dan,
                       organizational_unit=oj, year=dan.year)
        obrisana = save_entry(number=99, account="52000", journal_type="LD", debit=Decimal("5"), credit=Decimal("0"), booking_date=date(2026, 3, 2))
        obrisana.active = False  # stavka obrisana u izvoru ne ulazi u bilans
        obrisana.save()

    def test_opening_turnover_and_off_balance(self):
        redovi = {r["knt"]: r for r in izvrsi(2026, date(2026, 1, 1), date(2026, 12, 31))}
        self.assertEqual((redovi["52000"]["dug0"], redovi["52000"]["dug1"], redovi["52000"]["dugu"]), (Decimal("100"), Decimal("30"), Decimal("130")))
        self.assertEqual((redovi["43500"]["potu"], redovi["43500"]["tip"]), (Decimal("30"), 1))
        self.assertEqual((redovi["89310"]["tip"], redovi["89310"]["naz_tip"]), (2, "VANBILANSNA AKTIVA I PASIVA"))

    def test_filters_like_the_procedure(self):
        self.assertEqual({r["knt"] for r in izvrsi(2026, date(2026, 1, 1), date(2026, 12, 31), konto_do="5")}, {"43500", "52000"})
        self.assertEqual({r["knt"] for r in izvrsi(2026, date(2026, 1, 1), date(2026, 12, 31), konto_od="5", konto_do="5")}, {"52000"})
        po_oj = {r["knt"]: r for r in izvrsi(2026, date(2026, 1, 1), date(2026, 12, 31), oj_od=41, oj_do=41)}
        self.assertEqual(po_oj["52000"]["dugu"], Decimal("150"))
        self.assertEqual(izvrsi(2026, date(2026, 2, 1), date(2026, 3, 31))[1]["dugu"], Decimal("50"))


@patch("finansije.services.bruto_bilans.izvrsi", return_value=REDOVI)
class BrutoBilansViewTests(TestCase):
    PARAMETRI = {"godina": "2026", "datum_od": "01.01.2026", "datum_do": "31.12.2026", "nivo": "konto"}

    def setUp(self):
        self.user = get_user_model().objects.create_user("bilans-reader", password="test")
        self.role = Role.objects.create(name="Bilans test", slug="bilans-test")
        self.user.roles.add(self.role)
        self.client.force_login(self.user)
        self.url = reverse("finansije:bruto_bilans")

    def grant(self, *codes):
        for code in codes:
            self.role.permissions.add(PermissionCode.objects.get_or_create(code=code)[0])

    def test_requires_permission_and_whole_company(self, procedura):
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.grant("finansije:bruto_bilans")
        self.assertEqual(self.client.get(self.url, {**self.PARAMETRI, "generisi": "1"}).status_code, 403)
        procedura.assert_not_called()

    def test_opening_page_does_not_run_procedure(self, procedura):
        self.grant("finansije:bruto_bilans", "finansije:view_all")
        response = self.client.get(self.url)
        self.assertContains(response, "Generiši izveštaj")
        self.assertContains(response, 'href="%s"' % self.url)
        procedura.assert_not_called()

    def test_generate_passes_parameters_and_shows_totals(self, procedura):
        self.grant("finansije:bruto_bilans", "finansije:view_all")
        response = self.client.get(self.url, {**self.PARAMETRI, "oj_od": "20", "konto_do": "5", "generisi": "1"})
        procedura.assert_called_once_with(2026, date(2026, 1, 1), date(2026, 12, 31), 20, None, "", "5")
        self.assertContains(response, "Tekući račun")
        self.assertContains(response, "Svega · BILANSNA AKTIVA I PASIVA")

    def test_period_outside_year_is_rejected(self, procedura):
        self.grant("finansije:bruto_bilans", "finansije:view_all")
        response = self.client.get(self.url, {**self.PARAMETRI, "datum_do": "01.01.2027", "generisi": "1"})
        self.assertContains(response, "unutar izabrane godine")
        procedura.assert_not_called()

    def test_excel_export(self, procedura):
        self.grant("finansije:bruto_bilans", "finansije:view_all")
        response = self.client.get(self.url, {**self.PARAMETRI, "nivo": "k3", "izvoz": "xlsx"})
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        vrednosti = [c for r in load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True) for c in r if c is not None]
        self.assertIn("241", vrednosti)
        self.assertIn("Svega", vrednosti)
        self.assertIn(170.0, vrednosti)
