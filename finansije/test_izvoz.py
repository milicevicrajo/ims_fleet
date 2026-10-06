"""Izvoz ekrana Finansija u Excel i PDF, nazivi konta i otvaranje konta po dubini (05.10.2026.)."""
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .models import FinanceJob, LedgerEntry, SyncRun
from .tests import job, save_entry

PLAN = {"5": "RASHODI", "51": "TROŠKOVI MATERIJALA", "510": "Nabavna vrednost materijala", "51000": "Materijal",
        "6": "PRIHODI", "61": "PRIHODI OD PRODAJE", "610": "Prihodi od prodaje robe", "61000": "Prihod od usluga"}


def vrednosti(odgovor):
    return [c for red in load_workbook(BytesIO(odgovor.content)).active.iter_rows(values_only=True) for c in red if c is not None]


class IzvozTests(TestCase):
    def setUp(self):
        self.client.force_login(get_user_model().objects.create_superuser("fin-izvoz", "izvoz@example.invalid", "x"))
        FinanceJob.objects.create(**job())
        FinanceJob.objects.create(**job("420001", "42", name="Posao B"))
        save_entry()                                                                          # 41: prihod 100
        save_entry(number=2, account="51000", debit=Decimal("40"), credit=Decimal("0"))         # 41: rashod 40
        save_entry(number=5, center="42", job_code="420001", job_name="Posao B", credit=Decimal("120"))
        save_entry(number=6, account="20400", debit=Decimal("500"), credit=Decimal("0"), partner_name="=1+1")
        self.params = {"date_from": "2026-02-01", "date_to": "2026-02-28", "group": "center", "kind": "pnl"}
        patch("finansije.services.kontni_plan.nazivi_konta", return_value=PLAN).start()
        patch("finansije.services.job_additional.payroll_headcounts",
              return_value={"counts": {}, "closed_months": [], "open_months": []}).start()
        self.addCleanup(patch.stopall)

    def izvoz(self, **parametri):
        return self.client.get(reverse("finansije:export"), {**self.params, **parametri})

    def redovi(self, **parametri):
        odgovor = self.client.get(reverse("finansije:report"), {**self.params, "group": "account", **parametri})
        self.assertEqual(odgovor.status_code, 200)
        return {r["code"]: r for r in odgovor.context["rows"]}, odgovor

    def test_konta_se_otvaraju_po_dubini_sa_nazivima(self):
        redovi, _ = self.redovi()
        self.assertEqual({k: r["label"] for k, r in redovi.items()}, {"5": "RASHODI", "6": "PRIHODI"})
        self.assertIn("account=5", redovi["5"]["drill_url"])
        self.assertNotIn("account_exact", redovi["5"]["drill_url"])
        redovi, odgovor = self.redovi(account="5")
        self.assertEqual(set(redovi), {"51"})
        self.assertEqual(odgovor.context["konto_nivo"], "2")
        self.assertContains(odgovor, "5 RASHODI")
        redovi, _ = self.redovi(account="510")
        self.assertEqual((set(redovi), redovi["51000"]["label"]), ({"51000"}, "Materijal"))
        self.assertNotIn("drill_url", redovi["51000"])
        self.assertIn("account_exact=1", redovi["51000"]["ledger_url"])
        redovi, _ = self.redovi(nivo="3")
        self.assertEqual(set(redovi), {"510", "610"})

    def test_pdf_postoji_za_sve_ekrane(self):
        for parametri, naslov in (({"report": "overview"}, "Finansijski pregled"),
                                  ({}, "Izveštaj po centrima"),
                                  ({"group": "account", "account": "5"}, "Struktura po kontima"),
                                  ({"group": "month"}, "Mesečni pregled"),
                                  ({"group": "job", "analysis": "standard"}, "Šifre posla"),
                                  ({"group": "job", "analysis": "additional"}, "Šifre posla — dodatne analize"),
                                  ({"report": "ledger"}, "Finansijska knjiženja")):
            with self.subTest(naslov=naslov):
                odgovor = self.izvoz(format="pdf", **parametri)
                self.assertEqual(odgovor.status_code, 200)
                self.assertTemplateUsed(odgovor, "finansije/izvoz_stampa.html")
                self.assertContains(odgovor, f"<h1>{naslov}</h1>", html=False)
                self.assertContains(odgovor, "Sačuvaj kao PDF")
        self.assertContains(self.izvoz(format="pdf", group="account", account="5"), "Otvoreno: 5 RASHODI")

    def test_pdf_knjizenja_je_ogranicen_a_zbir_je_za_sve(self):
        with patch("finansije.services.izvoz_ekrani.MAKS_KNJIZENJA_PDF", 1):
            odgovor = self.izvoz(format="pdf", report="ledger", kind="all")
        self.assertContains(odgovor, "Prikazano prvih 1 od 4 knjiženja")
        self.assertRegex(odgovor.content.decode(), r"540[.,]00")  # duguje svih knjiženja: 40 + 500

    def test_excel_finansijskog_pregleda(self):
        odgovor = self.izvoz(report="overview")
        self.assertEqual(odgovor.status_code, 200)
        self.assertIn("finansijski_pregled.xlsx", odgovor["Content-Disposition"])
        sadrzaj = vrednosti(odgovor)
        self.assertIn("Institut za ispitivanje materijala a.d. · Finansijski pregled", sadrzaj)
        self.assertIn("Šifre posla — direktni rezultat", sadrzaj)
        self.assertIn(220, sadrzaj)  # prihodi ukupno, kao broj

    def test_detalj_sifre_excel_i_pdf_sa_nazivima_konta(self):
        SyncRun.objects.create(company=1, year_from=2026, year_to=2026, status="success")
        parametri = {"report": "job_card", "job": "410001", "year": "2026", "month": ""}
        odgovor = self.client.get(reverse("finansije:export"), parametri)
        self.assertEqual(odgovor.status_code, 200)
        sadrzaj = vrednosti(odgovor)
        for tekst in ("Struktura rashoda · sintetička konta", "510", "Nabavna vrednost materijala", "Izdate fakture · IF"):
            self.assertIn(tekst, sadrzaj)
        celija = next(c for red in load_workbook(BytesIO(odgovor.content)).active.iter_rows() for c in red
                      if isinstance(c.value, str) and c.value.endswith("=1+1"))
        self.assertEqual(celija.data_type, "s")
        pdf = self.client.get(reverse("finansije:export"), {**parametri, "format": "pdf"})
        self.assertContains(pdf, "Detalj šifre posla 410001")
        self.assertContains(pdf, "Nabavna vrednost materijala")
        stranica = self.client.get(reverse("finansije:job_card"), {"job": "410001", "year": "2026"})
        self.assertContains(stranica, "report=job_card")

    def test_detalj_sifre_bez_pristupa_ili_nepostojeca(self):
        self.assertEqual(self.client.get(reverse("finansije:export"), {"report": "job_card", "job": "999999",
                                                                       "year": "2026"}).status_code, 404)
        self.assertEqual(self.client.get(reverse("finansije:export"), {"report": "job_card"}).status_code, 400)

    def test_struktura_rashoda_ima_naziv_konta(self):
        SyncRun.objects.create(company=1, year_from=2026, year_to=2026, status="success")
        odgovor = self.client.get(reverse("finansije:job_table", args=["expenses"]), {"job": "410001", "year": "2026", "month": ""})
        red = odgovor.json()["data"][0]
        self.assertEqual((red[0]["sort"], red[1]["sort"]), ("510", "Nabavna vrednost materijala"))

    def test_konta_u_bocnom_meniju(self):
        odgovor = self.client.get(reverse("finansije:dashboard"))
        self.assertContains(odgovor, "?group=account")
        odgovor = self.client.get(reverse("finansije:report"), {"group": "account"})
        self.assertTrue(odgovor.context["valid"], odgovor.context["form"].errors)
        self.assertContains(odgovor, 'class="sidebar-link active" data-finance-section="accounts"')

    def test_dugmad_za_izvoz_na_ekranima(self):
        for ime, parametri in (("finansije:dashboard", {}), ("finansije:report", {}), ("finansije:ledger", {}),
                               ("finansije:report", {"group": "job"})):
            with self.subTest(ime=ime, parametri=parametri):
                self.assertContains(self.client.get(reverse(ime), {**self.params, **parametri}), "format=pdf")
        self.assertEqual(LedgerEntry.objects.count(), 4)
