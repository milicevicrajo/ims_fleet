from datetime import date
from decimal import Decimal as D
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.db import DatabaseError
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import PermissionCode, Role
from . import test_job_overview as fixtures
from .services.job_additional import enrich_additional, METRICS, EXTRA_KEYS, EXTRA_LABELS


def people(company, codes, start, end):
    return {'counts': {code: {m: 2 for m in range(start.month,end.month+1)} for code in codes},
            'closed_months':list(range(start.month,end.month+1)), 'open_months':[]}


class AdditionalCalculationTests(SimpleTestCase):
    def row(self):
        return {'code':'a','metrics':{key:dict(value=D(value),complete=True,note='')
                for key,value in zip(METRICS,(100,-40,60,-10,50,120,-150,-30))}}

    def calculate(self, data, start=date(2026,1,1), end=date(2026,2,28)):
        with patch('finansije.services.job_additional.payroll_headcounts',return_value=data):
            return enrich_additional([self.row()],start,end)[0]['additional']

    def test_ratios_and_per_person_use_monthly_average_not_sum(self):
        extra=self.calculate({'counts':{'a':{1:2,2:4}},'closed_months':[1,2],'open_months':[]})
        self.assertEqual(extra['people']['value'],3)
        self.assertEqual(extra['margin_after']['value'],50)
        self.assertEqual(extra['expense_share']['value'],40)
        self.assertEqual(extra['shared_share']['value'],10)
        self.assertEqual(extra['shared_cost_person']['value'],D(-10)/3)
        self.assertEqual(extra['net_cash_person']['value'],-10)
        self.assertTrue(extra['people']['complete'])

    def test_zero_employee_month_included_but_unclosed_month_is_not_assumed_zero(self):
        extra=self.calculate({'counts':{'a':{1:4}},'closed_months':[1,2],'open_months':[]})
        self.assertEqual(extra['people']['value'],2)
        extra=self.calculate({'counts':{'a':{1:4}},'closed_months':[1],'open_months':[2]})
        self.assertEqual(extra['people']['value'],4)
        self.assertFalse(extra['revenue_person']['complete'])
        self.assertIn('1 od 2',extra['people']['note'])

    def test_missing_and_zero_headcounts_do_not_make_infinite_or_zero_rates(self):
        for months in ([],[1,2]):
            extra=self.calculate({'counts':{},'closed_months':months,'open_months':[]})
            self.assertIsNone(extra['revenue_person']['value'])
            self.assertEqual(extra['margin']['value'],60)

    def test_missing_shared_cost_and_nonpositive_revenue_block_only_dependent_ratios(self):
        row=self.row();row['metrics']['shared_cost']['value']=None
        row['metrics']['after_shared']['value']=None
        row['metrics']['revenue']['value']=D(0)
        with patch('finansije.services.job_additional.payroll_headcounts',side_effect=people):
            result=enrich_additional([row],date(2026,1,1),date(2026,1,31))[0]['additional']
        self.assertIsNone(result['margin']['value'])
        self.assertIsNone(result['shared_cost_person']['value'])
        self.assertEqual(result['expense_person']['value'],-20)

    def test_multiple_years_source_error_and_partial_month_are_flagged(self):
        with patch('finansije.services.job_additional.payroll_headcounts',side_effect=[
            {'counts':{'a':{12:4}},'closed_months':[12],'open_months':[]},DatabaseError('private source')]) as source:
            with self.assertLogs('finansije.services.job_additional',level='ERROR'):
                result=enrich_additional([self.row()],date(2025,12,10),date(2026,1,20))[0]['additional']
        self.assertEqual(source.call_count,2)
        self.assertEqual(result['people']['value'],4)
        self.assertFalse(result['people']['complete'])
        self.assertNotIn('private source',result['people']['note'])


class AdditionalOverviewTests(TestCase):
    def setUp(self):
        fixtures.JobOverviewTests.setUp(self)
        self.params['analysis']='additional'
        self.people=patch('finansije.services.job_additional.payroll_headcounts',side_effect=people).start()

    def test_page_has_separate_lazy_tab_and_no_external_reads(self):
        response=self.client.get(reverse('finansije:report'),self.params)
        self.assertContains(response,'Dodatne analize')
        self.assertContains(response,'id="FinanceAdditionalTable"')
        self.assertContains(response,'data-lazy-tab="true"')
        self.assertContains(response,'id="jobs-additional" class="tab-pane active"')
        self.people.assert_not_called();self.shared.assert_not_called()

    def test_additional_ajax_preserves_absolute_metrics_and_limits_people_query(self):
        data=self.client.get(reverse('finansije:jobs_data'),self.params).json()['data']
        row=next(r for r in data if r[0]['sort']=='410001')
        self.assertEqual(len(row),3+len(METRICS)+len(EXTRA_KEYS))
        self.assertEqual([D(c['sort']) for c in row[3:11]],list(map(D,(100,-40,60,-10,50,120,-150,-30))))
        extra=dict(zip(EXTRA_KEYS,row[11:]))
        self.assertEqual(D(extra['people']['sort']),2)
        self.assertEqual(D(extra['shared_cost_person']['sort']),-5)
        self.people.assert_called_once_with(1,{'410001','410002'},date(2026,2,1),date(2026,2,28))

    def test_export_matches_additional_table_and_keeps_numeric_values(self):
        response=self.client.get(reverse('finansije:export'),self.params)
        records=list(load_workbook(BytesIO(response.content),read_only=True).active.values)
        self.assertEqual(records[2][11:11+len(EXTRA_LABELS)],EXTRA_LABELS)
        self.assertEqual(records[3][11+EXTRA_KEYS.index('shared_cost_person')],-5)

    def test_additional_excel_has_full_native_table_and_percentage_point_format(self):
        response=self.client.get(reverse('finansije:export'),self.params)
        sheet=load_workbook(BytesIO(response.content)).active
        table=sheet.tables['DodatneAnalize']
        self.assertEqual(table.ref,f'A3:Z{sheet.max_row}')
        self.assertEqual(table.autoFilter.ref,table.ref)
        self.assertEqual(len(table.tableColumns),26)
        self.assertEqual(sheet['L4'].value,60)
        self.assertEqual(sheet['L4'].data_type,'n')
        self.assertIn('"%"',sheet['L4'].number_format)
        self.assertEqual(sheet.freeze_panes,'D4')

    def test_sidebar_links_directly_to_additional_and_marks_active_section(self):
        from types import SimpleNamespace
        from finansije.templatetags.finance import finance_sidebar_section
        request=SimpleNamespace(resolver_match=SimpleNamespace(view_name='finansije:report'),GET=self.params)
        self.assertEqual(finance_sidebar_section(request),'additional')
        request.GET=dict(self.params,analysis='standard')
        self.assertEqual(finance_sidebar_section(request),'jobs')
        response=self.client.get(reverse('finansije:report'),self.params)
        self.assertContains(response,'data-finance-section="additional"')
        self.assertContains(response,'?group=job&amp;analysis=additional')
        self.assertContains(response,'Izvezi u Excel',count=2)

    def test_restricted_finance_role_cannot_read_people_and_invalid_center_reads_nothing(self):
        user=get_user_model().objects.create_user('additional-limited',allowed_center_codes='41')
        role=Role.objects.create(name='Additional test',slug='additional-test')
        role.permissions.add(PermissionCode.objects.create(code='finansije:dashboard'))
        user.roles.add(role);self.client.force_login(user)
        response=self.client.get(reverse('finansije:jobs_data'),self.params)
        self.assertEqual(response.status_code,200)
        data=response.json()['data']
        self.assertEqual({row[0]['sort'] for row in data},{'410001','410002'})
        self.assertTrue(all(row[11+EXTRA_KEYS.index('people')]['sort']=='' for row in data))
        self.people.assert_not_called()
        self.shared.reset_mock()
        self.assertEqual(self.client.get(reverse('finansije:jobs_data'),dict(self.params,center='42')).status_code,400)
        self.shared.assert_not_called()
