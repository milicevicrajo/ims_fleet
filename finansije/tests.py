from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.db import IntegrityError, transaction
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.models import OrganizationalUnit, PermissionCode, Role
from .models import FinanceJob, LedgerEntry, SyncRun
from .services.reports import base_querysets, summary
from .services.source import SOURCE_FIELDS, fetch_source, fingerprint, totals_for
from .services.sync import SyncBusy, sync_ledger, sync_lock


def posting(number=1, **changes):
    row = dict(
        company=1, year=2026, journal_type="IF", journal_number=number, line_number=1,
        organizational_unit=41, organizational_unit_name="OJ 41", account="61000", account_name="Prihod",
        partner_group=1, partner_code=123, partner_name="Partner", document_date=date(2026, 1, 8),
        document_reference=f"IF-{number}", debit=Decimal("0.00"), credit=Decimal("100.00"),
        currency="RSD", foreign_amount=None, description="", linked_line=None, due_date=None,
        change_status="O", job_code="410001", job_name="Posao A", center="41", booking_date=date(2026, 2, 10),
        debit_credit_flag="P", source_paid_amount=None,
    )
    row.update(changes)
    return row


def job(code="410001", center="41", **changes):
    row = dict(company=1, code=code, name="Posao A", center=center, active=True, profit_type="P")
    row.update(changes)
    return row


def save_entry(**changes):
    row = posting(**changes)
    return LedgerEntry.objects.create(**row, source_hash=fingerprint(row), changed_at=timezone.now())


class SyncTests(TestCase):
    def run_sync(self, rows, jobs=None, **kwargs):
        with patch("finansije.services.sync.fetch_source", return_value=(rows, jobs if jobs is not None else [job()])):
            return sync_ledger(year_from=2025, year_to=2026, **kwargs)

    def test_repeat_update_remove_and_reactivate_preserve_identity(self):
        run = self.run_sync([posting(), posting(2)])
        self.assertEqual((run.created, run.source_rows, run.status), (2, 2, "success"))
        entry = LedgerEntry.objects.get(journal_number=1)
        original_pk, original_changed = entry.pk, entry.changed_at
        run = self.run_sync([posting(), posting(2)])
        self.assertEqual((run.created, run.updated, run.unchanged), (0, 0, 2))
        entry.refresh_from_db()
        self.assertEqual(entry.changed_at, original_changed)
        run = self.run_sync([posting(credit=Decimal("150.00"))])
        self.assertEqual((run.updated, run.removed), (1, 1))
        entry.refresh_from_db()
        self.assertEqual(entry.pk, original_pk)
        self.assertEqual(entry.credit, Decimal("150.00"))
        self.assertFalse(LedgerEntry.objects.get(journal_number=2).active)
        run = self.run_sync([posting(credit=Decimal("150.00")), posting(2)])
        self.assertEqual(run.updated, 1)
        restored = LedgerEntry.objects.get(journal_number=2)
        self.assertTrue(restored.active)
        self.assertIsNone(restored.removed_at)

    def test_current_year_refresh_does_not_remove_prior_year(self):
        self.run_sync([posting(year=2025, booking_date=date(2025, 2, 1)), posting(2)])
        with patch("finansije.services.sync.fetch_source", return_value=([posting(3)], [job()])):
            sync_ledger(year_from=2026, year_to=2026)
        self.assertTrue(LedgerEntry.objects.get(year=2025).active)
        self.assertFalse(LedgerEntry.objects.get(journal_number=2).active)

    def test_empty_source_does_not_erase_existing_scope(self):
        self.run_sync([posting()])
        with self.assertRaisesMessage(ValueError, "Izvor je prazan"):
            self.run_sync([])
        self.assertTrue(LedgerEntry.objects.get().active)
        self.assertEqual(SyncRun.objects.first().status, "failed")

    def test_duplicate_source_fails_before_publishing(self):
        self.run_sync([posting()])
        with self.assertRaisesMessage(ValueError, "Dupliran"):
            self.run_sync([posting(), posting()])
        self.assertEqual(LedgerEntry.objects.count(), 1)
        self.assertEqual(SyncRun.objects.first().status, "failed")

    def test_source_failure_preserves_previous_report(self):
        self.run_sync([posting()])
        with patch("finansije.services.sync.fetch_source", side_effect=ConnectionError("Izvor nedostupan")):
            with self.assertRaises(ConnectionError):
                sync_ledger(year_from=2025, year_to=2026)
        self.assertEqual(LedgerEntry.objects.get().credit, Decimal("100"))
        self.assertEqual(SyncRun.objects.filter(status="success").count(), 1)

    def test_failed_reconciliation_rolls_back_entries_and_dimensions(self):
        self.run_sync([posting()])
        wrong = totals_for([posting(credit=Decimal("999"))])
        with patch("finansije.services.sync.totals_for", return_value=wrong):
            with self.assertRaisesMessage(ValueError, "Kontrolni zbir"):
                self.run_sync([posting(credit=Decimal("200"))], [job(name="Izmenjen naziv")])
        self.assertEqual(LedgerEntry.objects.get().credit, Decimal("100"))
        self.assertEqual(FinanceJob.objects.get().name, "Posao A")

    def test_missing_dimension_is_not_dropped(self):
        self.run_sync([posting(job_code="", job_name="", center="", account_name="")], [])
        self.assertEqual(LedgerEntry.objects.count(), 1)

    def test_database_enforces_source_identity(self):
        self.run_sync([posting()])
        with self.assertRaises(IntegrityError), transaction.atomic():
            save_entry()

    def test_overlapping_run_is_rejected(self):
        with sync_lock():
            with self.assertRaises(SyncBusy):
                self.run_sync([posting()])
        self.assertEqual(SyncRun.objects.count(), 0)

    def test_out_of_scope_rows_rejected(self):
        with self.assertRaisesMessage(ValueError, "van zahtevanog"):
            self.run_sync([posting(company=2)])
        self.assertFalse(LedgerEntry.objects.exists())

    def test_changed_center_and_name_are_refreshed(self):
        self.run_sync([posting()])
        run = self.run_sync([posting(center="42", job_name="Novi naziv")], [job(center="42", name="Novi naziv")])
        self.assertEqual(run.updated, 1)
        self.assertEqual(LedgerEntry.objects.get().center, "42")

    def test_mixed_nullable_amounts_survive_roundtrip(self):
        self.run_sync([
            posting(foreign_amount=Decimal("1000.00"), debit=Decimal("5160400.00")),
            posting(2, foreign_amount=None, source_paid_amount=None),
        ])
        self.assertEqual(LedgerEntry.objects.get(journal_number=1).foreign_amount, Decimal("1000.00"))
        self.assertIsNone(LedgerEntry.objects.get(journal_number=2).foreign_amount)
        self.run_sync([
            posting(foreign_amount=None, debit=Decimal("5160400.00")),
            posting(2, foreign_amount=Decimal("1000.00"), source_paid_amount=None),
        ])
        self.assertIsNone(LedgerEntry.objects.get(journal_number=1).foreign_amount)
        self.assertEqual(LedgerEntry.objects.get(journal_number=2).foreign_amount, Decimal("1000.00"))


class SourceReadTests(TestCase):
    def mock_cursor(self, connections, expected_credit=Decimal("100")):
        cursor = connections.__getitem__.return_value.cursor.return_value.__enter__.return_value
        cursor.fetchall.side_effect = [
            [("410001 ", "Posao A", "41 ", "D", "P")],
            [("61000 ", "Prihod")], [("2026", 41, "OJ 41")], [(1, 123, "Partner")],
            [("2026", "61000 ", 1, Decimal("0"), expected_credit)],
        ]
        row = posting()
        raw = [row[name] for name in SOURCE_FIELDS]
        raw[1], raw[6], raw[19] = "2026", "61000 ", "410001 "
        cursor.fetchmany.side_effect = [[tuple(raw)], []]
        return cursor

    @patch("finansije.services.source.connections")
    def test_source_normalizes_keys_and_uses_only_parameterized_reads(self, connections):
        cursor = self.mock_cursor(connections)
        rows, jobs = fetch_source(1, 2025, 2026)
        self.assertEqual(rows[0]["center"], "41")
        self.assertEqual(rows[0]["account"], "61000")
        self.assertEqual(jobs[0]["code"], "410001")
        for call in cursor.execute.call_args_list:
            self.assertTrue(call.args[0].lstrip().startswith("SELECT"))
            self.assertTrue(call.args[1])

    @patch("finansije.services.source.connections")
    def test_source_change_during_read_is_detected(self, connections):
        self.mock_cursor(connections, Decimal("101"))
        with self.assertRaisesMessage(ValueError, "Izvor je izmenjen"):
            fetch_source(1, 2025, 2026)


class FinancePermissionSetupTests(TestCase):
    def test_setup_is_scoped_idempotent_and_restricted_by_default(self):
        from core.permissions import sync_finance_permissions

        other = Role.objects.create(name="Drugi modul", slug="other-module")
        other.permissions.add(PermissionCode.objects.create(code="other:view"))
        sync_finance_permissions()
        sync_finance_permissions()
        finance = Role.objects.get(slug="finansije")
        self.assertSetEqual(set(finance.permissions.values_list("code", flat=True)), {
            "finansije:dashboard", "finansije:ledger", "finansije:export",
        })
        self.assertTrue(Role.objects.get(slug="uprava").permissions.filter(code="finansije:view_all").exists())
        self.assertEqual(list(other.permissions.values_list("code", flat=True)), ["other:view"])


class ManualSyncTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser("sync-admin", "sync@example.invalid", "test")
        self.client.force_login(self.user)
        self.year = timezone.localdate().year
        self.url = reverse("finansije:sync_run")

    def post(self, scope="current", **extra):
        return self.client.post(self.url, {"scope": scope}, HTTP_X_REQUESTED_WITH="XMLHttpRequest", **extra)

    def test_manual_sync_calls_service_directly_without_celery_and_logs_results(self):
        with patch("finansije.services.sync.fetch_source", return_value=([posting(year=self.year)], [job()])) as fetch, \
                patch("celery.app.task.Task.apply_async", side_effect=AssertionError("Manual sync must not enqueue a task")):
            response = self.post()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "success")
        fetch.assert_called_once_with(1, self.year, self.year)
        self.assertEqual(SyncRun.objects.get().status, "success")
        self.assertEqual(LedgerEntry.objects.count(), 1)
        self.assertIn("Preuzeto: 1", response.json()["message"])

    def test_all_years_scope_and_plain_form_redirect(self):
        with patch("finansije.services.sync.fetch_source", return_value=([posting(year=self.year)], [job()])) as fetch:
            response = self.client.post(self.url, {"scope": "all"})
        fetch.assert_called_once_with(1, 2025, self.year)
        self.assertRedirects(response, reverse("finansije:sync_status"))

    def test_invalid_scope_and_get_do_not_start_import(self):
        with patch("finansije.views.sync_ledger") as sync:
            self.assertEqual(self.client.get(self.url).status_code, 405)
            self.assertEqual(self.post("unexpected").status_code, 400)
            self.assertEqual(self.client.post(self.url, {}, HTTP_X_REQUESTED_WITH="XMLHttpRequest").status_code, 400)
        sync.assert_not_called()

    def test_overlapping_manual_sync_returns_conflict_without_new_run(self):
        with sync_lock():
            response = self.post()
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["status"], "busy")
        self.assertFalse(SyncRun.objects.exists())

    def test_manual_failure_keeps_previous_entries_and_is_not_reported_as_success(self):
        save_entry()
        with patch("finansije.services.sync.fetch_source", side_effect=ConnectionError("Source unavailable")), \
                self.assertLogs("finansije.views", level="ERROR"):
            response = self.post()
        self.assertEqual(response.status_code, 500)
        self.assertEqual(response.json()["status"], "error")
        self.assertEqual(SyncRun.objects.get().status, "failed")
        self.assertEqual(LedgerEntry.objects.get().credit, Decimal("100"))

    def test_manual_sync_requires_global_scope_permission_and_csrf(self):
        user = get_user_model().objects.create_user("sync-limited", password="test", allowed_center_codes="41")
        role = Role.objects.create(name="Sync test", slug="sync-test")
        role.permissions.add(PermissionCode.objects.create(code="finansije:sync_status"))
        user.roles.add(role)
        self.client.force_login(user)
        with patch("finansije.views.sync_ledger") as sync:
            self.assertEqual(self.post().status_code, 403)
            user.roles.clear()
            self.assertEqual(self.post().status_code, 403)
            self.client.logout()
            self.assertEqual(self.post().status_code, 302)
            csrf_client = Client(enforce_csrf_checks=True)
            csrf_client.force_login(self.user)
            self.assertEqual(csrf_client.post(self.url, {"scope": "current"}).status_code, 403)
        sync.assert_not_called()

    def test_celery_entry_points_use_the_same_function_and_sync_queue(self):
        from ims_erp.celery import app
        from .tasks import sync_all_years, sync_current_year, sync_ledger_task

        run = SyncRun(year_from=2025, year_to=self.year, source_rows=1, created=1)
        with patch("finansije.tasks.sync_ledger", return_value=run) as sync:
            sync_ledger_task.run(year_from=2025, year_to=self.year)
            sync.assert_called_with(year_from=2025, year_to=self.year)
            sync_current_year.run()
            sync.assert_called_with(year_from=self.year, year_to=None)
            sync_all_years.run()
            sync.assert_called_with(year_from=2025, year_to=None)
        self.assertIn("finansije.tasks.sync_ledger_task", app.tasks)
        route = app.amqp.router.route({}, "finansije.tasks.sync_ledger_task")
        self.assertEqual(route["queue"].name, "sync")
        with patch("finansije.tasks.sync_ledger", side_effect=SyncBusy):
            self.assertIn("skipped", sync_ledger_task.run())
        with patch("finansije.tasks.sync_ledger", side_effect=ConnectionError):
            with self.assertRaises(ConnectionError):
                sync_ledger_task.run()


class ReportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser("finance-admin", "finance@example.invalid", "test")
        FinanceJob.objects.create(**job())
        FinanceJob.objects.create(**job("420001", "42", name="Posao B"))
        FinanceJob.objects.create(**job("410002", "41", name="Bez prometa"))
        save_entry()
        save_entry(number=2, account="51000", debit=Decimal("40"), credit=Decimal("0"))
        save_entry(number=3, account="51000", debit=Decimal("0"), credit=Decimal("10"))
        save_entry(number=4, debit=Decimal("20"), credit=Decimal("0"))
        save_entry(number=5, center="42", job_code="420001", job_name="Posao B", credit=Decimal("120"))
        save_entry(number=6, account="20400", debit=Decimal("500"), credit=Decimal("0"))
        self.params = {"date_from": "2026-02-01", "date_to": "2026-02-28", "group": "center", "kind": "pnl"}
        self.client.force_login(self.user)

    def test_net_amounts_and_shares_with_reversals(self):
        response = self.client.get(reverse("finansije:report"), self.params)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["totals"], {"revenue": Decimal("200"), "expense": Decimal("30"), "result": Decimal("170"), "count": 5})
        row = response.context["rows"][0]
        self.assertEqual(row["revenue_share"], Decimal("40"))
        self.assertEqual(row["expense_share"], Decimal("100"))
        self.assertContains(response, "Finansijska analitika")

    def test_booking_date_drives_period_not_document_date(self):
        response = self.client.get(reverse("finansije:report"), dict(self.params, date_from="2026-01-01", date_to="2026-01-31"))
        self.assertEqual(response.context["totals"]["count"], 0)

    def test_overview_only_filters_period_and_links_keep_selected_dates(self):
        response = self.client.get(reverse("finansije:dashboard"), dict(self.params, center="42", kind="expense"))
        self.assertNotContains(response, 'class="finance-filter-grid"')
        self.assertNotContains(response, '<table')
        self.assertEqual(response.context["period_from"], date(2026, 2, 1))
        self.assertContains(response, 'name="date_from"')
        self.assertEqual(response.context["totals"]["revenue"], Decimal("200"))
        groups = response.context["chart_groups"]
        self.assertEqual([group["dimension"] for group in groups], ["center", "job"])
        for group in groups:
            for metric in ("revenue", "expense", "result"):
                self.assertEqual(sum(row[metric] for row in group["rows"]), response.context["totals"][metric])
            for row in group["rows"]:
                self.assertEqual(len(row["metrics"]), 3)
                detail = self.client.get(row["url"])
                self.assertEqual(detail.status_code, 200)
                self.assertEqual(detail.context["totals"]["result"], row["result"])
                self.assertEqual(detail.context["form"].cleaned_data["date_from"], date(2026, 2, 1))
                self.assertEqual(detail.context["form"].cleaned_data["date_to"], date(2026, 2, 28))

    def test_overview_period_updates_totals_charts_and_rankings(self):
        save_entry(number=7, year=2025, booking_date=date(2025, 6, 1), credit=Decimal("900"))
        response = self.client.get(reverse("finansije:dashboard"), {"date_from": "2025-01-01", "date_to": "2025-12-31"})
        self.assertEqual(response.context["totals"]["revenue"], Decimal("900"))
        self.assertEqual(response.context["chart_groups"][0]["rows"][0]["revenue"], Decimal("900"))
        self.assertEqual(response.context["rankings"][0]["rows"][0]["code"], "410001")
        empty = self.client.get(reverse("finansije:dashboard"), {"date_from": "2026-01-01", "date_to": "2026-01-31"})
        self.assertEqual(empty.context["totals"]["count"], 0)
        self.assertTrue(all(not group["rows"] for group in empty.context["chart_groups"]))
        self.assertTrue(all(not ranking["rows"] for ranking in empty.context["rankings"]))

    def test_overview_invalid_period_does_not_display_unfiltered_numbers(self):
        for changes in ({"date_from": "2024-01-01"}, {"date_from": "bad"}, {"date_from": "2026-03-01", "date_to": "2026-02-01"}, {"date_to": ""}):
            response = self.client.get(reverse("finansije:dashboard"), changes)
            self.assertEqual(response.status_code, 200)
            self.assertFalse(response.context["valid"])
            self.assertEqual(response.context["totals"]["count"], 0)
            self.assertNotContains(response, 'class="finance-overview-row"')
            self.assertContains(response, "Proverite izabrani period")

    def test_calendar_dates_and_iso_links_render_and_filter_the_same_period(self):
        for lower, upper in (("01.02.2026", "28.02.2026"), ("01.02.2026.", "28.02.2026."), ("2026-02-01", "2026-02-28")):
            for route in ("dashboard", "report", "ledger"):
                with self.subTest(route=route, lower=lower):
                    response = self.client.get(reverse(f"finansije:{route}"), dict(self.params, date_from=lower, date_to=upper))
                    self.assertEqual(response.status_code, 200)
                    self.assertTrue(response.context["valid"])
                    self.assertEqual(response.context["totals"]["revenue"], Decimal("200"))
                    self.assertContains(response, 'value="01.02.2026"')
                    self.assertContains(response, 'value="28.02.2026"')
                    self.assertContains(response, 'class="form-control js-date"')
                    self.assertContains(response, "od 01.02.2026. do 28.02.2026.")
                    self.assertIn("date_from=2026-02-01", response.context["params"])
        response = self.client.get(reverse("finansije:dashboard"), {"date_from": "01.01.2026", "date_to": "31.01.2026"})
        self.assertEqual(response.context["totals"]["count"], 0)
        self.assertContains(response, "od 01.01.2026. do 31.01.2026.")

    def test_localized_dates_reject_impossible_or_reversed_ranges(self):
        for lower, upper in (("31.02.2026", "28.02.2026"), ("01.03.2026", "28.02.2026")):
            response = self.client.get(reverse("finansije:dashboard"), {"date_from": lower, "date_to": upper})
            self.assertFalse(response.context["valid"])
            self.assertNotContains(response, 'class="finance-overview-row"')
            self.assertNotContains(response, 'class="finance-analysis-period"')
        self.assertNotContains(response, "Enter a valid date")

    def test_closing_and_result_transfers_do_not_cancel_analytical_revenue(self):
        closing = [
            ("ZAT", "61000", "200", "0"), ("ZAT", "51000", "0", "30"),
            ("ZAT", "69900", "0", "200"), ("ZAT", "59900", "30", "0"),
            ("ON", "69900", "200", "0"), ("ON", "59900", "0", "30"),
        ]
        for number, (kind, account, debit, credit) in enumerate(closing, 7):
            save_entry(number=number, journal_type=kind, account=account, debit=Decimal(debit), credit=Decimal(credit), center="3", job_code="111111")
        for route in ("dashboard", "report"):
            response = self.client.get(reverse(f"finansije:{route}"), self.params)
            self.assertEqual(response.context["totals"]["revenue"], Decimal("200"))
            self.assertEqual(response.context["totals"]["expense"], Decimal("30"))
            self.assertEqual(response.context["totals"]["count"], 5)
        overview = self.client.get(reverse("finansije:dashboard"), self.params)
        self.assertNotIn("3", [row["code"] for row in overview.context["chart_groups"][0]["rows"]])
        shares = [row["metrics"][0]["share"] for row in overview.context["chart_groups"][0]["rows"]]
        self.assertEqual(sum(shares), Decimal("100"))
        ledger = self.client.get(reverse("finansije:ledger"), dict(self.params, kind="all"))
        self.assertEqual(ledger.context["page_obj"].paginator.count, 12)
        self.assertEqual(ledger.context["totals"]["revenue"], Decimal("200"))
        self.assertTrue(any(row.journal_type == "ZAT" for row in ledger.context["page_obj"]))
        export = self.client.get(reverse("finansije:export"), self.params)
        sheet = load_workbook(BytesIO(export.content)).active
        self.assertEqual(sheet.max_row, 5)
        self.assertEqual(sum(sheet.cell(row, 4).value for row in (4, 5)), 200)
        self.assertEqual(LedgerEntry.objects.count(), 12)

    def test_regular_on_postings_and_prior_year_corrections_remain_included(self):
        save_entry(number=7, journal_type="ON", account="59120", debit=Decimal("7"), credit=Decimal("0"))
        save_entry(number=8, journal_type="ON", account="69120", debit=Decimal("0"), credit=Decimal("5"))
        response = self.client.get(reverse("finansije:dashboard"), self.params)
        self.assertEqual(response.context["totals"]["revenue"], Decimal("205"))
        self.assertEqual(response.context["totals"]["expense"], Decimal("37"))

    def test_overview_negative_values_keep_sign_on_percentage_scale(self):
        LedgerEntry.objects.filter(journal_number=5).update(credit=Decimal("0"), debit=Decimal("20"))
        response = self.client.get(reverse("finansije:dashboard"))
        group = response.context["chart_groups"][0]
        negative = next(row for row in group["rows"] if row["code"] == "42")
        self.assertEqual(negative["result"], Decimal("-20"))
        self.assertTrue(negative["metrics"][2]["negative"])
        for row in group["rows"]:
            for metric in row["metrics"]:
                width = Decimal(metric["width"])
                self.assertGreaterEqual(width, 0)
                self.assertLessEqual(width, Decimal("100"))
                self.assertNotIn(",", metric["width"])
        self.assertNotContains(response, 'class="finance-chart-zero"')
        self.assertEqual(negative["metrics"][0]["width"], "33.3333")
        self.assertEqual(negative["metrics"][2]["width"], "66.6667")

    def test_overview_zero_total_has_values_without_percentage(self):
        LedgerEntry.objects.filter(journal_number=5).update(credit=Decimal("0"), debit=Decimal("80"))
        response = self.client.get(reverse("finansije:dashboard"))
        group = response.context["chart_groups"][0]
        self.assertEqual(response.context["totals"]["revenue"], 0)
        self.assertEqual(len(group["rows"]), 2)
        self.assertTrue(all(row["metrics"][0]["share"] is None for row in group["rows"]))
        self.assertTrue(all(Decimal(row["metrics"][0]["width"]) == 0 for row in group["rows"]))

    def test_chart_bar_fill_matches_own_metric_percentage_including_net(self):
        response = self.client.get(reverse("finansije:dashboard"), self.params)
        for group in response.context["chart_groups"]:
            best = group["rows"][0]
            self.assertEqual([metric["width"] for metric in best["metrics"]], ["60.0000", "0.0000", "70.5882"])
            self.assertEqual(best["metrics"][0]["share"], Decimal("60"))
            self.assertAlmostEqual(best["metrics"][2]["share"], Decimal("70.5882352941"), places=8)

    def test_net_share_above_100_keeps_actual_percentage_and_marks_capped_bar(self):
        save_entry(number=7, account="51000", debit=Decimal("100"), credit=Decimal("0"))
        response = self.client.get(reverse("finansije:dashboard"), self.params)
        for group in response.context["chart_groups"]:
            best, worst = group["rows"]
            net = best["metrics"][2]
            self.assertEqual(net["width"], "100.0000")
            self.assertAlmostEqual(net["share"], Decimal("171.4285714286"), places=8)
            self.assertTrue(net["overflow"])
            self.assertTrue(group["has_overflow"])
            self.assertEqual(worst["metrics"][2]["width"], "71.4286")
        self.assertContains(response, "finance-chart-overflow-marker")
        save_entry(number=8, account="51000", debit=Decimal("70"), credit=Decimal("0"))
        response = self.client.get(reverse("finansije:dashboard"), self.params)
        self.assertEqual(response.context["totals"]["result"], 0)
        for row in response.context["chart_groups"][0]["rows"]:
            self.assertIsNone(row["metrics"][2]["share"])
            self.assertEqual(row["metrics"][2]["width"], "0.0000")

    def test_overview_expands_all_metrics_and_ranks_by_net_result(self):
        for number in range(10, 20):
            save_entry(number=number, job_code=str(number), job_name=f"Posao {number}", credit=Decimal(str(number)))
        save_entry(number=20, job_code="", job_name="", center="", credit=Decimal("1000"))
        save_entry(number=21, job_code="loss", job_name="Gubitak", center="43", account="51000", debit=Decimal("90"), credit=Decimal("0"))
        response = self.client.get(reverse("finansije:dashboard"))
        centers, jobs = response.context["chart_groups"]
        self.assertEqual(len(jobs["visible"]), 8)
        self.assertEqual(len(jobs["remaining"]), 6)
        self.assertTrue(all(len(row["metrics"]) == 3 for row in jobs["remaining"]))
        self.assertEqual(centers["remaining"], [])
        self.assertContains(response, 'class="finance-overview-expand"', count=1)
        self.assertContains(response, "Proširi sve šifre (14)")
        rankings = response.context["rankings"]
        self.assertEqual([row["code"] for row in rankings[0]["rows"]], ["420001", "410001", "19"])
        self.assertEqual([row["code"] for row in rankings[1]["rows"]], ["loss", "10", "11"])
        self.assertEqual(rankings[2]["rows"][0]["code"], "41")
        self.assertEqual(rankings[3]["rows"][0]["code"], "43")
        for index, ranking in enumerate(rankings):
            for row in ranking["rows"]:
                self.assertTrue(row["url"].startswith(reverse("finansije:job_card" if index < 2 else "finansije:report")))
        missing = next(row for row in jobs["rows"] if not row["code"])
        detail = self.client.get(missing["url"])
        self.assertEqual(detail.context["totals"]["revenue"], Decimal("1000"))

    def test_month_report_and_detail_drilldown(self):
        response = self.client.get(reverse("finansije:report"), dict(self.params, group="month"))
        row = response.context["rows"][0]
        self.assertEqual(row["code"], "2026-02")
        detail = self.client.get(row["ledger_url"])
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.context["page_obj"].paginator.count, 5)

    def test_account_link_matches_exact_account(self):
        save_entry(number=7, account="610001", credit=Decimal("777"))
        response = self.client.get(reverse("finansije:report"), dict(self.params, group="account"))
        row = next(r for r in response.context["rows"] if r["code"] == "61000")
        detail = self.client.get(row["ledger_url"])
        self.assertEqual(detail.context["totals"]["revenue"], Decimal("200"))

    def test_jobs_include_empty_and_removed_kind_filter_cannot_hide_revenue(self):
        response = self.client.get(reverse("finansije:report"), dict(self.params, group="job", include_empty="1", kind="expense"))
        self.assertEqual(len(response.context["rows"]), 3)
        self.assertEqual(response.context["totals"]["revenue"], Decimal("200"))
        response = self.client.get(reverse("finansije:report"), dict(self.params, group="account", kind="expense"))
        self.assertTrue(all(r["revenue_share"] is None for r in response.context["rows"]))

    def test_invalid_period_and_account_do_not_show_unfiltered_results(self):
        for changes in ({"date_from": "2026-03-01"}, {"date_from": "2024-01-01"}, {"account": "5';DELETE"}):
            response = self.client.get(reverse("finansije:report"), dict(self.params, **changes))
            self.assertFalse(response.context["valid"])
            self.assertEqual(response.context["totals"]["count"], 0)

    def test_all_postings_option_keeps_pnl_separate_from_turnover(self):
        response = self.client.get(reverse("finansije:ledger"), dict(self.params, kind="all"))
        self.assertEqual(response.context["page_obj"].paginator.count, 6)
        self.assertEqual(response.context["totals"]["revenue"], Decimal("200"))
        self.assertEqual(response.context["movements"]["debit"], Decimal("560"))

    def test_ledger_datatable_sorts_dates_across_months_and_years_before_paging(self):
        save_entry(number=7, year=2025, booking_date=date(2025, 12, 31), document_date=date(2026, 3, 1))
        save_entry(number=8, booking_date=date(2026, 1, 2), document_date=date(2025, 12, 31))
        params = dict(self.params, date_from="2025-01-01", date_to="2026-12-31", draw="1", start="0", length="1")
        params.update({"order[0][column]": "0", "order[0][dir]": "asc"})
        data = self.client.get(reverse("finansije:ledger"), params).json()
        self.assertEqual(data["recordsTotal"], 7)
        self.assertEqual(data["data"][0][0], "31.12.2025.")
        data = self.client.get(reverse("finansije:ledger"), dict(params, start="1")).json()
        self.assertEqual(data["data"][0][0], "02.01.2026.")
        params.update({"order[0][column]": "8", "order[0][dir]": "desc"})
        data = self.client.get(reverse("finansije:ledger"), params).json()
        self.assertEqual(data["data"][0][8], "01.03.2026.")
        params.update({"search[value]": "31.12.2025."})
        data = self.client.get(reverse("finansije:ledger"), params).json()
        self.assertEqual(data["recordsFiltered"], 2)

    def test_ledger_datatable_searches_all_rows_and_sorts_amounts_numerically(self):
        for number, amount in ((7, "2"), (8, "10"), (9, "100")):
            save_entry(number=number, account="51000", debit=Decimal(amount), credit=Decimal("0"), partner_name="Sorting probe")
        params = dict(self.params, draw="2", length="2", start="1")
        params.update({"search[value]": "Sorting probe", "order[0][column]": "9", "order[0][dir]": "asc"})
        data = self.client.get(reverse("finansije:ledger"), params).json()
        self.assertEqual(data["draw"], 2)
        self.assertEqual(data["recordsTotal"], 8)
        self.assertEqual(data["recordsFiltered"], 3)
        self.assertEqual(len(data["data"]), 2)
        self.assertIn("2026/IF/8", data["data"][0][1])
        self.assertIn("2026/IF/9", data["data"][1][1])

    def test_ledger_datatable_escapes_source_text_and_bounds_input(self):
        LedgerEntry.objects.filter(journal_number=1).update(partner_name='<img src=x onerror="alert(1)">', debit=Decimal("-15"))
        params = dict(self.params, draw="<script>", start="-3", length="999999")
        params.update({"order[0][column]": "account;DROP TABLE", "search[value]": "<img"})
        data = self.client.get(reverse("finansije:ledger"), params).json()
        self.assertEqual(data["draw"], 0)
        self.assertEqual(len(data["data"]), 1)
        self.assertNotIn("<img", data["data"][0][6])
        self.assertIn("&lt;img", data["data"][0][6])
        self.assertIn("finance-amount-negative", data["data"][0][9])
        invalid = self.client.get(reverse("finansije:ledger"), dict(params, date_from="2024-01-01")).json()
        self.assertEqual(invalid["recordsTotal"], 0)

    def test_sync_datatable_orders_real_timestamps_and_keeps_error_in_same_row(self):
        first = SyncRun.objects.create(company=1, year_from=2025, year_to=2025, status="success")
        second = SyncRun.objects.create(company=1, year_from=2026, year_to=2026, status="failed", error="<script>bad</script>")
        SyncRun.objects.filter(pk=first.pk).update(started_at=timezone.make_aware(datetime(2025, 12, 31, 10)))
        SyncRun.objects.filter(pk=second.pk).update(started_at=timezone.make_aware(datetime(2026, 1, 2, 10)))
        SyncRun.objects.create(company=2, year_from=2026, year_to=2026, status="success")
        params = {"draw": "1", "length": "1", "order[0][column]": "0", "order[0][dir]": "asc"}
        data = self.client.get(reverse("finansije:sync_status"), params).json()
        self.assertEqual(data["recordsTotal"], 2)
        self.assertTrue(data["data"][0][0].startswith("31.12.2025."))
        data = self.client.get(reverse("finansije:sync_status"), dict(params, **{"search[value]": "Neuspešno"})).json()
        self.assertEqual(data["recordsFiltered"], 1)
        self.assertEqual(len(data["data"][0]), 9)
        self.assertIn("&lt;script&gt;", data["data"][0][3])

    def test_successful_sync_banner_is_absent_from_all_reports(self):
        SyncRun.objects.create(company=1, year_from=2025, year_to=2026, status="success", finished_at=timezone.now())
        for route in ("dashboard", "report", "ledger"):
            response = self.client.get(reverse(f"finansije:{route}"), self.params)
            self.assertNotContains(response, "Poslednja uspešna sinhronizacija")
            self.assertNotContains(response, "Istorija sinhronizacije</a>")

    def test_excel_amounts_are_numeric_and_external_text_is_not_formula(self):
        LedgerEntry.objects.filter(journal_number=1).update(partner_name="=1+1")
        response = self.client.get(reverse("finansije:export"), dict(self.params, report="ledger"))
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content))
        sheet = book.active
        self.assertEqual(sheet.cell(4, 13).value, "=1+1")
        self.assertEqual(sheet.cell(4, 13).data_type, "s")
        self.assertEqual(sheet.cell(4, 17).data_type, "n")

    def test_removed_rows_are_not_reported(self):
        LedgerEntry.objects.filter(journal_number=5).update(active=False)
        response = self.client.get(reverse("finansije:report"), self.params)
        self.assertEqual(response.context["totals"]["revenue"], Decimal("80"))

    def test_incomplete_years_are_visible(self):
        SyncRun.objects.create(company=1, year_from=2026, year_to=2026, status="success", finished_at=timezone.now())
        response = self.client.get(reverse("finansije:report"), dict(self.params, date_from="2025-01-01"))
        self.assertEqual(response.context["missing_years"], [2025])

    def test_first_import_shows_progress_without_reading_unpublished_rows(self):
        SyncRun.objects.create(company=1, year_from=2025, year_to=2026, status="running")
        response = self.client.get(reverse("finansije:report"), self.params)
        self.assertContains(response, "Sinhronizacija je u toku")
        self.assertEqual(response.context["totals"]["count"], 0)

    def test_scope_applies_to_totals_choices_links_and_export(self):
        user = get_user_model().objects.create_user("finance-limited", password="test", allowed_center_codes=" 41 ; ")
        role = Role.objects.create(name="Finansije test", slug="finance-test")
        for name in ("dashboard", "ledger", "export", "sync_status"):
            role.permissions.add(PermissionCode.objects.create(code=f"finansije:{name}"))
        user.roles.add(role)
        self.client.force_login(user)
        response = self.client.get(reverse("finansije:report"), self.params)
        self.assertEqual(response.context["totals"]["revenue"], Decimal("80"))
        self.assertNotContains(response, "Posao B")
        overview = self.client.get(reverse("finansije:dashboard"))
        self.assertNotContains(overview, "Posao B")
        for group in overview.context["chart_groups"]:
            self.assertEqual(len(group["rows"]), 1)
            self.assertEqual(group["rows"][0]["metrics"][0]["share"], Decimal("100"))
        self.assertEqual(overview.context["rankings"][0]["rows"][0]["code"], "410001")
        self.assertNotIn(("42", "42"), response.context["form"].fields["center"].choices)
        data = self.client.get(reverse("finansije:ledger"), dict(self.params, draw="1", **{"search[value]": "Posao B"})).json()
        self.assertEqual(data["recordsTotal"], 4)
        self.assertEqual(data["recordsFiltered"], 0)
        self.assertEqual(data["data"], [])
        self.assertEqual(self.client.get(reverse("finansije:sync_status"), {"draw": "1"}).status_code, 403)
        forged = self.client.get(reverse("finansije:ledger"), dict(self.params, center="42"))
        self.assertEqual(forged.context["page_obj"].paginator.count, 0)
        export = self.client.get(reverse("finansije:export"), self.params)
        sheet = load_workbook(BytesIO(export.content)).active
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet.cell(4, 4).value, 80)
        self.assertEqual(self.client.get(reverse("finansije:sync_status")).status_code, 403)
        user.allowed_center_codes = ""
        user.save()
        response = self.client.get(reverse("finansije:report"), self.params)
        self.assertEqual(response.context["totals"]["count"], 0)

    def test_m2m_centers_with_source_padding_are_supported(self):
        user = get_user_model().objects.create_user("finance-unit", password="test")
        unit = OrganizationalUnit.objects.create(code="410001 ", name="Unit", center="41   ")
        user.allowed_centers.add(unit)
        entries, _ = base_querysets(user)
        self.assertEqual(summary(entries)["revenue"], Decimal("80"))

    def test_no_role_is_denied_and_anonymous_redirects(self):
        user = get_user_model().objects.create_user("finance-no-role", password="test")
        self.client.force_login(user)
        for route in ("dashboard", "report", "ledger", "export", "sync_status"):
            self.assertEqual(self.client.get(reverse(f"finansije:{route}")).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(reverse("finansije:report")).status_code, 302)

    def test_view_all_permission_allows_complete_report(self):
        user = get_user_model().objects.create_user("finance-management", password="test")
        role = Role.objects.create(name="Uprava test", slug="finance-management")
        for name in ("dashboard", "view_all"):
            role.permissions.add(PermissionCode.objects.create(code=f"finansije:{name}"))
        user.roles.add(role)
        self.client.force_login(user)
        response = self.client.get(reverse("finansije:report"), self.params)
        self.assertEqual(response.context["totals"]["revenue"], Decimal("200"))
