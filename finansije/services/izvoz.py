"""Izvoz ekrana Finansija u Excel i PDF (od 05.10.2026.).

Svaki ekran se opiše kao dokument: naslov, period, sažetak i tabele. Isti dokument se upisuje u
Excel (jedan list, tabele jedna ispod druge, brojevi ostaju brojevi) ili se prikazuje kao A4 strana
za štampu („Sačuvaj kao PDF”), kao spisak zaposlenih u Kadrovima. Brojevi se ne preračunavaju:
ovde dolaze vrednosti koje ekran već prikazuje.
"""
from dataclasses import dataclass, field
from datetime import date, datetime

from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from django.utils.formats import number_format
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INSTITUT = "Institut za ispitivanje materijala a.d."
ADRESA = "Bulevar vojvode Mišića 43, 11000 Beograd"
PLAVA = "14395B"
IZNOS = "#,##0.00;[Red]-#,##0.00;0.00"
PROCENAT = '0.00"%";[Red]-0.00"%";0.00"%"'


@dataclass
class Kolona:
    naziv: str
    vrsta: str = "tekst"  # tekst, iznos, procenat, broj, datum

    @property
    def broj(self):
        return self.vrsta in ("iznos", "procenat", "broj")


@dataclass
class Tabela:
    naslov: str
    kolone: list
    redovi: list
    zbir: list = None
    opis: str = ""


@dataclass
class Dokument:
    naslov: str
    period: str = ""
    filteri: str = ""
    napomena: str = ""
    sazetak: list = field(default_factory=list)  # [(naziv, vrednost, vrsta)]
    tabele: list = field(default_factory=list)


def period_tekst(od, do):
    return f"Period od {od:%d.%m.%Y.} do {do:%d.%m.%Y.}"


def prikaz(vrednost, vrsta):
    """Tekst za štampu: iznosi sa razdvajanjem hiljada, crta za nepoznato."""
    if vrednost is None or vrednost == "":
        return "—"
    if vrsta == "iznos":
        return number_format(vrednost, decimal_pos=2, use_l10n=True, force_grouping=True)
    if vrsta == "procenat":
        return number_format(vrednost, decimal_pos=2, use_l10n=True, force_grouping=True) + "%"
    if vrsta == "broj":
        return number_format(vrednost, decimal_pos=0, use_l10n=True, force_grouping=True)
    if vrsta == "datum" and isinstance(vrednost, (date, datetime)):
        return vrednost.strftime("%d.%m.%Y.")
    return str(vrednost)


def klasa(vrednost, kolona):
    if not kolona.broj:
        return ""
    return "n neg" if vrednost is not None and vrednost != "" and vrednost < 0 else "n"


# --- Excel ---------------------------------------------------------------------------------------

def _excel_vrednost(vrednost, vrsta):
    if vrednost is None:
        return None
    if vrsta in ("iznos", "procenat", "broj"):
        return float(vrednost) if not isinstance(vrednost, int) else vrednost
    if vrsta == "datum" and isinstance(vrednost, (date, datetime)):
        return vrednost
    return str(vrednost)


def excel_odgovor(dokument, ime_fajla):
    knjiga = Workbook()
    list_ = knjiga.active
    list_.title = "Izveštaj"
    list_.sheet_view.showGridLines = False
    list_.page_setup.orientation = "landscape"
    list_.page_setup.paperSize = list_.PAPERSIZE_A4
    list_.page_setup.fitToWidth, list_.page_setup.fitToHeight = 1, 0
    list_.sheet_properties.pageSetUpPr.fitToPage = True
    tanka = Side(style="thin", color="D7DCE5")
    sirine = {}

    def upisi(red, kolona, vrednost, *, vrsta="tekst", font=None, fill=None, desno=False, granica=None):
        celija = list_.cell(red, kolona, _excel_vrednost(vrednost, vrsta))
        if isinstance(celija.value, str):
            celija.data_type = "s"  # spoljašnji tekst nikad ne postaje formula
        if vrsta == "iznos":
            celija.number_format = IZNOS
        elif vrsta == "procenat":
            celija.number_format = PROCENAT
        elif vrsta == "datum":
            celija.number_format = "DD.MM.YYYY"
        if font:
            celija.font = font
        if fill:
            celija.fill = PatternFill("solid", fgColor=fill)
        celija.alignment = Alignment(horizontal="right" if desno else "left", vertical="top",
                                     wrap_text=isinstance(celija.value, str) and len(celija.value) > 60)
        if granica:
            celija.border = granica
        duzina = len(prikaz(vrednost, vrsta)) if vrednost is not None else 1
        sirine[kolona] = max(sirine.get(kolona, 8), min(duzina + 2, 60))
        return celija

    upisi(1, 1, f"{INSTITUT} · {dokument.naslov}", font=Font(bold=True, size=14, color=PLAVA))
    red = 2
    for tekst in (dokument.period, dokument.filteri, dokument.napomena):
        if tekst:
            list_.cell(red, 1, tekst).font = Font(italic=True, color="475569")
            list_.cell(red, 1).data_type = "s"
            red += 1
    list_.cell(red, 1, f"Izrađeno: {timezone.localtime():%d.%m.%Y. %H:%M}").font = Font(color="64748B", size=9)
    red += 2
    if dokument.sazetak:
        for naziv, vrednost, vrsta in dokument.sazetak:
            upisi(red, 1, naziv, font=Font(bold=True))
            upisi(red, 2, vrednost, vrsta=vrsta, desno=True)
            red += 1
        red += 1
    for tabela in dokument.tabele:
        upisi(red, 1, tabela.naslov, font=Font(bold=True, size=12, color=PLAVA))
        red += 1
        if tabela.opis:
            list_.cell(red, 1, tabela.opis).font = Font(italic=True, size=9, color="64748B")
            list_.cell(red, 1).data_type = "s"
            red += 1
        for i, kolona in enumerate(tabela.kolone, 1):
            upisi(red, i, kolona.naziv, font=Font(bold=True, color="FFFFFF"), fill=PLAVA, desno=kolona.broj)
        red += 1
        for vrednosti in tabela.redovi:
            for i, (kolona, vrednost) in enumerate(zip(tabela.kolone, vrednosti), 1):
                upisi(red, i, vrednost, vrsta=kolona.vrsta, desno=kolona.broj, granica=Border(bottom=tanka))
            red += 1
        if not tabela.redovi:
            upisi(red, 1, "Nema podataka za izabrani period.", font=Font(italic=True, color="64748B"))
            red += 1
        if tabela.zbir:
            for i, (kolona, vrednost) in enumerate(zip(tabela.kolone, tabela.zbir), 1):
                upisi(red, i, vrednost, vrsta=kolona.vrsta, desno=kolona.broj, font=Font(bold=True), fill="F1F5F9",
                      granica=Border(top=Side(style="medium", color=PLAVA)))
            red += 1
        red += 1
    for kolona, sirina in sirine.items():
        list_.column_dimensions[get_column_letter(kolona)].width = sirina
    list_.column_dimensions["A"].width = min(max(sirine.get(1, 12), 14), 34)
    odgovor = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    odgovor["Content-Disposition"] = f'attachment; filename="{ime_fajla}.xlsx"'
    knjiga.save(odgovor)
    return odgovor


# --- PDF (A4 strana za štampu) ---------------------------------------------------------------------

def pdf_odgovor(request, dokument):
    tabele = []
    for tabela in dokument.tabele:
        tabele.append({
            "naslov": tabela.naslov, "opis": tabela.opis,
            "kolone": tabela.kolone,
            "redovi": [[(prikaz(v, k.vrsta), klasa(v, k)) for k, v in zip(tabela.kolone, red)] for red in tabela.redovi],
            "zbir": [(prikaz(v, k.vrsta) if v not in (None, "") or k.broj else "", klasa(v, k))
                     for k, v in zip(tabela.kolone, tabela.zbir)] if tabela.zbir else None,
        })
    sazetak = [(naziv, prikaz(v, vrsta), "neg" if vrsta != "tekst" and v is not None and v < 0 else "")
               for naziv, v, vrsta in dokument.sazetak]
    najvise_kolona = max((len(t.kolone) for t in dokument.tabele), default=0)
    return render(request, "finansije/izvoz_stampa.html", {
        "dokument": dokument, "tabele": tabele, "sazetak": sazetak,
        "polozeno": najvise_kolona > 7, "sitno": najvise_kolona > 11,
        "institut": INSTITUT, "adresa": ADRESA, "izradjeno": timezone.localtime(),
    })


def odgovor(request, dokument, ime_fajla):
    if request.GET.get("format") == "pdf":
        return pdf_odgovor(request, dokument)
    return excel_odgovor(dokument, ime_fajla)
