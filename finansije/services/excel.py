"""Presentation of the existing job report values as a native Excel table."""
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo


NUMBER_FORMAT = '#,##0.00;[Red]-#,##0.00;0.00'


def prepare_job_sheet(sheet, headers):
    # Keep the report identifiers visible while scrolling across the metrics.
    sheet.freeze_panes = 'D4'
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.print_title_rows = '1:3'
    sheet.row_dimensions[3].height = 48
    for index in range(1, len(headers) + 1):
        width = {1: 16, 2: 42, 3: 14, len(headers)-1: 60, len(headers): 45}.get(index, 23)
        sheet.column_dimensions[get_column_letter(index)].width = width


def job_excel_row(sheet, values, headers, *, heading=False):
    cells = []
    for index, value in enumerate(values):
        item = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            item.data_type = 's'  # Keep external text literal, including leading '='.
        if heading:
            item.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
            item.fill = PatternFill('solid', fgColor='235B83')
            item.alignment = Alignment(vertical='center', wrap_text=True)
        else:
            item.font = Font(name='Calibri', size=11)
            numeric = 3 <= index < len(headers)-2
            item.alignment = Alignment(vertical='top', horizontal='right' if numeric else 'left',
                                       wrap_text=index in (1, len(headers)-2))
            if numeric:
                # Ratios are already stored as percentage points (50 = 50%).
                item.number_format = '0.00"%";[Red]-0.00"%";0.00"%"' if headers[index].endswith('%') else NUMBER_FORMAT
            elif index < 3:
                item.number_format = '@'
        cells.append(item)
    return cells


def finish_job_table(sheet, headers, row_count, *, additional=False):
    reference = f'A3:{get_column_letter(len(headers))}{3 + row_count}'
    table = Table(displayName='DodatneAnalize' if additional else 'SifrePosla', ref=reference,
                  tableColumns=[TableColumn(id=i, name=label) for i, label in enumerate(headers, 1)],
                  autoFilter=AutoFilter(ref=reference),
                  tableStyleInfo=TableStyleInfo(name='TableStyleMedium2', showRowStripes=True))
    # Match table metadata to the visible header, including the additional metrics.
    sheet.add_table(table)
