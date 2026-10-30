import re
from datetime import date
from decimal import Decimal as D
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import PermissionCode, Role
from .services import bilansi
from .services.bilansi import DOBITAK_TEKUCE, GUBITAK_TEKUCE, izracunaj, salda_stanja, salda_uspeha
from .services.bilansi_obrazac import NAPOMENE, NENEGATIVNE, POZICIJE
from .tests import save_entry


class ObrazacTests(SimpleTestCase):
    def test_formulas_reference_positions_of_the_same_report(self):
        for p in POZICIJE:
            for ref in re.findall(r"\d{4}", p.formula):
                self.assertIn(ref, {q.aop for q in POZICIJE if q.izvestaj == p.izvestaj}, p.aop)
        self.assertTrue(NENEGATIVNE <= {p.aop for p in POZICIJE})

    def test_each_account_has_one_position_except_profit_loss_pairs(self):
        vidjeno = {}
        for p in POZICIJE:
            for oznaka in p.konta.split():
                kljuc = (p.izvestaj, oznaka)
                if kljuc in vidjeno:
                    self.assertIn({vidjeno[kljuc], p.aop}, ({"1047", "1048"}, {"1052", "1053"}), oznaka)
                vidjeno[kljuc] = p.aop

    def test_every_empty_position_is_explained(self):
        for p in POZICIJE:
            if not p.konta and not p.formula:
                self.assertIn(p.aop, NAPOMENE)
        self.assertTrue(set(NAPOMENE) <= {p.aop for p in POZICIJE})

    def test_deo_accounts_follow_decisions(self):
        def gde(izv, konto, saldo=D("1")):
            return [p.aop for p in bilansi.rasporedi(konto, saldo, bilansi.pravila(izv))]
        self.assertEqual(gde("BS", "02920"), ["0015"])
        self.assertEqual(gde("BS", "23661"), ["0054"])
        self.assertEqual(gde("BS", "28004"), ["0058"])
        self.assertEqual(gde("BS", "28800"), ["0029"])
        self.assertEqual(gde("BS", "49700"), ["0454"])
        self.assertEqual(gde("BS", "49800"), ["0429"])
        self.assertEqual(gde("BS", "42500"), ["0437"])
        self.assertEqual(gde("BS", "33100", D("5")), ["0407"])
        self.assertEqual(gde("BS", "33100", D("-5")), ["0406"])
        self.assertEqual(gde("BU", "59120"), ["1047", "1048"])


class IzracunTests(SimpleTestCase):
    def test_balance_sheet_sides_and_totals(self):
        r = izracunaj("BS", {"02200": D("1000"), "24100": D("300"), "43500": D("-200"), "30000": D("-1100"),
                             "88000": D("7"), "89000": D("-7")})
        i = r["iznosi"]
        self.assertEqual((i["0010"], i["0057"], i["0445"], i["0402"]), (D("1000"), D("300"), D("200"), D("1100")))
        self.assertEqual((i["0002"], i["0030"], i["0059"], i["0456"]), (D("1000"), D("300"), D("1300"), D("1300")))
        self.assertEqual((i["0060"], i["0457"]), (D("7"), D("7")))
        self.assertEqual(r["konta"]["0010"], [("02200", D("1000"))])

    def test_wrong_side_balance_stays_negative_and_keeps_the_balance(self):
        # Dobavljač u inostranstvu sa dugovnim saldom: pozicija i zbir su negativni, aktiva = pasiva.
        i = izracunaj("BS", {"24100": D("100"), "43600": D("40"), "30000": D("-140")})["iznosi"]
        self.assertEqual((i["0446"], i["0442"], i["0431"]), (D("-40"), D("-40"), D("-40")))
        self.assertEqual(i["0059"], i["0456"])

    def test_negative_equity_goes_to_loss_above_equity(self):
        i = izracunaj("BS", {"24100": D("100"), "43500": D("-150"), "30000": D("-10"), GUBITAK_TEKUCE: D("60")})["iznosi"]
        self.assertEqual((i["0414"], i["0401"], i["0455"]), (D("60"), D("0"), D("50")))
        self.assertEqual((i["0059"], i["0456"]), (D("100"), D("100")))

    def test_income_statement_pairs_never_negative(self):
        i = izracunaj("BU", {"61200": D("-500"), "52000": D("700"), "66200": D("-30"), "59120": D("50"), "69120": D("-20"),
                             "721": D("0"), "722": D("-15")})["iznosi"]
        self.assertEqual((i["1001"], i["1013"], i["1025"], i["1026"]), (D("500"), D("700"), D("0"), D("200")))
        self.assertEqual((i["1037"], i["1038"], i["1047"], i["1048"]), (D("30"), D("0"), D("0"), D("30")))
        self.assertEqual((i["1046"], i["1050"]), (D("170"), D("200")))
        self.assertEqual((i["1052"], i["1053"], i["1055"], i["1056"]), (D("0"), D("15"), D("0"), D("185")))

    def test_unmapped_accounts_are_reported(self):
        self.assertEqual(izracunaj("BS", {"29900": D("5")})["nerasporedjeno"], {"29900": D("5")})


def knjizi(broj, konto, dug="0", pot="0", vrsta="ON", dan=date(2026, 3, 1), linija=1):
    return save_entry(number=broj, line_number=linija, journal_type=vrsta, account=konto, debit=D(dug), credit=D(pot),
                      booking_date=dan, year=dan.year)


class KnjizenjaTests(TestCase):
    """Zaključak godine kao u Stankovom programu: ZAT prenosi klase 5/6 na 59900/69900, a nalog ON od 31.12.
    zatvara 59900/69900 kroz klasu 7, knjiži porez (721/481) i odloženi porez (722/498) i prenosi dobitak na 341."""

    def setUp(self):
        knjizi(1, "24100", dug="1000", vrsta="POC", dan=date(2025, 1, 1))
        knjizi(1, "30000", pot="1000", vrsta="POC", dan=date(2025, 1, 1), linija=2)
        knjizi(2, "24100", dug="500", vrsta="IF", dan=date(2025, 5, 1))
        knjizi(2, "61200", pot="500", vrsta="IF", dan=date(2025, 5, 1), linija=2)
        knjizi(3, "52000", dug="300", vrsta="LD", dan=date(2025, 6, 1))
        knjizi(3, "24100", pot="300", vrsta="LD", dan=date(2025, 6, 1), linija=2)
        kraj = date(2025, 12, 31)
        for linija, (konto, dug, pot) in enumerate((("61200", "500", "0"), ("69900", "0", "500"), ("59900", "300", "0"),
                                                    ("52000", "0", "300")), 1):
            knjizi(1, konto, dug, pot, vrsta="ZAT", dan=kraj, linija=linija)
        for linija, (konto, dug, pot) in enumerate((
                ("69900", "500", "0"), ("59900", "0", "300"), ("72000", "0", "200"),  # dobitak pre oporezivanja 200
                ("72100", "30", "0"), ("48100", "0", "30"),  # porez 30
                ("49800", "10", "0"), ("72200", "0", "10"),  # odloženi poreski prihod 10
                ("72000", "200", "0"), ("72100", "0", "30"), ("72200", "10", "0"), ("72400", "0", "180"),
                ("72400", "180", "0"), ("34100", "0", "180")), 1):
            knjizi(9, konto, dug, pot, dan=kraj, linija=linija)
        knjizi(5, "49800", pot="10", vrsta="POC", dan=date(2025, 1, 1))
        knjizi(5, "30000", dug="10", vrsta="POC", dan=date(2025, 1, 1), linija=2)

    def test_income_statement_of_closed_year(self):
        s = salda_uspeha(date(2025, 1, 1), date(2025, 12, 31))
        self.assertEqual((s["61200"], s["52000"], s["721"], s["722"]), (D("-500"), D("300"), D("30"), D("-10")))
        self.assertNotIn("59900", s)
        i = izracunaj("BU", s)["iznosi"]
        self.assertEqual((i["1049"], i["1051"], i["1053"], i["1055"]), (D("200"), D("30"), D("10"), D("180")))

    def test_income_statement_before_closing_has_no_tax(self):
        i = izracunaj("BU", salda_uspeha(date(2025, 1, 1), date(2025, 6, 30)))["iznosi"]
        self.assertEqual((i["1049"], i["1051"], i["1055"]), (D("200"), D("0"), D("200")))

    def test_balance_sheet_during_and_after_the_year(self):
        tokom = salda_stanja(date(2025, 6, 30))
        self.assertEqual(tokom[DOBITAK_TEKUCE], D("-200"))
        i = izracunaj("BS", tokom)["iznosi"]
        self.assertEqual((i["0057"], i["0410"], i["0059"], i["0456"]), (D("1200"), D("200"), D("1200"), D("1200")))
        posle = salda_stanja(date(2025, 12, 31))
        self.assertNotIn(DOBITAK_TEKUCE, posle)
        i = izracunaj("BS", posle)["iznosi"]
        self.assertEqual((i["0410"], i["0452"], i["0429"], i["0059"], i["0456"]), (D("180"), D("30"), D("0"), D("1200"), D("1200")))

    def test_comparative_year_end_from_next_year_opening(self):
        knjizi(1, "24100", dug="1200", vrsta="POC", dan=date(2026, 1, 1))
        knjizi(1, "30000", pot="1200", vrsta="POC", dan=date(2026, 1, 1), linija=2)
        # 2024. nije preuzeta: kraj godine je početno stanje 2025, a za 2023. i Bilans uspeha 2024. nema podataka.
        self.assertEqual(salda_stanja(date(2024, 12, 31)), {"24100": D("1000"), "30000": D("-990"), "49800": D("-10")})
        self.assertIsNone(salda_stanja(date(2023, 12, 31)))
        self.assertIsNone(salda_uspeha(date(2024, 1, 1), date(2024, 12, 31)))
        with patch("finansije.services.bilansi.ima_godinu", side_effect=lambda godina: godina == 2026):
            self.assertEqual(salda_stanja(date(2025, 12, 31)), {"24100": D("1200"), "30000": D("-1200")})


@patch("finansije.services.bilansi.nazivi_konta", return_value={"24100": "Tekući račun"})
class BilansViewTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("bilansi-reader", password="test")
        self.role = Role.objects.create(name="Bilansi test", slug="bilansi-test")
        self.user.roles.add(self.role)
        self.client.force_login(self.user)
        knjizi(1, "24100", dug="1000", vrsta="POC", dan=date(2026, 1, 1))
        knjizi(1, "30000", pot="1000", vrsta="POC", dan=date(2026, 1, 1), linija=2)
        knjizi(2, "24100", dug="500", vrsta="IF", dan=date(2026, 5, 1))
        knjizi(2, "61200", pot="500", vrsta="IF", dan=date(2026, 5, 1), linija=2)

    def grant(self, *codes):
        for code in codes:
            self.role.permissions.add(PermissionCode.objects.get_or_create(code=code)[0])

    def test_requires_permission_and_whole_company(self, _):
        url = reverse("finansije:bilans_stanja")
        self.assertEqual(self.client.get(url).status_code, 403)
        self.grant("finansije:bilans_stanja")
        self.assertEqual(self.client.get(url, {"na_dan": "30.06.2026", "generisi": "1"}).status_code, 403)

    def test_balance_sheet_shows_positions_accounts_and_controls(self, _):
        self.grant("finansije:bilans_stanja", "finansije:view_all")
        url = reverse("finansije:bilans_stanja")
        self.assertNotContains(self.client.get(url), "Kontrole")
        response = self.client.get(url, {"na_dan": "30.06.2026", "generisi": "1"})
        self.assertContains(response, "GOTOVINA I GOTOVINSKI EKVIVALENTI")
        self.assertContains(response, "Tekući račun")
        self.assertContains(response, "Aktiva (0059) − pasiva (0456) = 0", count=2)
        self.assertContains(response, "Na dan 31.12.2025.")  # iz početnog stanja 2026.
        self.assertContains(response, 'href="%s"' % url)

    def test_income_statement_period_and_excel(self, _):
        self.grant("finansije:bilans_uspeha", "finansije:view_all")
        url = reverse("finansije:bilans_uspeha")
        response = self.client.get(url, {"datum_od": "01.01.2026", "datum_do": "31.12.2027", "generisi": "1"})
        self.assertContains(response, "unutar jedne godine")
        response = self.client.get(url, {"datum_od": "01.01.2026", "datum_do": "30.06.2026", "izvoz": "xlsx"})
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        vrednosti = [c for r in load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True) for c in r if c is not None]
        self.assertIn("1055", vrednosti)
        self.assertIn("Ć. NETO DOBITAK", vrednosti)
        self.assertIn(500.0, vrednosti)

    def test_methodology_lists_rules_and_positions(self, _):
        url = reverse("finansije:bilansi_metodologija")
        self.assertEqual(self.client.get(url).status_code, 403)
        self.grant("finansije:bilansi_metodologija")
        response = self.client.get(url, {"izvestaj": "BU"})
        self.assertContains(response, "SPFINizv12")
        self.assertContains(response, 'id="bs-0059"')
        self.assertContains(response, '<a href="#bs-0029">0029</a>')
        self.assertContains(response, "samo potražni saldo")
        self.assertContains(response, "razgraničene obaveze za PDV")
        self.assertContains(response, 'id="metod-bu" role="tabpanel"')
        self.assertContains(response, "show active", count=1)
